"""Office Open XML local. Không cần Microsoft Office để đọc/tạo file."""
import os
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path
from .excel import digest


class OfficeTools:
    def __init__(self, files):
        self.files = files

    def check(self, path):
        p = self.files.path(path)
        if not p.is_file() or p.suffix.lower() not in {'.docx', '.pptx'}:
            raise ValueError('Chỉ hỗ trợ DOCX/PPTX; chuyển DOC/PPT cũ sang định dạng mới bằng Office.')
        if p.stat().st_size > 20 * 1024**2: raise ValueError('Office tối đa 20 MiB.')
        with zipfile.ZipFile(p) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024**2:
                raise ValueError('Tài liệu giải nén vượt 100 MiB.')
        return p

    def office_read(self, path):
        p = self.check(path)
        if p.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(p)
            lines = [para.text for para in doc.paragraphs]
            for table in doc.tables:
                lines += [' | '.join(c.text for c in r.cells) for r in table.rows]
        else:
            from pptx import Presentation
            deck = Presentation(p); lines = []
            for n, slide in enumerate(deck.slides, 1):
                lines.append(f'Slide {n}')
                for shape in slide.shapes:
                    if shape.has_text_frame: lines.append(shape.text)
                    if shape.has_table:
                        lines += [' | '.join(c.text for c in r.cells) for r in shape.table.rows]
        text = '\n'.join(lines)
        return {'path': str(p), 'content': text[:16000], 'truncated': len(text) > 16000,
                'note': 'Đọc text và bảng; không OCR ảnh, không phân tích toàn bộ bố cục.'}

    def prepare(self, name, args):
        p = self.files.path(args['path'], exists=False)
        if p.suffix.lower() not in {'.docx', '.pptx'} or not p.parent.is_dir():
            raise ValueError('Cần đường dẫn DOCX/PPTX trong whitelist và thư mục cha có sẵn.')
        if name == 'office_create' and p.exists(): raise ValueError('Tạo file cần tên chưa tồn tại.')
        if name == 'word_replace':
            if p.suffix.lower() != '.docx': raise ValueError('Sửa nội dung chỉ hỗ trợ DOCX.')
            self.check(str(p))
            if not args['search'] or len(args['search']) > 6000 or len(args['replacement']) > 10000:
                raise ValueError('Chuỗi tìm/thay không hợp lệ.')
        elif name in ('office_create','office_overwrite'):
            if name=='office_overwrite':
                if p.suffix.lower()!='.docx':raise ValueError('Ghi đè chỉ hỗ trợ DOCX.')
                if p.exists():self.check(str(p))
            if not args.get('content') or len(args['content']) > 30000 or len(args.get('title','')) > 500:
                raise ValueError('Nội dung 1–30000 ký tự, tiêu đề tối đa 500.')
            if p.suffix.lower()=='.docx' and 'font_name' in args:
                if not isinstance(args['font_name'],str) or not 1<=len(args['font_name'])<=80:raise ValueError('Tên font không hợp lệ.')
                if not 8<=float(args.get('font_size',13))<=72:raise ValueError('Cỡ chữ phải từ 8 đến 72.')
                if not 1<=float(args.get('line_spacing',1.15))<=3:raise ValueError('Giãn dòng phải từ 1 đến 3.')
                if args.get('alignment','justify') not in ('left','center','right','justify'):raise ValueError('Căn đoạn không hợp lệ.')
        else: raise ValueError('Thao tác Office không hợp lệ.')
        return {'action': name, **args, 'path': str(p), 'sha256': digest(p) if p.exists() else None,
                'approval_id': uuid.uuid4().hex,
                'note': 'Sửa Word: định dạng ký tự của đoạn chứa chuỗi sẽ đồng nhất. Có backup trước khi sửa.' if name == 'word_replace' else 'Tạo tài liệu/slide với bố cục cơ bản.'}

    def commit(self, plan):
        p = self.files.path(plan['path'], exists=plan['sha256'] is not None)
        if (digest(p) if p.exists() else None) != plan['sha256']:
            raise RuntimeError('File đã thay đổi sau preview; xin duyệt lại.')
        backup = None
        if p.exists():
            self.check(str(p))
            backup = self.files.backups / f'{p.stem}_{uuid.uuid4().hex}{p.suffix}'
            shutil.copy2(p, backup)
            if digest(backup) != plan['sha256']: raise RuntimeError('Backup không khớp.')
        self.files.audit('office_started', {'path': str(p), 'action': plan['action']})
        fd, raw = tempfile.mkstemp(prefix='.office_', suffix=p.suffix, dir=p.parent)
        os.close(fd); tmp = Path(raw)
        try:
            if plan['action'] == 'word_replace':
                from docx import Document
                doc = Document(p)
                paragraphs = list(doc.paragraphs)
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells: paragraphs.extend(cell.paragraphs)
                matches = [(para, para.text.count(plan['search'])) for para in paragraphs]
                if sum(n for _, n in matches) != 1:
                    raise ValueError('Chuỗi tìm phải xuất hiện đúng một lần trong một đoạn/bảng. Header/footer chưa hỗ trợ.')
                for para, count in matches:
                    if count:
                        # Text spanning runs is supported; formatting of this paragraph becomes uniform.
                        para.text = para.text.replace(plan['search'], plan['replacement'], 1)
                doc.save(tmp)
            elif p.suffix.lower() == '.docx':
                from docx import Document
                from .word_layout import resignation_paragraphs,format_resignation
                letter=resignation_paragraphs(plan.get('title',''),plan['content'])
                doc=Document()
                if letter:
                    for line in letter:doc.add_paragraph(line)
                else:
                    doc.add_heading(plan.get('title','Tài liệu'),0)
                    for line in plan['content'].splitlines():doc.add_paragraph(line)
                if 'font_name' in plan:
                    from docx.shared import Pt,Cm
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    from docx.oxml.ns import qn
                    align={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}
                    section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
                    section.top_margin=section.bottom_margin=Cm(2);section.left_margin=Cm(3);section.right_margin=Cm(2)
                    normal=doc.styles['Normal'];normal.font.name=plan['font_name'];normal.font.size=Pt(float(plan.get('font_size',13)))
                    for index,paragraph in enumerate(doc.paragraphs):
                        paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER if index==0 else align[plan.get('alignment','justify')]
                        paragraph.paragraph_format.line_spacing=float(plan.get('line_spacing',1.15))
                        paragraph.paragraph_format.space_after=Pt(6)
                        for run in paragraph.runs:
                            run.font.name=plan['font_name'];run.font.size=Pt(float(plan.get('font_size',13)))
                            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),plan['font_name'])
                            if index==0:run.bold=True
                if letter:format_resignation(doc,plan.get('font_name','Times New Roman'),float(plan.get('font_size',13)))
                doc.save(tmp)
            else:
                from pptx import Presentation
                deck = Presentation()
                for index, block in enumerate(plan['content'].split('\n---\n')[:30]):
                    slide = deck.slides.add_slide(deck.slide_layouts[1])
                    lines = block.strip().splitlines()
                    slide.shapes.title.text = lines[0] if lines else plan.get('title', 'Slide')
                    slide.placeholders[1].text = '\n'.join(lines[1:])
                deck.save(tmp)
            if (digest(p) if p.exists() else None) != plan['sha256']:
                raise RuntimeError('File thay đổi trước khi ghi.')
            if plan['sha256'] is None:
                with p.open('xb') as out, tmp.open('rb') as inp: shutil.copyfileobj(inp, out)
            else: os.replace(tmp, p)
            result = {'ok': True, 'path': str(p), 'backup': str(backup) if backup else None}
            self.files.audit('office_success', result); return result
        except Exception as exc:
            self.files.audit('office_failed', {'path': str(p), 'error': str(exc), 'backup': str(backup)})
            raise
        finally: tmp.unlink(missing_ok=True)
