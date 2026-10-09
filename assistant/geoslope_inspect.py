"""Read engineering inputs and existing GSZ results; never invoke a solver."""
import csv
import io
import math
import re
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
import json
import uuid
import hashlib
from collections import Counter


def label(value):
    text=str(value or '').lower().replace('đ','d')
    return ''.join(c for c in unicodedata.normalize('NFD',text) if not unicodedata.combining(c))


def parallel_layers(profile,layers):
    """Vertical thickness, not a perpendicular CAD offset or a guessed horizon."""
    if not isinstance(profile,list) or not 2<=len(profile)<=1000:raise ValueError('Đường tự nhiên cần 2–1000 điểm.')
    points=[]
    for p in profile:
        if len(p)!=2 or any(isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) for v in p):raise ValueError('Điểm mặt cắt phải là cặp số hữu hạn.')
        points.append(list(map(float,p)))
    if points[-1][0]<points[0][0]:points.reverse()
    if any(b[0]<=a[0] for a,b in zip(points,points[1:])):raise ValueError('Đường tự nhiên phải đơn trị theo X; không tự sắp xếp hoặc làm mất đoạn đứng.')
    if not isinstance(layers,list) or not 1<=len(layers)<=30:raise ValueError('Cần 1–30 lớp từ BTH.')
    depth=0;boundaries=[{'depth':0,'points':points}];regions=[]
    for layer in layers:
        if not isinstance(layer,dict):raise ValueError('Mỗi lớp cần code và thickness.')
        code=layer.get('code');thickness=layer.get('thickness')
        if not isinstance(code,str) or not 1<=len(code)<=40:raise ValueError('Tên lớp không hợp lệ.')
        if isinstance(thickness,bool) or not isinstance(thickness,(int,float)) or not math.isfinite(thickness) or not 0<thickness<=1000:raise ValueError('Bề dày phải là số dương hữu hạn; không đổi ô thiếu/lỗi thành 0.')
        depth+=thickness
        bottom=[[x,y-depth] for x,y in points]
        regions.append({'code':code,'thickness':thickness,'source':layer.get('source'),
                        'points':boundaries[-1]['points']+list(reversed(bottom))})
        boundaries.append({'depth':depth,'points':bottom})
    return {'boundaries':boundaries,'regions':regions,'total_thickness':depth,
            'note':'Địa tầng phác dựng từ bề dày đứng của BTH, không phải khảo sát ranh giới thực; chưa gán vật liệu hoặc tính ổn định.'}


def section_table(sheet, rows):
    """Keep section/layer provenance, including repeated layer codes and alternatives."""
    from .dxf_profile import parse_station
    headers=[c for row in rows[:20] for c in row
             if 'mat cat tinh toan' in label(c.value)
             or 'mæt c¾t tÝnh'.lower() in str(c.value or '').lower()]
    if not headers:return None
    h=max(headers,key=lambda c:c.row)
    layer_heads=[c for row in rows[:20] for c in row
                 if 'be day lop dat' in label(c.value) or 'bÒ dµy líp ®Êt'.lower() in str(c.value or '').lower()]
    if not layer_heads:return None
    lh=max(layer_heads,key=lambda c:c.row)
    by_row={i:row for i,row in enumerate(rows,1)}
    codes=[]
    for r in range(lh.row+1,min(lh.row+4,len(rows)+1)):
        for c in by_row.get(r,())[lh.column-1:]:
            if re.fullmatch(r'\d+[a-z]?',str(c.value or '').strip(),re.I):codes.append(c)
            elif codes:break
        if codes:break
    if not codes:return None
    records=[]
    for row in rows[max(h.row,codes[0].row):]:
        station_cell=row[h.column-1]
        station=parse_station(str(station_cell.value or ''))
        if station is None:continue
        layers=[]
        for code in codes:
            c=row[code.column-1]
            if c.value is None:continue
            layers.append({'code':str(code.value),'thickness':c.value,'cell':c.coordinate,
                           'header_cell':code.coordinate,'error':c.data_type=='e'})
        records.append({'section':station,'section_cell':station_cell.coordinate,'row':station_cell.row,
                        'layers':layers,'row_cells':[{'cell':c.coordinate,'value':c.value,'error':c.data_type=='e'} for c in row if c.value is not None]})
    return {'sheet':sheet.title,'section_header':h.coordinate,'layer_header':lh.coordinate,
            'header_cells':[{'cell':c.coordinate,'value':c.value} for row in rows[:max(h.row,codes[0].row)] for c in row if c.value is not None],
            'sections':records,'scanned_rows':len(rows),'truncated':sheet.max_row>len(rows)}


def inspect_workbook(path, sheet='', cell_range='A1:J20'):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    values=load_workbook(path,data_only=True,read_only=True)
    formulas=load_workbook(path,data_only=False,read_only=True)
    try:
        sheets=[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in values]
        if sheet:
            if sheet not in values.sheetnames:raise ValueError('Sheet không tồn tại; chọn tên từ sheets.')
            left,top,right,bottom=range_boundaries(cell_range)
            if None in (left,top,right,bottom) or min(left,top)<1 or right<left or bottom<top or (right-left+1)*(bottom-top+1)>5000:
                raise ValueError('cell_range cần vùng hữu hạn tối đa 5000 ô.')
            cached={(c.row,c.column):c for row in values[sheet].iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right) for c in row if c.value is not None}
            cells=[]
            for r,row in enumerate(formulas[sheet].iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right),top):
                for col,c in enumerate(row,left):
                    v=cached.get((r,col))
                    if c.value is not None or v is not None:
                        cells.append({'cell':c.coordinate,'value':v.value if v else None,
                                      'formula':c.value if c.data_type=='f' else None,
                                      'error':bool((v and v.data_type=='e') or c.data_type=='e')})
            return {'sheets':sheets,'sheet':sheet,'range':cell_range,'cells':cells,
                    'note':'Giá trị công thức là cache đã lưu, không tự tính lại Excel. Cache thiếu/lỗi không được đổi thành 0.'}
        tables=[];sections=[]
        for s in values:
            # Locate indicators by labels, not a fixed sheet or fixed row number.
            rows=list(s.iter_rows(min_row=1,max_row=min(s.max_row,120),max_col=min(s.max_column,250)))
            sections_table=section_table(s,rows)
            if sections_table:sections.append(sections_table)
            indicators={}
            for row in rows:
                for c in row:
                    name=label(c.value)
                    if 'dung trong' in name:indicators['unit_weight']=c.row
                    elif 'luc dinh' in name:indicators['cohesion']=c.row
                    elif 'goc ma sat' in name:indicators['phi']=c.row
            if set(indicators)!= {'unit_weight','cohesion','phi'}:continue
            header_row=min(indicators.values())-1
            if header_row<1:continue
            row_by_number={i:row for i,row in enumerate(rows,1)}
            header=row_by_number[header_row]
            unit_columns=[c.column for c in header if label(c.value).strip()=='don vi']
            unit_column=unit_columns[0] if len(unit_columns)==1 else None
            materials=[]
            for h in header[unit_column:] if unit_column else header:
                if h.value is None:continue
                if label(h.value).strip() in {'ten chi tieu','ky hieu','don vi'}:continue
                fields={}
                for key,r in indicators.items():
                    c=row_by_number[r][h.column-1]
                    fields[key]={'value':c.value,'cell':c.coordinate,
                                 'declared_unit':row_by_number[r][unit_column-1].value if unit_column else None,
                                 'row_context':[str(x.value) for x in row_by_number[r][:unit_column-1] if x.value is not None] if unit_column else []}
                materials.append({'name':str(h.value),'header_cell':h.coordinate,'fields':fields})
            tables.append({'sheet':s.title,'materials':materials})
        return {'sheets':sheets,'material_tables':tables,'section_tables':sections,
                'note':'Bảng nhận diện theo nhãn chỉ tiêu. Không tự chọn c tổng/c hiệu quả, quy đổi đơn vị hay ghép tên lớp; đọc vùng ô để kiểm tra.'}
    finally:
        values.close();formulas.close()


def inspect_dxf(path, layer='', handle='', start=0, limit=30):
    import ezdxf
    from ezdxf.lldxf.encoding import decode_dxf_unicode
    doc=ezdxf.readfile(path);model=doc.modelspace()
    counts={}
    for e in model:
        counts.setdefault(e.dxf.layer,Counter())[e.dxftype()]+=1
    candidates=[e for e in model if e.dxftype() in {'LWPOLYLINE','POLYLINE','LINE','TEXT','MTEXT','INSERT'}
                and (not layer or e.dxf.layer==layer) and (not handle or e.dxf.handle.upper()==handle.upper())]
    entities=[]
    for e in candidates[start:start+limit]:
        item={'handle':e.dxf.handle,'layer':e.dxf.layer,'type':e.dxftype()}
        if e.dxftype()=='LWPOLYLINE':
            item.update(points=[list(map(float,p)) for p in e.get_points('xyb')],closed=bool(e.closed),elevation=float(e.dxf.elevation))
        elif e.dxftype()=='POLYLINE':
            item.update(points=[list(map(float,v.dxf.location)) for v in e.vertices],closed=bool(e.is_closed))
        elif e.dxftype()=='LINE':item['points']=[list(map(float,e.dxf.start)),list(map(float,e.dxf.end))]
        elif e.dxftype() in {'TEXT','MTEXT'}:
            item['text']=decode_dxf_unicode(e.dxf.text if e.dxftype()=='TEXT' else e.plain_text())[:2000]
            item['insert']=list(map(float,e.dxf.insert))
        else:
            item.update(block=e.dxf.name,insert=list(map(float,e.dxf.insert)),
                        scale=[e.dxf.xscale,e.dxf.yscale,e.dxf.zscale],rotation=e.dxf.rotation)
        if len(item.get('points',[]))>1000:
            item['point_count']=len(item['points']);item['points']=item['points'][:1000];item['points_truncated']=True
        entities.append(item)
    return {'insunits':doc.header.get('$INSUNITS',0),'measurement':doc.header.get('$MEASUREMENT'),
            'layers':[{'name':name,'entities':dict(count)} for name,count in counts.items()],
            'entities':entities,'total_matching':len(candidates),
            'next_start':start+limit if start+limit<len(candidates) else None,
            'note':'INSUNITS chỉ là khai báo của DXF, phải đối chiếu kích thước/cao độ trước khi đổi tỷ lệ. INSERT chưa bung block; đường cong bulge không được coi là đoạn thẳng. Không ghép các mặt cắt khác nhau.'}


def inspect_gsz(path):
    with zipfile.ZipFile(path) as archive:
        entries=archive.infolist()
        if len(entries)>10000 or sum(e.file_size for e in entries)>200*1024**2:
            raise ValueError('GSZ vượt giới hạn đọc 200 MiB sau giải nén hoặc 10000 mục.')
        roots=[e.filename for e in entries if '/' not in e.filename and e.filename.lower().endswith('.xml')]
        if len(roots)!=1:raise ValueError('GSZ cần đúng một XML dự án gốc; không chọn XML kết quả thay cho mô hình.')
        data=archive.read(roots[0])
        if len(data)>20*1024**2 or re.search(br'<!\s*(?:DOCTYPE|ENTITY)',data,re.I):raise ValueError('XML dự án không được hỗ trợ.')
        root=ET.fromstring(data)
        if root.tag!='GSIData':raise ValueError('XML không phải GSIData của GeoStudio.')
        analyses=[{c.tag:c.text for c in a if c.tag in {'ID','Name','Kind','Method','GeometryId'}} for a in root.findall('./Analyses/Analysis')]
        assignments=[]
        for context in root.findall('./Contexts/Context'):
            assignments.extend({'analysis_id':context.findtext('AnalysisID'),**e.attrib} for e in context.findall('./GeometryUsesMaterials/GeometryUsesMaterial'))
        active={e.get('Entry') for e in assignments}
        materials=[]
        for m in root.findall('./Materials/Material'):
            materials.append({'id':m.findtext('ID'),'name':m.findtext('Name'),'model':m.findtext('SlopeModel'),
                              'used':m.findtext('ID') in active,
                              'parameters':{c.tag:c.text for c in m.findall('./StressStrain/*') if c.text and c.text.strip()}})
        geometries=[]
        for g in root.findall('./Geometries/Geometry'):
            geometries.append({'name':g.findtext('Name'),'points':[p.attrib for p in g.findall('./Points/Point')],
                               'regions':[{'id':r.findtext('ID'),'point_ids':r.findtext('PointIDs')} for r in g.findall('./Regions/Region')]})
        results=[]
        for e in entries:
            if not e.filename.lower().endswith('/slip_surface.csv'):continue
            count=valid=0;best=None
            with archive.open(e) as raw:
                for row in csv.DictReader(io.TextIOWrapper(raw,encoding='utf-8-sig')):
                    count+=1
                    try:fs=float(row.get('SlipFOS',''))
                    except (ValueError,TypeError):continue
                    if not math.isfinite(fs) or fs<=0:continue
                    valid+=1
                    if best is None or fs<best['factor_of_safety']:
                        best={'factor_of_safety':fs,'slip_number':row.get('SlipNum'),
                              'center_x':row.get('SlipCenterX'),'center_y':row.get('SlipCenterY'),'radius':row.get('SlipRadiusX')}
            results.append({'entry':e.filename,'rows':count,'positive_finite_rows':valid,'minimum':best})
        def xml(tag):
            elem=root.find(tag)
            return ET.tostring(elem,encoding='unicode') if elem is not None else None
        return {'version':root.attrib,'project_xml':roots[0],'analyses':analyses,'materials':materials,
                'assignments':assignments,'geometries':geometries,'coordinates_xml':xml('Coordinates'),
                'stability_xml':xml('StabilityItems'),'water_xml':xml('WaterItems'),'stored_results':results,
                'new_solve_performed':False,
                'note':'Đây là mô hình/kết quả đã lưu trong file, không phải vừa chạy Solve. Fs nhỏ nhất đọc toàn bộ CSV; chưa chứng minh hội tụ, miền tìm kiếm đầy đủ hoặc đạt tiêu chuẩn thiết kế.'}


class GeoslopeInspect:
    def __init__(self,files):self.files=files
    def prepare(self,name,args):
        path=self.files.path(args['path'])
        if not path.is_file() or path.suffix.lower() not in {'.xlsx','.dxf','.gsz'}:raise ValueError('Cần XLSX, DXF hoặc GSZ trong thư mục được phép.')
        if path.stat().st_size>40*1024**2:raise ValueError('Tệp tối đa 40 MiB.')
        if name=='geoslope_profile':
            if path.suffix.lower()!='.dxf':raise ValueError('geoslope_profile cần DXF.')
            if args.get('units')!='m':raise ValueError('Cần xác minh đơn vị m trước khi dựng địa tầng.')
            try:layers=json.loads(args['layers'])
            except (ValueError,KeyError,TypeError):raise ValueError('layers cần chuỗi JSON các lớp từ BTH.') from None
            if not isinstance(args.get('handle'),str) or not args['handle']:raise ValueError('Chọn handle đường tự nhiên từ geoslope_inspect trước.')
            result=inspect_dxf(path,handle=args['handle'],limit=1)
            if not result['entities']:raise ValueError('Không tìm thấy handle đường tự nhiên.')
            entity=result['entities'][0]
            if entity['type']!='LWPOLYLINE' or entity['closed'] or entity.get('points_truncated') or any(p[2]!=0 for p in entity['points']):raise ValueError('Cần LWPOLYLINE hở, không bulge và đầy đủ điểm; không tự làm thẳng đường cong.')
            geometry=parallel_layers([p[:2] for p in entity['points']],layers)
            output=self.files.path(str(self.files.roots[0]/('GeoSlope-BTH-profile-'+uuid.uuid4().hex[:8]+'.dxf')),exists=False)
            return {'action':name,'source':str(path),'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                    'source_handle':args['handle'],'path':str(output),'geometry':geometry,'units':'m'}
        start=args.get('start',0);limit=args.get('limit',30)
        if type(start) is not int or not 0<=start<=1000000 or type(limit) is not int or not 1<=limit<=100:raise ValueError('start/limit không hợp lệ.')
        return {'action':'geoslope_inspect','path':str(path),'options':{k:v for k,v in args.items() if k!='path'}}
    def commit(self,plan):
        if plan['action']=='geoslope_profile':
            import ezdxf
            source=self.files.path(plan['source'])
            if source.stat().st_size>40*1024**2 or hashlib.sha256(source.read_bytes()).hexdigest()!=plan['source_sha256']:raise RuntimeError('DXF nguồn đã thay đổi; đọc lại trước khi vẽ.')
            output=self.files.path(plan['path'],exists=False)
            doc=ezdxf.new('R2010');doc.header['$INSUNITS']=6
            doc.styles.new('GEOSLOPE_TEXT',dxfattribs={'font':'arial.ttf'})
            model=doc.modelspace();geometry=plan['geometry']
            doc.layers.new('NATURAL',dxfattribs={'color':3,'lineweight':35})
            doc.layers.new('NOTE',dxfattribs={'color':7})
            model.add_lwpolyline(geometry['boundaries'][0]['points'],dxfattribs={'layer':'NATURAL'})
            palette=[30,140,200,60,170]
            for i,(region,boundary) in enumerate(zip(geometry['regions'],geometry['boundaries'][1:]),1):
                name='SOIL_'+str(i);doc.layers.new(name,dxfattribs={'color':palette[(i-1)%len(palette)],'lineweight':25})
                model.add_lwpolyline(region['points'],close=True,dxfattribs={'layer':name})
                midpoint=region['points'][len(boundary['points'])//2]
                model.add_text('Lop '+region['code']+' | h='+str(region['thickness'])+' m',
                               dxfattribs={'layer':'NOTE','style':'GEOSLOPE_TEXT','height':.6,'insert':(midpoint[0],midpoint[1]-region['thickness']/2)})
            top=geometry['boundaries'][0]['points'];left=min(x for x,y in top);high=max(y for x,y in top)
            model.add_text('DIA TANG THEO BTH - BE DAY DUNG - DON VI m',dxfattribs={'layer':'NOTE','style':'GEOSLOPE_TEXT','height':.8,'insert':(left,high+2)})
            model.add_text('CHUA GAN VAT LIEU / CHUA CHAY SOLVE',dxfattribs={'layer':'NOTE','style':'GEOSLOPE_TEXT','height':.6,'insert':(left,high+1)})
            text=io.StringIO();doc.write(text)
            with output.open('xb') as f:f.write(text.getvalue().encode(doc.output_encoding,errors='dxfreplace'))
            return {'ok':True,'path':str(output),'files':[str(output)],'layer_count':len(geometry['regions']),
                    'total_thickness':geometry['total_thickness'],'source_handle':plan['source_handle'],
                    'new_solve_performed':False,'note':geometry['note']}
        path=self.files.path(plan['path']);opts=plan['options']
        if path.stat().st_size>40*1024**2:raise ValueError('Tệp đã thay đổi và vượt giới hạn đọc.')
        if path.suffix.lower()=='.xlsx':result=inspect_workbook(path,opts.get('sheet',''),opts.get('cell_range','A1:J20'))
        elif path.suffix.lower()=='.dxf':result=inspect_dxf(path,opts.get('layer',''),opts.get('handle',''),opts.get('start',0),opts.get('limit',30))
        else:result=inspect_gsz(path)
        return {'ok':True,'path':str(path),'read_only':True,**result}
