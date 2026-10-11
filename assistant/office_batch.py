"""Structured Office changes, saved once and read back before reporting success."""
import json
import os
import shutil
import tempfile
import uuid
from pathlib import Path
from .excel import digest,scalar
from .tools import schema,TEXT
from .windows_apps import _STOP,wait_automation

SCHEMAS=[
    schema('excel_create','Tạo XLSX có dữ liệu, công thức, định dạng, lọc và biểu đồ rồi mở nếu open=true. rows là JSON ma trận; formulas là JSON {"D2":"=B2*C2"}; options JSON {"bold_header":true,"auto_filter":true,"chart":{"range":"A1:B4","title":"..."}}. Công thức chỉ tính lại khi mở Excel/COM, không báo cache là kết quả mới.',
           {'path':TEXT,'sheet':TEXT,'rows':TEXT,'formulas':TEXT,'options':TEXT,'open':{'type':'boolean'}},['path','sheet','rows']),
    schema('word_inspect','Đọc DOCX gồm chỉ số đoạn, bảng, ảnh, công thức và header/footer để sửa đúng vị trí.',{'path':TEXT},['path']),
    schema('word_edit','Tạo/sửa DOCX theo lô; backup và đọc lại. operations JSON gồm append{text}, replace{search,text}, paragraph{index,text,font,size,bold,alignment}, table{rows}, table_cell{table,row,column,text}, image{path,width_cm}, equation{latex}, header/footer{section,text}. Chỉ số từ 0. create=true tạo mới. Công thức là Office Math có thể chỉnh sửa.',
           {'path':TEXT,'operations':TEXT,'create':{'type':'boolean'}},['path','operations'])]


def operations(raw):
    rows=json.loads(raw) if isinstance(raw,str) else raw
    if not isinstance(rows,list) or not 1<=len(rows)<=100:raise ValueError('Cần 1–100 thao tác.')
    if any(not isinstance(row,dict) for row in rows):raise ValueError('Mỗi thao tác phải là object.')
    return rows


def word_summary(doc):
    return {'paragraphs':[{'index':i,'text':p.text,'style':p.style.name} for i,p in enumerate(doc.paragraphs)][:200],
            'tables':[{'index':i,'rows':[[c.text for c in r.cells] for r in t.rows][:100]} for i,t in enumerate(doc.tables)][:30],
            'images':len(doc.inline_shapes),'equations':len(doc.element.xpath('//*[local-name()="oMath"]')),
            'sections':[{'index':i,'header':s.header.paragraphs[0].text,'footer':s.footer.paragraphs[0].text} for i,s in enumerate(doc.sections)]}


class OfficeBatch:
    def __init__(self,files):self.files=files;self.on_status=None

    def prepare(self,name,args):
        create=name=='excel_create' or args.get('create') is True
        p=self.files.path(args['path'],exists=not create)
        if p.suffix.lower()!=('.xlsx' if name=='excel_create' else '.docx'):raise ValueError('Định dạng Office không hợp lệ.')
        if create and p.exists():raise FileExistsError('File đích đã tồn tại; dùng tên mới.')
        if not p.parent.is_dir():raise ValueError('Thư mục đích chưa có.')
        plan={'action':name,'path':str(p),'arguments':dict(args),'sha256':digest(p) if p.exists() else None}
        if name=='excel_create':
            rows=json.loads(args['rows'])
            if not isinstance(rows,list) or not rows or any(not isinstance(r,list) for r in rows) or sum(len(r) for r in rows)>10000:raise ValueError('Dữ liệu tối đa 10000 ô.')
            if any(type(v) not in (str,int,float,bool,type(None)) for r in rows for v in r):raise ValueError('Ô cần giá trị đơn.')
            sheet=args['sheet']
            if not isinstance(sheet,str) or not 1<=len(sheet)<=31 or any(c in sheet for c in '[]:*?/\\'):raise ValueError('Tên sheet không hợp lệ.')
            plan['rows']=rows;plan['formulas']=json.loads(args.get('formulas','{}'));plan['options']=json.loads(args.get('options','{}'))
            if not isinstance(plan['formulas'],dict) or not isinstance(plan['options'],dict):raise ValueError('formulas/options cần JSON object.')
            from openpyxl.utils.cell import coordinate_to_tuple
            for cell,value in plan['formulas'].items():
                row,col=coordinate_to_tuple(cell)
                if row>100000 or col>16384 or not isinstance(value,str) or not value.startswith('='):raise ValueError('Công thức cần ô hợp lệ và bắt đầu =.')
            plan['formulas']={k.upper():v for k,v in plan['formulas'].items()}
        elif name=='word_edit':
            plan['operations']=operations(args['operations'])
            allowed={'append','replace','paragraph','table','table_cell','image','equation','header','footer'}
            for op in plan['operations']:
                if op.get('action') not in allowed:raise ValueError('Thao tác Word không được hỗ trợ.')
                if op['action']=='image':
                    source=self.files.path(op['path']);op['path']=str(source);op['sha256']=digest(source)
        return plan

    def commit(self,plan):
        p=self.files.path(plan['path'],exists=plan['sha256'] is not None)
        if (digest(p) if p.exists() else None)!=plan['sha256']:raise RuntimeError('File đã đổi sau preview.')
        if _STOP.is_set():raise PermissionError('Đã dừng nhiệm vụ.')
        name=plan['action'];backup=None
        if name=='word_inspect':
            from docx import Document
            return {'ok':True,'path':str(p),**word_summary(Document(p))}
        if p.exists():
            backup=self.files.backups/(uuid.uuid4().hex+p.suffix);shutil.copy2(p,backup)
        fd,raw=tempfile.mkstemp(suffix=p.suffix,dir=p.parent);os.close(fd);temp=Path(raw)
        try:
            self.files.audit('office_batch_started',{'action':name,'path':str(p),'backup':str(backup) if backup else None})
            if name=='excel_create':
                from openpyxl import Workbook,load_workbook
                from openpyxl.styles import Font
                from openpyxl.chart import BarChart,Reference
                from openpyxl.utils.cell import range_boundaries
                doc=Workbook();sheet=doc.active;sheet.title=plan['arguments']['sheet']
                for row in plan['rows']:
                    for value in row:
                        if isinstance(value,float) and not __import__('math').isfinite(value):raise ValueError('Số không hữu hạn.')
                    sheet.append(row)
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value,str):cell.data_type='s'
                for cell,formula in plan['formulas'].items():sheet[cell]=formula
                options=plan['options']
                if options.get('bold_header'):
                    for cell in sheet[1]:cell.font=Font(bold=True)
                if options.get('auto_filter'):sheet.auto_filter.ref=sheet.dimensions
                if options.get('chart'):
                    config=options['chart'];left,top,right,bottom=range_boundaries(config['range'])
                    if left==right:raise ValueError('Biểu đồ cần cột nhãn và dữ liệu.')
                    chart=BarChart();chart.title=config.get('title','Chart')
                    chart.add_data(Reference(sheet,min_col=left+1,max_col=right,min_row=top,max_row=bottom),titles_from_data=True)
                    chart.set_categories(Reference(sheet,min_col=left,min_row=top+1,max_row=bottom));sheet.add_chart(chart,config.get('anchor','F2'))
                doc.save(temp);doc.close()
                check=load_workbook(temp,data_only=False);ws=check[plan['arguments']['sheet']]
                expected={(i+1,j+1):v for i,r in enumerate(plan['rows']) for j,v in enumerate(r)}
                for (i,j),v in expected.items():
                    if ws.cell(i,j).coordinate not in plan['formulas'] and ws.cell(i,j).value!=(None if v=='' else v):raise RuntimeError('Dữ liệu Excel đọc lại không khớp.')
                for cell,v in plan['formulas'].items():
                    if ws[cell].value!=v or ws[cell].data_type!='f':raise RuntimeError('Công thức lưu không khớp.')
                evidence={'rows':ws.max_row,'columns':ws.max_column,'formulas':plan['formulas'],'charts':len(ws._charts),'formula_values_verified':False};check.close()
            else:
                from docx import Document
                from docx.shared import Pt,Cm
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                doc=Document(p) if p.exists() else Document();checks=[]
                for index,op in enumerate(plan['operations'],1):
                    wait_automation()
                    if _STOP.is_set():raise PermissionError('Đã dừng nhiệm vụ.')
                    kind=op['action']
                    if self.on_status:self.on_status('Word · bước '+str(index)+'/'+str(len(plan['operations']))+' · '+{'append':'thêm nội dung','replace':'thay nội dung','paragraph':'định dạng đoạn','table':'tạo bảng','table_cell':'ghi ô bảng','image':'chèn ảnh','equation':'chèn công thức','header':'ghi header','footer':'ghi footer'}[kind])
                    if kind=='append':doc.add_paragraph(op['text']);checks.append(('text',op['text']))
                    elif kind=='replace':
                        paras=list(doc.paragraphs)+[c.paragraphs[0] for t in doc.tables for r in t.rows for c in r.cells]
                        matches=[q for q in paras if op['search'] in q.text]
                        if sum(q.text.count(op['search']) for q in matches)!=1:raise ValueError('Chuỗi thay cần xuất hiện đúng một lần.')
                        matches[0].text=matches[0].text.replace(op['search'],op['text']);checks.append(('text',op['text']))
                    elif kind=='paragraph':
                        para=doc.paragraphs[int(op['index'])]
                        if 'text' in op:para.text=op['text'];checks.append(('text',op['text']))
                        if 'alignment' in op:para.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}[op['alignment']]
                        for run in para.runs:
                            if 'font' in op:run.font.name=op['font']
                            if 'size' in op:run.font.size=Pt(float(op['size']))
                            if 'bold' in op:run.bold=bool(op['bold'])
                    elif kind=='table':
                        rows=op['rows']
                        if not rows or sum(len(r) for r in rows)>5000:raise ValueError('Bảng tối đa 5000 ô.')
                        table=doc.add_table(rows=len(rows),cols=max(map(len,rows)))
                        for i,row in enumerate(rows):
                            for j,value in enumerate(row):table.cell(i,j).text=str(value)
                        checks.append(('tables',len(doc.tables)))
                    elif kind=='table_cell':doc.tables[int(op['table'])].cell(int(op['row']),int(op['column'])).text=op['text'];checks.append(('text',op['text']))
                    elif kind=='image':
                        image=self.files.path(op['path'])
                        if digest(image)!=op['sha256']:raise RuntimeError('Ảnh đã thay đổi sau preview.')
                        doc.add_picture(str(image),width=Cm(float(op.get('width_cm',10))));checks.append(('images',len(doc.inline_shapes)))
                    elif kind=='equation':
                        from .math_format import word_math
                        before=len(doc.element.xpath('//*[local-name()="oMath"]'));word_math(doc.add_paragraph(),'$'+op['latex']+'$')
                        if len(doc.element.xpath('//*[local-name()="oMath"]'))<=before:raise RuntimeError('Chưa chuyển được công thức sang Office Math.')
                        checks.append(('equations',before+1))
                    else:
                        section=doc.sections[int(op.get('section',0))];part=section.header if kind=='header' else section.footer
                        part.is_linked_to_previous=False;part.paragraphs[0].text=op['text'];checks.append(('text',op['text']))
                doc.save(temp);read=Document(temp);evidence=word_summary(read)
                all_text='\n'.join([p.text for p in read.paragraphs]+[c.text for t in read.tables for r in t.rows for c in r.cells]+[v for s in evidence['sections'] for v in (s['header'],s['footer'])])
                for kind,value in checks:
                    if kind=='text' and value not in all_text:raise RuntimeError('Nội dung Word đọc lại không khớp.')
                    if kind!='text' and (len(evidence['tables']) if kind=='tables' else evidence[kind])<value:raise RuntimeError('Cấu trúc Word đọc lại không khớp.')
            if (digest(p) if p.exists() else None)!=plan['sha256']:raise RuntimeError('File đổi trước khi lưu.')
            if plan['sha256'] is None:
                with p.open('xb') as out,temp.open('rb') as src:shutil.copyfileobj(src,out)
            else:os.replace(temp,p)
            result={'ok':True,'path':str(p),'backup':str(backup) if backup else None,'verified':True,'readback':evidence}
            if plan['arguments'].get('open'):
                if os.name!='nt':raise RuntimeError('Mở ứng dụng cần Windows; file đã được tạo.')
                os.startfile(str(p));result['opened']=True
            self.files.audit('office_batch_result',result);return result
        finally:temp.unlink(missing_ok=True)
