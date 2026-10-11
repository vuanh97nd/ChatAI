"""Native Office file tools exposed through online automation's preview/commit flow."""
from .excel import digest


class OnlineOffice:
    def __init__(self, excel, office, files,on_status=None):
        self.excel,self.office,self.files=excel,office,files
        from .task_tools import TaskTools,SCHEMAS as task_schemas
        from .office_batch import OfficeBatch,SCHEMAS as office_schemas
        from .com_tasks import ComTools,SCHEMAS as com_schemas
        self.task=TaskTools(files);self.batch=OfficeBatch(files);self.com=ComTools(files)
        self.task.on_status=self.batch.on_status=self.com.on_status=on_status
        self.schemas=task_schemas+office_schemas+com_schemas
        self.task_names={s['function']['name'] for s in task_schemas}
        self.batch_names={s['function']['name'] for s in office_schemas}
        self.com_names={s['function']['name'] for s in com_schemas}

    def prepare(self,name,args):
        if name in self.task_names:return self.task.prepare(name,args)
        if name in self.batch_names:return self.batch.prepare(name,args)
        if name in self.com_names:return self.com.prepare(name,args)
        if name=='excel_edit_batch':return self.excel.prepare_batch(**args)
        if name=='excel_edit_cell':return {'action':name,**self.excel.prepare_edit(**args)}
        if name=='excel_create_from_template':return self.excel.prepare_from_template(**args)
        if name in ('office_create','word_replace'):
            if name=='office_create' and args['path'].lower().endswith('.pptx') and len(args['content'].split('\n---\n'))>30:raise ValueError('Tạo tối đa 30 slide mỗi file; chia bài thành nhiều phần.')
            return self.office.prepare(name,args)
        if name=='file_list':
            self.files.path(args.get('path','.'))
            return {'action':name,'arguments':dict(args)}
        if name=='excel_list_files':return {'action':name,'arguments':dict(args)}
        if name in ('excel_list_sheets','excel_read','excel_summary','office_read'):
            path=self.office.check(args['path']) if name=='office_read' else self.excel.path(args['path'])
            return {'action':name,'arguments':dict(args,path=str(path)),'sha256':digest(path)}
        raise ValueError('Công cụ Office không hợp lệ.')

    def commit(self,plan):
        name=plan['action']
        if name in self.task_names:return self.task.commit(plan)
        if name in self.batch_names:return self.batch.commit(plan)
        if name in self.com_names:return self.com.commit(plan)
        if name in ('excel_edit_cell','excel_edit_batch'):
            result=self.excel.commit_edit(plan)
            from .excel import load_workbook
            from .math_format import excel_math
            workbook=load_workbook(self.excel.path(plan['path']),data_only=False)
            try:
                sheet=workbook[plan['sheet']];rows=plan.get('fills',[{'cell':plan.get('cell'),'value':plan.get('value')}])
                for item in rows:
                    if sheet[item['cell']].value!=excel_math(item['value']):raise RuntimeError('Ô Excel đọc lại không khớp yêu cầu.')
                result.update(verified=True,readback=[{'cell':i['cell'],'value':sheet[i['cell']].value} for i in rows[:50]])
            finally:workbook.close()
            return result
        if name=='excel_create_from_template':return self.excel.commit_from_template(plan)
        if name in ('office_create','word_replace'):
            result=self.office.commit(plan)
            p=self.files.path(plan['path'])
            if p.suffix.lower()=='.pptx':
                from pptx import Presentation
                deck=Presentation(p);blocks=plan['content'].split('\n---\n')
                if len(blocks)>30 or len(deck.slides)!=len(blocks):raise RuntimeError('Số slide đọc lại không khớp; chưa hoàn tất toàn bộ yêu cầu.')
                for slide,block in zip(deck.slides,blocks):
                    text='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
                    if any(line.strip() and line.strip() not in text for line in block.strip().splitlines()):raise RuntimeError('Nội dung slide đọc lại không khớp.')
                result['readback']={'slides':len(deck.slides)}
            else:
                from docx import Document
                doc=Document(p);text='\n'.join([q.text for q in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
                if name=='word_replace' and plan['replacement'] not in text:raise RuntimeError('Nội dung thay thế Word đọc lại không khớp.')
                result['readback']={'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)}
            return dict(result,verified=True)
        args=plan['arguments']
        if 'sha256' in plan:
            path=self.office.check(args['path']) if name=='office_read' else self.excel.path(args['path'])
            if digest(path)!=plan['sha256']:raise RuntimeError('File đã thay đổi; đọc lại trước khi tiếp tục.')
        target=self.files if name=='file_list' else self.office if name=='office_read' else self.excel
        return {'ok':True,**getattr(target,name)(**args)}

    def cache_key(self,plan):
        import hashlib,json
        if plan['action'] not in ('excel_read','excel_list_sheets','excel_summary','office_read','word_inspect','file_read','document_read'):return None
        if not plan.get('sha256'):return None
        args=plan.get('arguments',{});p=self.files.path(plan.get('path') or args['path'])
        if digest(p)!=plan['sha256']:raise RuntimeError('Dữ liệu đã thay đổi; đọc lại.')
        return hashlib.sha256(json.dumps({'action':plan['action'],'args':args,'sha256':plan['sha256']},sort_keys=True,default=str).encode()).hexdigest()
