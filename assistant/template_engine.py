import hashlib
import json
import os
import re
import shutil
import tempfile
import uuid
from pathlib import Path

MAX_DOCX_FILE = 20 * 1024 * 1024
MAX_PLACEHOLDERS = 200
MAX_VALUE_LEN = 10000


def load_docx(path):
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError('Chưa có thư viện python-docx. Cài đặt trước.') from error
    return Document(str(path))


def digest(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _extract_placeholders_from_paragraphs(paragraphs, found):
    for para in paragraphs:
        full_text = para.text
        for match in re.finditer(r'\{\{(\w+)\}\}', full_text):
            found.add(match.group(1))


def _extract_placeholders_from_tables(tables, found):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                _extract_placeholders_from_paragraphs(cell.paragraphs, found)
                # nested tables
                _extract_placeholders_from_tables(cell.tables, found)


def _fill_paragraph(para, values):
    """Replace {{key}} in a paragraph, handling run-split placeholders."""
    full_text = para.text
    if '{{' not in full_text:
        return
    # Check if any placeholder is present
    placeholders_in_text = re.findall(r'\{\{(\w+)\}\}', full_text)
    if not placeholders_in_text:
        return

    # Merge all runs into first run, then clear the rest
    if not para.runs:
        return

    # Build new full text with replacements
    new_text = full_text
    for key in placeholders_in_text:
        if key in values:
            new_text = new_text.replace('{{' + key + '}}', str(values[key]))

    # Keep formatting of first run, put all text there
    first_run = para.runs[0]
    first_run.text = new_text
    for run in para.runs[1:]:
        run.text = ''


def _fill_tables(tables, values):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _fill_paragraph(para, values)
                _fill_tables(cell.tables, values)


class TemplateEngine:
    def __init__(self, roots, backups, audit):
        self.roots = [Path(p).resolve() for p in roots]
        self.backups = Path(backups).resolve()
        self.backups.mkdir(parents=True, exist_ok=True)
        self.audit = audit

    def path(self, raw):
        if not isinstance(raw, str) or not raw or len(raw) > 1000:
            raise ValueError("Đường dẫn không hợp lệ.")
        p = Path(raw)
        if not p.is_absolute():
            p = self.roots[0] / p
        p = p.resolve(strict=True)
        if not any(p.is_relative_to(root) for root in self.roots):
            raise PermissionError("File nằm ngoài whitelist.")
        if not p.is_file() or p.suffix.lower() != ".docx":
            raise ValueError("Chỉ hỗ trợ file .docx.")
        if p.stat().st_size > MAX_DOCX_FILE:
            raise ValueError("File vượt giới hạn 20 MiB.")
        return p

    def scan_template(self, path: str) -> dict:
        """Read .docx template, extract all {{placeholder}} names."""
        p = self.path(path)
        doc = load_docx(p)
        found = set()
        _extract_placeholders_from_paragraphs(doc.paragraphs, found)
        _extract_placeholders_from_tables(doc.tables, found)
        return {
            'path': str(p),
            'placeholders': sorted(found),
            'table_count': len(doc.tables),
            'paragraph_count': len(doc.paragraphs),
        }

    def prepare(self, name, args) -> dict:
        if name != 'template_fill':
            raise ValueError(f"Thao tác không hỗ trợ: {name}")
        template = args['template']
        output_name = args['output_name']
        values_str = args['values']

        tmpl = self.path(template)

        if not isinstance(output_name, str) or not output_name:
            raise ValueError("output_name không hợp lệ.")
        clean_name = Path(output_name).name
        if not clean_name or clean_name != output_name or clean_name.startswith('.'):
            raise ValueError("output_name phải là tên file đơn, không chứa thư mục.")
        if '/' in clean_name or '\\' in clean_name or os.sep in clean_name:
            raise ValueError("output_name không được chứa dấu phân cách thư mục.")
        if not clean_name.lower().endswith('.docx'):
            clean_name = clean_name + '.docx'
        output_path = tmpl.parent / clean_name
        if output_path.exists():
            raise FileExistsError(f"File đích đã tồn tại: {output_path}. Chọn tên khác.")

        if not isinstance(values_str, str) or len(values_str) > 2_000_000:
            raise ValueError("values phải là chuỗi JSON tối đa 2000000 ký tự.")
        try:
            values = json.loads(values_str)
        except Exception:
            raise ValueError("values không phải JSON hợp lệ.")
        if not isinstance(values, dict):
            raise ValueError("values phải là đối tượng JSON {key: value}.")
        if len(values) > MAX_PLACEHOLDERS:
            raise ValueError(f"Tối đa {MAX_PLACEHOLDERS} placeholder.")
        validated = {}
        for k, v in values.items():
            if not isinstance(k, str) or not re.fullmatch(r'\w+', k):
                raise ValueError(f"Tên placeholder không hợp lệ: {k!r}")
            if not isinstance(v, str):
                raise ValueError(f"Giá trị của '{k}' phải là chuỗi.")
            if len(v) > MAX_VALUE_LEN:
                raise ValueError(f"Giá trị của '{k}' vượt {MAX_VALUE_LEN} ký tự.")
            validated[k] = v

        return {
            'action': 'template_fill',
            'template': str(tmpl),
            'output': str(output_path),
            'values': validated,
            'sha256': digest(tmpl),
            'approval_id': uuid.uuid4().hex,
        }

    def commit(self, plan) -> dict:
        tmpl = self.path(plan['template'])
        if digest(tmpl) != plan['sha256']:
            raise RuntimeError("File template đã thay đổi sau khi xin duyệt. Hãy thử lại.")
        output_path = Path(plan['output'])
        if output_path.exists():
            raise FileExistsError(f"File đích đã tồn tại: {output_path}.")
        if not any(output_path.parent.resolve() == root for root in self.roots):
            raise PermissionError("Thư mục đích nằm ngoài whitelist.")

        self.audit('template_fill_started', {'template': str(tmpl), 'output': str(output_path),
                                              'placeholders': len(plan['values'])})
        temp_path = None
        try:
            fd, temp = tempfile.mkstemp(prefix='.tmpl_', suffix='.docx', dir=output_path.parent)
            os.close(fd)
            temp_path = Path(temp)
            shutil.copy2(tmpl, temp_path)

            doc = load_docx(temp_path)
            values = plan['values']

            for para in doc.paragraphs:
                _fill_paragraph(para, values)
            _fill_tables(doc.tables, values)

            saved = Path(temp_path.parent / (temp_path.name + '_v2.docx'))
            doc.save(str(saved))

            os.replace(saved, output_path)
            temp_path.unlink(missing_ok=True)
            temp_path = None
        except Exception as exc:
            self.audit('template_fill_failed', {'error': str(exc)})
            raise
        finally:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

        result = {'ok': True, 'output': str(output_path), 'template': str(tmpl),
                  'placeholders_filled': len(plan['values'])}
        self.audit('template_fill_success', result)
        return result
