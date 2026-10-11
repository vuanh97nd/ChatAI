"""Render explicit TeX delimiters; leave code and spreadsheet formulas untouched."""
import html
import io
import re
from functools import lru_cache

PARTS = re.compile(r'(```[\s\S]*?```|`[^`\n]+`)|\$\$([\s\S]+?)\$\$|\\\[([\s\S]+?)\\\]|\\\(([\s\S]+?)\\\)|(?<!\\)\$([^\$\n]+?)\$')


def formulas(text):
    for match in PARTS.finditer(text):
        if not match.group(1):
            yield match, next(value for value in match.groups()[1:] if value is not None)


@lru_cache(maxsize=128)
def math_png(tex, color='#202020', size=14):
    if len(tex)>4000:
        raise ValueError('Công thức quá dài.')
    # No external TeX process, network, or model-supplied commands are executed.
    from matplotlib.mathtext import math_to_image
    output=io.BytesIO()
    math_to_image('$'+tex.strip()+'$',output,format='png',dpi=160,color=color,
                  prop=__import__('matplotlib.font_manager',fromlist=['FontProperties']).FontProperties(size=size))
    return output.getvalue()


def word_math(paragraph, text):
    if not list(formulas(text)):
        paragraph.add_run(text)
        return paragraph
    from docx.oxml import parse_xml
    from latex2mathml.converter import convert
    from mathml2omml import convert as to_omml
    offset=0
    for match,tex in formulas(text):
        paragraph.add_run(text[offset:match.start()])
        try:
            equation=to_omml(convert(tex)).replace('<m:oMath>', '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">', 1)
            # Converter produces an editable native Office Math element.
            paragraph._p.append(parse_xml(equation))
        except Exception:
            paragraph.add_run(match.group(0))
        offset=match.end()
    paragraph.add_run(text[offset:])
    return paragraph


def excel_math(value):
    if not isinstance(value,str) or value.startswith('='):
        return value
    if not list(formulas(value)):
        return value
    from latex2mathml.converter import convert
    from xml.etree import ElementTree as ET
    def linear(node):
        name=node.tag.rsplit('}',1)[-1]
        children=list(node)
        if name=='mfrac':return '('+linear(children[0])+')/('+linear(children[1])+')'
        if name=='msqrt':return '√('+''.join(map(linear,children))+')'
        if name in ('msup','msub'):
            base,script=map(linear,children[:2])
            alphabet='0123456789+-=()n' if name=='msup' else '0123456789+-=()'
            glyphs='⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ' if name=='msup' else '₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎'
            if all(c in alphabet for c in script):return base+script.translate(str.maketrans(alphabet,glyphs))
            return base+('^' if name=='msup' else '_')+'('+script+')'
        return (node.text or '')+''.join(linear(c)+(c.tail or '') for c in children)
    offset=0;parts=[]
    for match,tex in formulas(value):
        parts.append(value[offset:match.start()])
        try:parts.append(linear(ET.fromstring(convert(tex))))
        except Exception:parts.append(match.group(0))
        offset=match.end()
    return ''.join(parts)+value[offset:]


def word_document_text(document, text):
    """Keep multiline equations together while preserving ordinary paragraphs."""
    paragraph=document.add_paragraph()
    offset=0
    def add_text(value, current):
        lines=value.split('\n')
        current.add_run(lines[0])
        for line in lines[1:]:
            current=document.add_paragraph();current.add_run(line)
        return current
    for match,_ in formulas(text):
        paragraph=add_text(text[offset:match.start()],paragraph)
        word_math(paragraph,match.group(0))
        offset=match.end()
    add_text(text[offset:],paragraph)
