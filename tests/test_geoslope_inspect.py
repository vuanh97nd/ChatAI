import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock

from assistant.files import FileTools
from assistant.geoslope_inspect import GeoslopeInspect,inspect_gsz,inspect_workbook,parallel_layers
from assistant.online_automation import OnlineAutomation,requested_automation,parse_plan
from assistant.tools import EXTRA_TOOLS


class GeoslopeInputTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.files=FileTools([str(self.root)],self.root/'backup',lambda *a:None)
        self.tool=GeoslopeInspect(self.files)
    def tearDown(self):self.temp.cleanup()

    def test_sheet_names_and_section_columns_are_not_fixed_and_sources_survive(self):
        from openpyxl import Workbook
        w=Workbook();s=w.active;s.title='Bang tong hop moi'
        s['B3']='Mặt cắt tính toán lựa chọn';s['E3']='Bề dày lớp đất (m)'
        s['E4']='1a';s['F4']='1c';s['G4']='3'
        s['B5']='Km134+300';s['E5']=4.5;s['F5']=3;s['G5']=5.5
        m=w.create_sheet('Chi tieu moi')
        m.append([]);m.append(['Tên chỉ tiêu','Ký hiệu','Đơn vị','Lớp 1a'])
        m.append(['Dung trọng','g','kN/m3',16.1]);m.append(['Lực dính','c','kPa',14.7])
        m.append(['Góc ma sát trong','phi','độ',0])
        p=self.root/'inputs.xlsx';w.save(p)
        r=inspect_workbook(p)
        section=r['section_tables'][0]['sections'][0]
        self.assertEqual(section['section'],134300)
        self.assertEqual([x['thickness'] for x in section['layers']],[4.5,3,5.5])
        self.assertEqual(section['layers'][0]['cell'],'E5')
        self.assertEqual(r['material_tables'][0]['materials'][0]['fields']['unit_weight']['cell'],'D3')
        self.assertEqual(r['material_tables'][0]['materials'][0]['fields']['phi']['value'],0)

    def test_missing_formula_cache_is_not_zero_and_formula_is_retained(self):
        from openpyxl import Workbook
        w=Workbook();s=w.active;s['B2']='=1+2';s['D3']='#REF!'
        p=self.root/'formulas.xlsx';w.save(p)
        r=inspect_workbook(p,s.title,'A1:E4')
        cells={c['cell']:c for c in r['cells']}
        self.assertIsNone(cells['B2']['value']);self.assertEqual(cells['B2']['formula'],'=1+2')
        self.assertTrue(cells['D3']['error'])

    def test_all_stored_slips_are_read_without_claiming_a_new_solve(self):
        p=self.root/'reference.gsz'
        xml='<GSIData AppVersion="25.1.1"><Analyses><Analysis><ID>1</ID><Kind>SLOPE/W</Kind><Method>Bishop</Method></Analysis></Analyses><Materials><Material><ID>1</ID><Name>Soil</Name><SlopeModel>MohrCoulomb</SlopeModel><StressStrain><UnitWeight>18</UnitWeight></StressStrain></Material></Materials><Contexts><Context><AnalysisID>1</AnalysisID><GeometryUsesMaterials><GeometryUsesMaterial ID="Regions-1" Entry="1"/></GeometryUsesMaterials></Context></Contexts></GSIData>'
        with zipfile.ZipFile(p,'w') as z:
            z.writestr('model.xml',xml)
            z.writestr('Slope/001/slip_surface.csv','SlipNum,SlipFOS\n'+''.join(f'{i},2\n' for i in range(130))+'130,0.5863037247541445\n131,nan\n132,-1\n')
        r=inspect_gsz(p)
        self.assertEqual(r['stored_results'][0]['rows'],133)
        self.assertEqual(r['stored_results'][0]['minimum']['slip_number'],'130')
        self.assertFalse(r['new_solve_performed']);self.assertTrue(r['materials'][0]['used'])

    def test_invented_gsz_schema_and_xml_entities_are_rejected(self):
        p=self.root/'bad.gsz'
        for xml in ['<SLOPE_MODEL/>','<!DOCTYPE GSIData [<!ENTITY x "x">]><GSIData/>']:
            with zipfile.ZipFile(p,'w') as z:z.writestr('model.xml',xml)
            with self.assertRaises(ValueError):inspect_gsz(p)

    def test_vertical_layer_boundaries_keep_x_and_cumulative_thickness(self):
        profile=[[-2,1],[0,0],[3,2]]
        r=parallel_layers(profile,[{'code':'1a','thickness':4.5},{'code':'1c','thickness':3},{'code':'3','thickness':5.5}])
        self.assertEqual([x['depth'] for x in r['boundaries']],[0,4.5,7.5,13])
        self.assertEqual(r['boundaries'][-1]['points'],[[-2.,-12.],[0.,-13.],[3.,-11.]])
        for profile,layers in [([[0,0],[0,2]],[{'code':'1','thickness':1}]),([[0,0],[1,2]],[{'code':'1','thickness':None}])]:
            with self.assertRaises(ValueError):parallel_layers(profile,layers)

    def test_profile_uses_selected_handle_and_never_changes_source_or_overwrites_output(self):
        import ezdxf
        p=self.root/'source.dxf';d=ezdxf.new();e=d.modelspace().add_lwpolyline([[-2,1],[0,0],[3,2]])
        d.saveas(p);before=p.read_bytes()
        plan=self.tool.prepare('geoslope_profile',{'path':str(p),'handle':e.dxf.handle,'units':'m',
                 'layers':json.dumps([{'code':'1a','thickness':4.5,'source':'BTH!E5'}])})
        result=self.tool.commit(plan)
        self.assertEqual(before,p.read_bytes());self.assertFalse(result['new_solve_performed'])
        out=ezdxf.readfile(result['path']);self.assertEqual(out.header['$INSUNITS'],6)
        self.assertTrue(any(e.closed for e in out.modelspace().query('LWPOLYLINE')))
        with self.assertRaises(FileExistsError):self.tool.commit(plan)

    def test_whitelist_and_geoslope_online_tool_routing(self):
        with self.assertRaises(PermissionError):self.tool.prepare('geoslope_inspect',{'path':str(Path(__file__).resolve())})
        agent=OnlineAutomation(Mock(),{},Mock(),'c',Mock(),Mock(),geoslope_inspector=self.tool)
        names={s['function']['name'] for s in agent.schemas}
        self.assertIn('geoslope_inspect',names);self.assertIn('geoslope_profile',names)
        self.assertNotIn('geoslope_create',names)
        self.assertIs(agent.component('geoslope_profile'),self.tool)
        self.assertTrue(requested_automation('Chạy GeoStudio từ BTH và DXF này'))
        self.assertFalse(requested_automation('GeoStudio là gì?'))
        plan=parse_plan(json.dumps({'tool':'geoslope_profile','arguments':{
            'path':'sample.dxf','handle':'AB','units':'m','layers':[{'code':'1a','thickness':4.5}]}}),[s for _,s in EXTRA_TOOLS])
        self.assertEqual(json.loads(plan['arguments']['layers'])[0]['thickness'],4.5)


if __name__=='__main__':unittest.main()
