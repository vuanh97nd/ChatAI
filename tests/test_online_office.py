import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from openpyxl import Workbook,load_workbook
from docx import Document
from pptx import Presentation

from assistant.excel import ExcelTools
from assistant.files import FileTools
from assistant.office import OfficeTools
from assistant.online_office import OnlineOffice
from assistant.online_automation import OnlineAutomation,requested_automation
from assistant.storage import Store
from assistant.procedure_memory import ProcedureMemory
from assistant.procedure_sync import ProcedureSync
from test_procedure_sync import Backend


class OnlineOfficeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.files=FileTools([str(self.root)],self.root/'backup',lambda *a:None)
        self.office=OnlineOffice(ExcelTools([self.root],self.root/'backup',lambda *a:None),OfficeTools(self.files),self.files)
        with patch.dict(os.environ,{'LOCALAPPDATA':''}):self.store=Store(self.root/'history.db')
        self.cid=self.store.create(False)
        self.agent=OnlineAutomation(Mock(),{'procedure_training_enabled':True},self.store,self.cid,Mock(),Mock(),office_tools=self.office)

    def execute(self,name,args):
        state=self.store.load(self.cid)
        state.update(account_username='admin',running=True,queue=[{'function':{'name':name,'arguments':args}}],
                     messages=[{'role':'user','content':'Tạo bảng Excel và tài liệu Word PowerPoint theo mẫu.'}])
        state['pending']={'plan':self.office.prepare(name,args),'decision_started':False}
        self.agent.approve(state,True,state['pending'].copy())
        result=json.loads(next(m for m in reversed(state['messages']) if m['role']=='tool')['content'])
        self.assertTrue(result['ok'],result)
        return result

    def test_online_excel_batch_writes_and_reads_actual_workbook(self):
        path=self.root/'data.xlsx';w=Workbook();w.save(path);w.close()
        result=self.execute('excel_edit_batch',{'path':str(path),'sheet':'Sheet','fills':json.dumps([{'cell':'B2','value':10},{'cell':'B3','value':5}])})
        w=load_workbook(path);self.assertEqual(w.active['B2'].value,10);w.close()
        self.assertTrue(Path(result['backup']).exists())
        result=self.execute('excel_read',{'path':str(path),'sheet':'Sheet','cell_range':'B2:B3'})
        self.assertEqual(result['rows'],[[10],[5]])

    def test_duplicate_reads_are_cached_but_external_changes_are_reread(self):
        path=self.root/'data.xlsx';w=Workbook();w.active['A1']=10;w.save(path);w.close()
        args={'path':str(path),'sheet':'Sheet','cell_range':'A1:A1'}
        with patch.object(self.office.excel,'excel_read',wraps=self.office.excel.excel_read) as read:
            self.execute('excel_read',args)
            second=self.execute('excel_read',args)
            self.assertTrue(second['cached']);self.assertEqual(read.call_count,1)
            w=load_workbook(path);w.active['A1']=20;w.save(path);w.close()
            third=self.execute('excel_read',args)
            self.assertEqual(third['rows'],[[20]]);self.assertEqual(read.call_count,2)

    def test_online_word_creation_and_replacement(self):
        path=self.root/'report.docx'
        self.execute('office_create',{'path':str(path),'title':'Report','content':'Old text'})
        self.execute('word_replace',{'path':str(path),'search':'Old text','replacement':'New text'})
        self.assertIn('New text','\n'.join(p.text for p in Document(path).paragraphs))

    def test_online_powerpoint_creates_two_slides(self):
        path=self.root/'slides.pptx'
        self.execute('office_create',{'path':str(path),'title':'Deck','content':'Slide 1\nFirst\n---\nSlide 2\nSecond'})
        self.assertEqual(len(Presentation(path).slides),2)

    def test_tools_are_exposed_and_office_requests_are_routed(self):
        names={s['function']['name'] for s in self.agent.schemas}
        self.assertTrue({'excel_edit_batch','excel_read','office_create','word_replace','office_read','file_list'}<=names)
        for prompt in ('Ghi số lượng vào Excel','Tạo PowerPoint báo cáo','Sửa nội dung Word'):
            self.assertTrue(requested_automation(prompt))

    def test_cloudflare_can_use_existing_server_planning_endpoint(self):
        from assistant.cloud import ServerApiClient
        from assistant.accounts import request_account
        client=ServerApiClient({'endpoint':'https://example.org','username':'alice','key':'fixture'},'cloudflare')
        with patch('assistant.accounts.request_account',return_value={'success':True,'answer':'{"answer":"OK","tool":"","arguments":"{}"}'}) as request:
            client.chat(client.model,[{'role':'user','content':'Tạo PowerPoint'}],format={'type':'object'})
        self.assertEqual(request.call_args.args[1],'/api/provider/model')
        self.assertEqual(request.call_args.args[2]['provider'],'cloudflare')

    def test_successful_office_operation_is_shared_and_recalled_by_another_user(self):
        self.execute('office_create',{'path':str(self.root/'report.docx'),'title':'Report','content':'Example'})
        backend=Backend();session={'username':'admin','key':'fixture','endpoint':'https://example.org','role':'system'}
        ProcedureSync(self.store,session,backend).cycle()
        shared=[row['record'] for (owner,_),row in backend.rows.items() if owner=='@shared']
        self.assertEqual(len(shared),1)
        self.assertNotIn(str(self.root),json.dumps(shared))
        self.assertNotIn('Example',json.dumps(shared))
        with patch.dict(os.environ,{'LOCALAPPDATA':''}):other=Store(self.root/'other.db')
        ProcedureSync(other,dict(session,username='bob',role='user'),backend).cycle()
        self.assertIn('office_create',ProcedureMemory(other).context('bob','office_create Word'))


if __name__=='__main__':unittest.main()
