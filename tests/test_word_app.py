import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
from assistant.files import FileTools
from assistant.word_app import WordApp


class WordAppTest(unittest.TestCase):
    def test_creates_real_document_before_open_and_never_overwrites(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);exe=root/'WINWORD.EXE';exe.write_bytes(b'fixture')
            files=FileTools([str(root)],root/'backups',lambda *a:None)
            windows=SimpleNamespace(check=lambda: {},allowed_path=lambda p:Path(p))
            tool=WordApp(windows,files)
            plan=tool.prepare('word_create_open',{'app':str(exe),'title':'Mẹ','content':'Mẹ luôn yêu thương tôi.'})
            with patch('assistant.word_app.subprocess.Popen') as launch:
                result=tool.commit(plan)
                self.assertTrue(result['document_created'])
                self.assertIn('Mẹ luôn yêu thương tôi.', '\n'.join(p.text for p in Document(result['path']).paragraphs))
                doc=Document(result['path'])
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                for paragraph in doc.paragraphs:
                    for run in paragraph.runs:
                        self.assertEqual(run.font.name,'Times New Roman');self.assertEqual(run.font.size.pt,13)
                self.assertEqual(doc.paragraphs[0].alignment,WD_ALIGN_PARAGRAPH.CENTER)
                self.assertEqual(doc.paragraphs[1].alignment,WD_ALIGN_PARAGRAPH.JUSTIFY)
                self.assertAlmostEqual(doc.sections[0].left_margin.cm,3,places=2)
                launch.assert_called_once_with([str(exe),result['path']],shell=False)
                with self.assertRaises(RuntimeError):tool.commit(plan)
            exe.write_bytes(b'changed')
            with self.assertRaises(PermissionError):tool.commit(plan)

    def test_named_file_and_overwrite_preserve_backup(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);exe=root/'WINWORD.EXE';exe.write_bytes(b'fixture')
            files=FileTools([str(root)],root/'backups',lambda *a:None)
            tool=WordApp(SimpleNamespace(check=lambda:{},allowed_path=lambda p:Path(p)),files)
            args={'app':str(exe),'path':'Don-xin-nghi.docx','title':'Đơn','content':'Nội dung cũ'}
            with patch('assistant.word_app.subprocess.Popen'):
                first=tool.commit(tool.prepare('word_create_open',args))
                self.assertEqual(Path(first['path']).name,'Don-xin-nghi.docx')
                old=Path(first['path']).read_bytes()
                second=tool.commit(tool.prepare('word_create_open',args))
                self.assertEqual(Path(second['path']).name,'Don-xin-nghi-2.docx')
                third=tool.prepare('word_create_open',args)
                self.assertEqual(Path(third['document']['path']).name,'Don-xin-nghi-3.docx')
                self.assertEqual(Path(first['path']).read_bytes(),old)
                args.update(mode='overwrite',content='Nội dung mới')
                result=tool.commit(tool.prepare('word_create_open',args))
                self.assertEqual(Path(result['backup']).read_bytes(),old)
                self.assertIn('Nội dung mới','\n'.join(p.text for p in Document(result['path']).paragraphs))
                plan=tool.prepare('word_create_open',args)
                Document().save(result['path'])
                with self.assertRaises(RuntimeError):tool.commit(plan)
