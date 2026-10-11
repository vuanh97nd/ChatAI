import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from assistant.excel import ExcelTools

class ExcelBatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.file=self.root/'data.xlsx'
        wb=Workbook();wb.active['A1']='old';wb.save(self.file);wb.close()
        self.tools=ExcelTools([self.root],self.root/'backups',lambda *args:None)
    def test_many_cells_are_saved_once_with_backup_and_literal_text(self):
        plan=self.tools.prepare_batch(str(self.file),'Sheet',json.dumps([
            {'cell':'A1','value':'new'},{'cell':'B1','value':42},{'cell':'C1','value':'=SUM(A1:B1)'}]))
        with patch('assistant.excel.load_workbook',wraps=load_workbook) as opened:
            result=self.tools.commit_edit(plan)
            self.assertEqual(opened.call_count,1)
        self.assertEqual(result['cells_written'],3)
        wb=load_workbook(self.file)
        self.assertEqual(wb.active['A1'].value,'new');self.assertEqual(wb.active['B1'].value,42)
        self.assertEqual(wb.active['C1'].data_type,'s');wb.close()
        wb=load_workbook(result['backup']);self.assertEqual(wb.active['A1'].value,'old');wb.close()
    def test_changed_file_is_not_overwritten(self):
        plan=self.tools.prepare_batch(str(self.file),'Sheet',[{'cell':'A1','value':'new'}])
        wb=load_workbook(self.file);wb.active['A1']='changed';wb.save(self.file);wb.close()
        with self.assertRaises(RuntimeError):self.tools.commit_edit(plan)
    def test_duplicate_or_invalid_cells_are_rejected(self):
        for fills in [[{'cell':'A1','value':1},{'cell':'a1','value':2}], [{'cell':'XFE1','value':1}], []]:
            with self.assertRaises(ValueError):self.tools.prepare_batch(str(self.file),'Sheet',fills)
