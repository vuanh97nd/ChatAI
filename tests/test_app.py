import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from assistant.agent import Agent, bounded_context
from assistant.excel import ExcelTools, digest
from assistant.locking import execution_lock
from assistant.storage import Store


class FakeCall:
    def __init__(self, name, args):
        self.value = {"function": {"name": name, "arguments": args}}

    def model_dump(self, **kwargs):
        return self.value


def chunk(text="", calls=None):
    return SimpleNamespace(message=SimpleNamespace(content=text, tool_calls=calls or []))


class FakeClient:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.requests = []

    def chat(self, **kwargs):
        if kwargs.get("stream") is False:
            return {"message": {"content": json.dumps({"category": "knowledge",
                "creative": False, "complex": False, "high_accuracy": False})}}
        self.requests.append(kwargs)
        return iter(next(self.responses))


class AppTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.allowed = self.root / "allowed"
        self.allowed.mkdir()
        self.book = self.allowed / "demo.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Tên", "Số lượng"])
        ws.append(["A", 10])
        ws.append(["B", 20])
        ws["C2"] = "=B2*2"
        wb.save(self.book)
        wb.close()
        self.store = Store(self.root / "history.sqlite3")
        self.cid = self.store.create()
        self.excel = ExcelTools([self.allowed], self.root / "backups",
                               lambda action, details: self.store.audit(self.cid, action, details))
        self.cfg = {"default_model": "qwen2.5:7b", "code_model": "qwen2.5-coder:7b",
                    "num_ctx": 4096, "num_predict": 1000, "temperature": 0.2, "max_rounds": 4}

    def tearDown(self):
        self.temp.cleanup()

    def value(self, cell):
        wb = load_workbook(self.book)
        try:
            return wb["Data"][cell].value
        finally:
            wb.close()

    def agent(self, responses):
        client = FakeClient(responses)
        return Agent(client, self.excel, self.cfg, self.store, self.cid), client

    def test_read_summary(self):
        self.assertEqual(self.excel.excel_list_sheets("demo.xlsx")["sheets"], ["Data"])
        rows = self.excel.excel_read("demo.xlsx", "Data", "B2:C2")["rows"]
        self.assertEqual(rows, [[10, "=B2*2"]])
        summary = self.excel.excel_summary("demo.xlsx", "Data")
        self.assertEqual(summary["data_rows"], 2)
        self.assertEqual(summary["columns"][1]["stats"]["mean"], 15)

    def test_whitelist_and_range(self):
        outside = self.root / "outside.xlsx"
        outside.write_bytes(self.book.read_bytes())
        for path in (str(outside), "../outside.xlsx"):
            with self.assertRaises(PermissionError):
                self.excel.excel_list_sheets(path)
        with self.assertRaises(ValueError):
            self.excel.excel_read("demo.xlsx", "Data", "A1:ZZ99999")

    def test_symlink_escape(self):
        outside = self.root / "outside.xlsx"
        outside.write_bytes(self.book.read_bytes())
        link = self.allowed / "link.xlsx"
        try:
            link.symlink_to(outside)
        except OSError:
            self.skipTest("Hệ điều hành không cho tạo symlink.")
        with self.assertRaises(PermissionError):
            self.excel.path(str(link))

    def test_edit_backup_and_literal(self):
        original = self.book.read_bytes()
        plan = self.excel.prepare_edit("demo.xlsx", "Data", "A2", "=HYPERLINK(\"https://example.com\")")
        result = self.excel.commit_edit(plan)
        self.assertEqual(Path(result["backup"]).read_bytes(), original)
        wb = load_workbook(self.book)
        try:
            self.assertEqual(wb["Data"]["A2"].data_type, "s")
            self.assertEqual(wb["Data"]["C2"].value, "=B2*2")
        finally:
            wb.close()
        self.assertTrue(any(e["action"] == "edit_success" for e in self.store.recent_audit(self.cid)))

    def test_changed_file_rejected(self):
        plan = self.excel.prepare_edit("demo.xlsx", "Data", "B2", 99)
        wb = load_workbook(self.book)
        wb["Data"]["B2"] = 33
        wb.save(self.book)
        wb.close()
        with self.assertRaises(RuntimeError):
            self.excel.commit_edit(plan)
        self.assertEqual(self.value("B2"), 33)

    def test_agent_tools_then_stream(self):
        call = FakeCall("excel_list_sheets", {"path": "demo.xlsx"})
        agent, client = self.agent([[chunk(calls=[call])], [chunk("Có "), chunk("sheet Data.")]])
        state = self.store.load(self.cid)
        agent.start(state, "Liệt kê sheet", self.cfg["default_model"])
        events = list(agent.run(state))
        self.assertEqual("".join(e["text"] for e in events if e["type"] == "token"), "Có sheet Data.")
        self.assertFalse(state["running"])
        self.assertEqual(client.requests[1]["messages"][-1]["role"], "tool")
        self.assertEqual(self.store.load(self.cid), state)

    def test_pending_survives_reload_and_denial(self):
        calls = [FakeCall("excel_edit_cell", {"path": "demo.xlsx", "sheet": "Data", "cell": "B2", "value": 99}),
                 FakeCall("excel_list_sheets", {"path": "demo.xlsx"})]
        agent, _ = self.agent([[chunk(calls=calls)], [chunk("Đã từ chối sửa.")]])
        state = self.store.load(self.cid)
        before = digest(self.book)
        agent.start(state, "Sửa B2 và liệt kê sheet", self.cfg["default_model"])
        list(agent.run(state))
        self.assertEqual(digest(self.book), before)
        reloaded = self.store.load(self.cid)
        self.assertIsNotNone(reloaded["pending"])
        self.assertEqual(len(reloaded["queue"]), 2)
        agent.approve(reloaded, False)
        list(agent.run(reloaded))
        self.assertEqual(digest(self.book), before)
        tools = [m for m in reloaded["messages"] if m["role"] == "tool"]
        self.assertEqual(len(tools), 2)
        self.assertTrue(json.loads(tools[0]["content"])["denied"])

    def test_approval_once(self):
        call = FakeCall("excel_edit_cell", {"path": "demo.xlsx", "sheet": "Data", "cell": "B2", "value": 99})
        agent, _ = self.agent([[chunk(calls=[call])], [chunk("Đã sửa.")]])
        state = self.store.load(self.cid)
        agent.start(state, "Sửa B2", self.cfg["default_model"])
        list(agent.run(state))
        agent.approve(state, True)
        self.assertEqual(self.value("B2"), 99)
        with self.assertRaises(RuntimeError):
            agent.approve(state, True)
        self.assertEqual(len(list((self.root / "backups").glob("*.xlsx"))), 1)
        list(agent.run(state))
        self.assertFalse(state["running"])

    def test_unknown_tool_returns_error(self):
        call = FakeCall("run_command", {"command": "whoami"})
        agent, _ = self.agent([[chunk(calls=[call])], [chunk("Tool chưa có.")]])
        state = self.store.load(self.cid)
        agent.start(state, "Chạy lệnh", self.cfg["default_model"])
        list(agent.run(state))
        result = next(m for m in state["messages"] if m["role"] == "tool")
        self.assertFalse(json.loads(result["content"])["ok"])

    def test_crash_does_not_replay_write(self):
        call = FakeCall("excel_edit_cell", {"path": "demo.xlsx", "sheet": "Data", "cell": "B2", "value": 99})
        agent, _ = self.agent([[chunk(calls=[call])]])
        state = self.store.load(self.cid)
        agent.start(state, "Sửa B2", self.cfg["default_model"])
        list(agent.run(state))
        state["pending"]["decision_started"] = True
        self.store.save(self.cid, state)
        agent.recover_uncertain(state)
        self.assertEqual(self.value("B2"), 10)
        self.assertIsNone(state["pending"])
        self.assertTrue(json.loads(state["messages"][-1]["content"])["uncertain"])

    def test_round_limit(self):
        call = FakeCall("excel_list_files", {})
        agent, _ = self.agent([[chunk(calls=[call])] for _ in range(4)])
        state = self.store.load(self.cid)
        agent.start(state, "Lặp mãi", self.cfg["default_model"])
        list(agent.run(state))
        self.assertEqual(state["rounds"], 4)
        self.assertFalse(state["running"])
        self.assertFalse(state["queue"])

    def test_context_preserves_tool_pairs(self):
        messages = [{"role": "user", "content": "x"*10000},
                    {"role": "assistant", "content": "x"*10000},
                    {"role": "user", "content": "sheets"},
                    {"role": "assistant", "content": "", "tool_calls": []},
                    {"role": "tool", "tool_name": "excel_list_sheets", "content": "{}"}]
        self.assertEqual(bounded_context(messages), messages[2:])

    def test_lock_blocks_second_execution(self):
        path = self.root / "execution.lock"
        with execution_lock(path):
            with self.assertRaises(RuntimeError):
                with execution_lock(path):
                    pass
        with execution_lock(path):
            pass

    def test_current_turn_context_is_bounded_without_changing_history(self):
        messages = [{"role": "user", "content": "Tóm tắt"},
                    {"role": "assistant", "content": "", "tool_calls": []},
                    {"role": "tool", "tool_name": "excel_read", "content": "x"*30000}]
        result = bounded_context(messages)
        self.assertLessEqual(len(json.dumps(result, ensure_ascii=False)), 9000)
        self.assertEqual(len(messages[-1]["content"]), 30000)
        self.assertTrue(json.loads(result[-1]["content"])["truncated_for_context"])

    def test_model_failure_keeps_history_and_stops(self):
        agent, _ = self.agent([])
        state = self.store.load(self.cid)
        agent.start(state, "Xin chào", self.cfg["default_model"])
        list(agent.run(state))
        self.assertFalse(state["running"])
        self.assertIn("Lỗi Ollama", state["messages"][-1]["content"])
        self.assertEqual(self.store.load(self.cid), state)


if __name__ == "__main__":
    unittest.main()
