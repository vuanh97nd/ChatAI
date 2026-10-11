import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,patch

from assistant.files import FileTools
from assistant.excel import ExcelTools
from assistant.office import OfficeTools
from assistant.office_batch import OfficeBatch
from assistant.online_office import OnlineOffice
from assistant.online_automation import OnlineAutomation
from assistant.task_tools import TaskTools
from assistant.com_tasks import ComTools,execute,check_identity
from assistant.windows_apps import _STOP,resume_automation


class OnlineTaskToolsTests(unittest.TestCase):
    def setUp(self):
        resume_automation();self.addCleanup(resume_automation)
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        self.files=FileTools([self.root],self.root/'backups',lambda *a:None)
        self.office=OnlineOffice(ExcelTools([self.root],self.root/'backups',lambda *a:None),OfficeTools(self.files),self.files)

    def test_create_excel_table_formula_filter_chart_and_verify(self):
        p=self.root/'data.xlsx'
        plan=self.office.prepare('excel_create',{'path':str(p),'sheet':'Data','rows':json.dumps([['Item','Qty'],['A',10],['B',5]]),
                    'formulas':json.dumps({'B4':'=SUM(B2:B3)'}),'options':json.dumps({'bold_header':True,'auto_filter':True,'chart':{'range':'A1:B3','title':'Qty'}}),'open':False})
        result=self.office.commit(plan)
        from openpyxl import load_workbook
        w=load_workbook(p);self.assertEqual(w['Data']['B4'].data_type,'f');self.assertTrue(w['Data']['A1'].font.bold)
        self.assertEqual(w['Data'].auto_filter.ref,'A1:B4');self.assertEqual(len(w['Data']._charts),1);w.close()
        self.assertTrue(result['verified']);self.assertFalse(result['readback']['formula_values_verified'])

    def test_word_batch_table_header_footer_style_and_equation(self):
        p=self.root/'report.docx'
        ops=[{'action':'append','text':'Report'},{'action':'paragraph','index':0,'bold':True,'size':16,'alignment':'center'},
             {'action':'table','rows':[['A','B'],['10','5']]},{'action':'table_cell','table':0,'row':1,'column':1,'text':'7'},
             {'action':'header','text':'Header'},{'action':'footer','text':'Footer'},{'action':'equation','latex':'x^2+1'}]
        result=self.office.commit(self.office.prepare('word_edit',{'path':str(p),'create':True,'operations':json.dumps(ops)}))
        self.assertTrue(result['verified']);self.assertEqual(result['readback']['equations'],1)
        self.assertEqual(result['readback']['tables'][0]['rows'][1],['10','7'])
        self.assertEqual(result['readback']['sections'][0]['header'],'Header')
        from docx import Document
        self.assertTrue(Document(p).paragraphs[0].runs[0].bold)

    def test_native_file_write_read_move_and_backup(self):
        tools=TaskTools(self.files);p=self.root/'code.ps1';p.write_text('old',encoding='utf-8')
        result=tools.commit(tools.prepare('file_write',{'path':str(p),'content':'Write-Output 15'}))
        self.assertTrue(result['verified']);self.assertEqual(Path(result['backup']).read_text(),'old')
        self.assertIn('15',tools.commit(tools.prepare('file_read',{'path':str(p)}))['content'])
        result=tools.commit(tools.prepare('file_move',{'path':str(p),'destination':str(self.root/'new.ps1')}))
        self.assertTrue(result['verified']);self.assertFalse(p.exists())

    def test_task_script_runs_after_preview_with_backup_and_readback(self):
        p=self.root/'result.txt';p.write_text('old',encoding='utf-8')
        args={'language':'python','path':str(self.root),'files':[str(p)],'code':"from pathlib import Path\nPath('result.txt').write_text('15', encoding='utf-8')\nprint('done')"}
        plan=self.office.prepare('task_script',args);self.assertTrue(plan['requires_code_approval']);self.assertEqual(p.read_text(),'old')
        result=self.office.commit(plan);self.assertTrue(result['ok'],result);self.assertEqual(p.read_text(),'15')
        self.assertEqual(Path(result['backups'][0]).read_text(),'old');self.assertTrue(result['files'][0]['exists'])

    def test_cancel_stops_script_and_returns_uncertain_not_success(self):
        import threading
        plan=self.office.prepare('task_script',{'language':'python','path':str(self.root),'files':[],'code':'import time\ntime.sleep(30)','timeout':60})
        timer=threading.Timer(.4,_STOP.set);timer.start()
        try:result=self.office.commit(plan)
        finally:timer.cancel()
        self.assertFalse(result['ok']);self.assertTrue(result['uncertain'])

    def test_task_grant_does_not_authorize_another_task_or_unreviewed_code(self):
        agent=OnlineAutomation(Mock(),{'task_autonomy':True},Mock(),None,Mock(),Mock(),office_tools=self.office)
        state={'pending':{'plan':{'action':'excel_create'}},'task_autonomy':True}
        self.assertTrue(agent.auto_execute_allowed(state))
        state['pending']['plan']['requires_code_approval']=True;self.assertFalse(agent.auto_execute_allowed(state))
        state['pending']['plan']={'action':'excel_create'};state['task_autonomy']=False
        self.assertFalse(agent.auto_execute_allowed(state))
        state['pending']['plan']={'action':'calculate'};self.assertTrue(agent.auto_execute_allowed(state))
        names={s['function']['name'] for s in agent.schemas}
        self.assertTrue({'calculate','document_read','file_write','task_script','excel_com','word_edit','vba_project','cad_com'}<=names)

    def test_com_checks_exact_document_and_readback_before_success(self):
        p=self.root/'book.xlsx';p.write_bytes(b'fixture');backup=self.root/'backup.xlsx'
        target=SimpleNamespace(Value2=10,Formula='=5+5');sheet=SimpleNamespace(Range=lambda key:target)
        def save_copy(path):Path(path).write_bytes(b'backup')
        doc=SimpleNamespace(FullName=str(p),SaveCopyAs=save_copy,Save=Mock(),Worksheets=SimpleNamespace(Item=lambda name:sheet))
        app=SimpleNamespace(Workbooks=SimpleNamespace(Count=1,Item=lambda index:doc),Calculate=Mock())
        plan=ComTools(self.files).prepare('excel_com',{'path':str(p),'operations':json.dumps([{'action':'formula','sheet':'Data','range':'B2','value':'=5+5'}])})
        result=execute(plan,app);self.assertTrue(result['verified']);self.assertEqual(result['readback'][0]['values'],10)
        self.assertTrue(Path(result['backup']).exists());doc.Save.assert_called_once()
        with self.assertRaises(RuntimeError):check_identity(doc,self.root/'other.xlsx')

    def test_macro_requires_expected_result_and_code_approval(self):
        p=self.root/'book.xlsm';p.write_bytes(b'fixture');tools=ComTools(self.files)
        with self.assertRaises(ValueError):tools.prepare('vba_project',{'application':'excel','path':str(p),'operation':'run','macro':'Module1.Fill'})
        plan=tools.prepare('vba_project',{'application':'excel','path':str(p),'operation':'run','macro':'Module1.Fill',
                 'verify':json.dumps({'sheet':'Data','range':'A1','expected':[[15]]})})
        self.assertTrue(plan['requires_code_approval'])

    def test_macro_result_mismatch_never_reports_success(self):
        p=self.root/'book.xlsm';p.write_bytes(b'fixture')
        target=SimpleNamespace(Value2=0)
        doc=SimpleNamespace(FullName=str(p),Name='book.xlsm',Worksheets=SimpleNamespace(Item=lambda name:SimpleNamespace(Range=lambda cell:target)),
                            SaveCopyAs=lambda path:Path(path).write_bytes(b'backup'),Save=Mock())
        app=SimpleNamespace(Workbooks=SimpleNamespace(Count=1,Item=lambda i:doc),Run=Mock(),Calculate=Mock())
        plan=ComTools(self.files).prepare('vba_project',{'application':'excel','path':str(p),'operation':'run','macro':'Module1.Fill',
                 'verify':json.dumps({'sheet':'Data','range':'A1','expected':[[15]]})})
        with self.assertRaisesRegex(RuntimeError,'không khớp'):execute(plan,app)
        doc.Save.assert_not_called()

    def test_cad_command_waits_for_marker_and_restores_drawing_variable(self):
        from assistant.com_tasks import run_cad_command
        doc=Mock();marker={'value':2}
        def send(code):
            import re
            marker['value']=int(re.search(r'USERI1" (\d+)',code).group(1))
        doc.SendCommand.side_effect=send;doc.GetVariable.side_effect=lambda name:marker['value']
        run_cad_command(doc,'(princ)',timeout=1)
        doc.SetVariable.assert_called_once_with('USERI1',2)

    def test_outside_root_and_changed_file_are_rejected(self):
        with self.assertRaises(PermissionError):self.office.prepare('excel_create',{'path':str(self.root.parent/'outside.xlsx'),'sheet':'Data','rows':'[[1]]'})
        p=self.root/'data.xlsx';p.write_bytes(b'old')
        plan=ComTools(self.files).prepare('excel_com',{'path':str(p),'operations':'[{"action":"read","sheet":"Data","range":"A1"}]'})
        p.write_bytes(b'new')
        with self.assertRaises(RuntimeError):ComTools(self.files).commit(plan)


if __name__=='__main__':unittest.main()


class AutomaticCodePermissionTests(unittest.TestCase):
    def test_global_auto_execution_includes_code_plans(self):
        from unittest.mock import Mock
        from assistant.online_automation import OnlineAutomation
        windows=Mock();windows.check.return_value={'ai_tools_auto_execute':True}
        agent=OnlineAutomation(Mock(),{'ai_tools_auto_execute':True},Mock(),None,windows,Mock())
        state={'pending':{'plan':{'action':'task_script','requires_code_approval':True}},'task_autonomy':True}
        self.assertTrue(agent.auto_execute_allowed(state))
        state['task_autonomy']=False
        self.assertFalse(agent.auto_execute_allowed(state))
