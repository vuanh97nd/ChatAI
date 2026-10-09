"""Giao diện desktop Windows bằng Qt; không khởi chạy HTTP server."""
import html
import math
import time
import re
import json
import sys
import traceback
from assistant.performance import measure,record
from threading import Event
from pathlib import Path
print("Chat AI Desktop 2.6.6: giao diện không chờ Ollama/SSL.",flush=True)

from assistant.runtime_compat import prepare_six
prepare_six()

from PySide6.QtCore import QThread, Signal, QTimer, Qt, QUrl, QByteArray, QBuffer, QIODevice, QRectF, QPropertyAnimation, QSize, QEasingCurve
from PySide6.QtGui import QDesktopServices, QTextDocument, QTextCursor, QIcon, QPixmap, QImage, QKeySequence, QPainter, QColor, QRadialGradient, QPainterPath, QTextTable, QTextCharFormat, QTextFormat
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
    QHBoxLayout, QListWidget, QListWidgetItem, QLabel, QPushButton, QComboBox,
    QPlainTextEdit, QTextBrowser, QTabWidget, QSplitter, QMessageBox, QDialog,
    QDialogButtonBox, QProgressBar, QFileDialog, QStackedWidget, QScrollArea, QFrame, QFormLayout, QGroupBox, QSpinBox, QDoubleSpinBox, QLineEdit, QCheckBox, QMenu, QGraphicsOpacityEffect, QGraphicsDropShadowEffect, QInputDialog, QStyle)
from assistant import initialize_runtime
initialize_runtime()

from assistant.config import ROOT, load_config
from assistant.locking import execution_lock
from assistant.modules import ModuleManager, MODULES, ALLOWED_MODELS, CHAT_MODELS, MEDIA_VARIANTS, media_model_dir
from assistant.storage import Store
from assistant.trial import GuestTrial, GUEST_OWNER


class WelcomeRobot(QWidget):
    def __init__(self,parent=None):
        super().__init__(parent)
        self.setFixedSize(112,112)
        self.pixmap=QPixmap(str(ROOT/'logo_chat_ai.png'))
        self.started=time.monotonic()
        self.timer=QTimer(self);self.timer.setInterval(66);self.timer.timeout.connect(self.update)
    def showEvent(self,event):
        super().showEvent(event);self.timer.start()
    def hideEvent(self,event):
        self.timer.stop();super().hideEvent(event)
    def paintEvent(self,event):
        elapsed=time.monotonic()-self.started
        pulse=(math.sin(elapsed*math.pi)+1)/2;y=56-3*math.sin(elapsed*math.pi)
        painter=QPainter(self);painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        glow=QRadialGradient(56,y,55)
        glow.setColorAt(0,QColor(104,66,255,int(35+45*pulse)))
        glow.setColorAt(.65,QColor(38,149,255,int(20+25*pulse)))
        glow.setColorAt(1,QColor(38,149,255,0))
        painter.setPen(Qt.PenStyle.NoPen);painter.setBrush(glow)
        painter.drawEllipse(QRectF(1,y-55,110,110))
        painter.drawPixmap(QRectF(16,y-40,80,80),self.pixmap,QRectF(self.pixmap.rect()))


class FadingStatusLabel(QLabel):
    """QLabel that fades in each time setText is called."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._opacity = QGraphicsOpacityEffect(self)
        self._opacity.setOpacity(1.0)
        self.setGraphicsEffect(self._opacity)
        self._anim = QPropertyAnimation(self._opacity, b'opacity', self)
        self._anim.setDuration(280)
        self._anim.setEasingCurve(QEasingCurve.Type.OutCubic)

    def setText(self, text):
        super().setText(text)
        self._anim.stop()
        self._anim.setStartValue(0.0)
        self._anim.setEndValue(1.0)
        self._anim.start()


class PromptEdit(QPlainTextEdit):
    sendRequested = Signal()
    imagePasted = Signal(object)
    filesPasted = Signal(object)

    def __init__(self):
        super().__init__()
        self.composing = False
        self.setTabChangesFocus(True)

    def insertFromMimeData(self, source):
        if source.hasImage():
            value=source.imageData()
            image=value.toImage() if isinstance(value,QPixmap) else QImage(value)
            if not image.isNull():self.imagePasted.emit(image); return
        if source.hasUrls():
            paths=[url.toLocalFile() for url in source.urls() if url.isLocalFile()]
            if paths:self.filesPasted.emit(paths);return
        super().insertFromMimeData(source)

    def inputMethodEvent(self, event):
        self.composing = bool(event.preeditString())
        super().inputMethodEvent(event)

    def keyPressEvent(self, event):
        if (event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter)
                and event.modifiers() in (Qt.KeyboardModifier.NoModifier, Qt.KeyboardModifier.KeypadModifier)
                and not self.composing):
            if not event.isAutoRepeat(): self.sendRequested.emit()
            event.accept()
            return
        super().keyPressEvent(event)


class ChatCancelled(BaseException):
    """Hủy hợp tác; không bị các nhánh bắt lỗi công cụ nuốt mất."""


class Worker(QThread):
    event = Signal(object)

    def __init__(self, fn):
        super().__init__()
        self.fn, self.result, self.failure = fn, None, None
        self.stop_requested=Event();self.cancelled=False;self.cancellable=False;self.app_countdown=False;self.app_countdown_done=False
        self.partial='';self.snapshot_count=0;self.chat_cid=None

    def emit_event(self,event):
        if self.cancellable and self.stop_requested.is_set():raise ChatCancelled()
        if event.get('type')=='app_activity' and self.app_countdown and not self.app_countdown_done:
            self.app_countdown_done=True
            for seconds in (3,2,1):
                self.event.emit({'type':'app_countdown','seconds':seconds})
                if self.stop_requested.wait(1):raise ChatCancelled()
            self.event.emit({'type':'app_countdown','seconds':0})
        if event.get('type')=='token':self.partial+=event.get('text','')
        elif event.get('type')=='snapshot':self.partial='';self.snapshot_count=len(event.get('messages',[]))
        self.event.emit(event)

    def run(self):
        try:
            self.result = self.fn(self.emit_event)
            self.cancelled=self.cancellable and self.stop_requested.is_set()
        except ChatCancelled:
            self.cancelled=True
        except Exception as exc:
            self.failure = str(exc)
            with (ROOT / 'data/startup.log').open('a', encoding='utf-8') as log:
                log.write(traceback.format_exc())


class ImagePreview(QLabel):
    def __init__(self):
        super().__init__(); self.image=None; self.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
        self.setToolTip('Bấm chọn ảnh, rồi Ctrl+C để sao chép')
    def keyPressEvent(self,event):
        if event.matches(QKeySequence.StandardKey.Copy) and self.image is not None:
            QApplication.clipboard().setImage(self.image); event.accept(); return
        super().keyPressEvent(event)
    def mousePressEvent(self,event):
        self.setFocus(); super().mousePressEvent(event)


class ChatView(QTextBrowser):
    codeActionRequested = Signal(object)
    def __init__(self):
        super().__init__()
        self.clipboard_images={}; self.selected_image=None;self.allowed_roots=[]
        self.setOpenLinks(False)
        self.setOpenExternalLinks(False)
        self.anchorClicked.connect(self.open_link)

    def enforce_text_size(self, pixels):
        # Apply to native text formats, including nested Markdown tables and spans.
        # CSS inheritance in QTextDocument does not reliably cover all fragments.
        points=pixels*72.0/max(1,self.logicalDpiY())
        font=self.font();font.setPointSizeF(points)
        self.document().setDefaultFont(font)
        cursor=QTextCursor(self.document())
        cursor.select(QTextCursor.SelectionType.Document)
        fmt=QTextCharFormat()
        fmt.setFontPointSize(points)
        cursor.mergeCharFormat(fmt)

    def wheelEvent(self,event):
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            event.accept();return
        super().wheelEvent(event)

    def paintEvent(self,event):
        super().paintEvent(event)
        # QTextDocument ignores CSS border-radius. Trim table corners natively
        # after text rendering; cell padding keeps text outside the corner area.
        painter=QPainter(self.viewport());painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setPen(Qt.PenStyle.NoPen);painter.setBrush(QColor('#131314'))
        layout=self.document().documentLayout()
        def trim(frame):
            if isinstance(frame,QTextTable) and frame.format().property(1048581)==True:
                rect=layout.frameBoundingRect(frame).translated(-self.horizontalScrollBar().value(),-self.verticalScrollBar().value())
                square=QPainterPath();square.addRect(rect)
                rounded=QPainterPath();rounded.addRoundedRect(rect,16,16)
                painter.drawPath(square.subtracted(rounded))
            for child in frame.childFrames():trim(child)
        try:
            trim(self.document().rootFrame())
        finally:
            painter.end()

    def loadResource(self, resource_type, url):
        if url.toString() in self.clipboard_images:
            return self.clipboard_images[url.toString()]
        # Chỉ logo và ảnh đính kèm trong bộ nhớ, không đọc file tùy ý.
        logo = ROOT / 'logo_chat_ai.png'
        if url.isLocalFile() and Path(url.toLocalFile()).resolve() == logo.resolve() and logo.is_file():
            return QPixmap(str(logo))
        # Nội dung model không tự tải ảnh từ Internet hoặc đọc file local.
        return None

    def mousePressEvent(self,event):
        cursor=self.cursorForPosition(event.position().toPoint()); fmt=cursor.charFormat()
        self.selected_image=fmt.toImageFormat().name() if fmt.isImageFormat() else None
        super().mousePressEvent(event)
    def keyPressEvent(self,event):
        if event.matches(QKeySequence.StandardKey.Copy) and self.selected_image in self.clipboard_images and not self.textCursor().hasSelection():
            QApplication.clipboard().setImage(self.clipboard_images[self.selected_image]);event.accept();return
        super().keyPressEvent(event)

    def open_link(self, url):
        if url.scheme()=='chatai-code':
            self.codeActionRequested.emit(url)
        elif url.scheme()=='chatai-media':
            try:
                import base64
                value=url.toString().partition(':')[2]
                value+='='*((4-len(value)%4)%4)
                path=Path(base64.urlsafe_b64decode(value.encode()).decode()).resolve(strict=True)
                if path.suffix.lower()!='.mp4' or not path.is_file() or not any(path.is_relative_to(Path(root).resolve()) for root in self.allowed_roots):return
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))
            except Exception:return
        elif url.scheme()=='chatai-file':
            try:
                import base64
                value=url.toString().partition(':')[2]
                value+='='*((4-len(value)%4)%4)
                path=Path(base64.urlsafe_b64decode(value.encode()).decode()).resolve(strict=True)
                if not path.is_file() or not any(path.is_relative_to(Path(root).resolve()) for root in self.allowed_roots):return
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))
            except Exception:return
        elif url.scheme()=='chatai-folder':
            try:
                import base64
                value=url.toString().partition(':')[2]
                value+='='*((4-len(value)%4)%4)
                path=Path(base64.urlsafe_b64decode(value.encode()).decode()).resolve(strict=True)
                if not any(path.is_relative_to(Path(root).resolve()) for root in self.allowed_roots):return
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(path.parent if path.is_file() else path)))
            except Exception:return
        elif url.scheme() in ('https', 'http'):
            QDesktopServices.openUrl(url)


class LocalOllamaClient:
    """Load Ollama only when a background operation first needs it."""
    def __init__(self,host,timeout=180):
        import threading
        self.host,self.timeout=host,timeout
        self._client=None;self._lock=threading.Lock()
    def __getattr__(self,name):
        with self._lock:
            if self._client is None:
                import ollama
                from assistant.vi_guard import install
                install()
                # The validated endpoint is HTTP loopback; no TLS or environment proxy is used.
                self._client=ollama.Client(host=self.host,timeout=self.timeout,verify=False,trust_env=False)
            client=self._client
        return getattr(client,name)


def prepare_context(progress=print):
    progress('Đang đọc cấu hình và mở lịch sử…')
    with measure('startup.config'):cfg = load_config()
    with measure('startup.history_database'):store = Store(ROOT / 'data/history.sqlite3')
    client = LocalOllamaClient(host=cfg['ollama_host'], timeout=180)
    progress('Đang nạp trạng thái module…')
    with measure('startup.module_state'):manager = ModuleManager(store, client, ROOT)
    cid = store.create(persist=False)
    rows = store.list(limit=30)
    return dict(cfg=cfg, store=store, client=client, manager=manager,
                rows=rows, cid=cid, state=store.load(cid), jobs=manager.jobs())


from assistant.support_ui import SupportMixin
from assistant.cloud import CLOUD_MODEL, NVIDIA_MODEL, REMOTE_MODELS, PROVIDER_NAMES, SWITCH_MESSAGE
from assistant.themes import style_sheet, recolor, chat_style, bubble_color
from assistant.settings_icon import SettingsButton
from assistant.profile_ui import ProfileMixin

from assistant.code_ui import CodeMixin
from assistant.admin_ui import AdminMixin,admin_session


class Window(QMainWindow, SupportMixin, ProfileMixin, CodeMixin, AdminMixin):
    model_probe_finished = Signal(object)
    account_enrichment_finished = Signal(object,object)
    login_restore_finished = Signal(object)
    support_badge_finished = Signal(object)
    history_sync_finished = Signal(object)
    session_verified = Signal(object)
    def __init__(self, context=None):
        super().__init__()
        context = context or prepare_context()
        self.cfg, self.store = context['cfg'], context['store']
        self.client, self.manager = context['client'], context['manager']
        self.initial_jobs = context['jobs']
        self.worker, self.callback, self.stream = None, None, ''
        self.chat_messages, self.html_cache = [], {}
        self.sent_prompt = None
        self.server_session=None; self.personal_memories=[]
        self.history_sync_loading=False
        self.history_sync_finished.connect(self.apply_history_sync)
        self.history_sync_timer=QTimer(self);self.history_sync_timer.setInterval(15000)
        self.history_sync_timer.timeout.connect(self.sync_history)
        self.history_sync_timer.start()
        self.login_restore_finished.connect(self.apply_restored_login)
        self.restore_login_loading=False
        self.support_badge_finished.connect(self.apply_support_badge)
        self.trial=GuestTrial(self.store)
        self.cfg.setdefault('chat_provider','nvidia')
        self.settings_last_tab=0; self.settings_navigation_guard=False
        import uuid
        self.device_id=uuid.uuid5(uuid.NAMESPACE_DNS,str(ROOT.resolve())).hex
        self.pending_image = None
        self.pending_documents=[]
        self.pending_media_paths=[]
        self.reply_active = False
        self.reply_frame = 0
        self.jobs_was_busy, self.jobs_loaded = False, False
        self.models = set()
        self.model_probe_running=False
        self.model_probe_finished.connect(self.models_probed)
        self.account_enrichment_finished.connect(self.apply_account_enrichment)
        self.session_verified.connect(self._on_session_verified)
        self.cid = context['cid']
        print('[4/5] Đang dựng cửa sổ chat…', flush=True)
        self.setWindowTitle('Chat AI · Desktop 2.6.6')
        logo = ROOT / 'logo_chat_ai.png'
        if logo.is_file():
            icon = QIcon(str(logo)); self.setWindowIcon(icon)
            QApplication.instance().setWindowIcon(icon)
        self.resize(1180, 790)
        self.setMinimumSize(820, 560)
        self.preview_theme=self.cfg.get('theme','dark');self.preview_font_size=self.cfg.get('font_size',13)
        self.setStyleSheet(style_sheet(self.preview_theme))
        splitter = QSplitter(); self.main_splitter = splitter
        sidebar = QWidget(); self.sidebar = sidebar; sidebar.setFixedWidth(260); sidebar.setObjectName('sidebar'); side = QVBoxLayout(sidebar)
        side.setContentsMargins(14, 22, 14, 18); side.setSpacing(12)
        brand_row = QHBoxLayout(); brand_row.addWidget(self.logo_label(28))
        title = QLabel('Chat AI'); title.setStyleSheet('font-size:18px;font-weight:600;padding:4px;')
        brand_row.addWidget(title, 1); self.sidebar_brand_title=title
        self.sidebar_toggle=QPushButton('☰');self.sidebar_toggle.setFixedSize(32,32);self.sidebar_toggle.setToolTip('Thu gọn / mở thanh bên');self.sidebar_toggle.clicked.connect(self.toggle_sidebar)
        brand_row.addWidget(self.sidebar_toggle);side.addLayout(brand_row)
        self.sidebar_stack=QStackedWidget(); side.addWidget(self.sidebar_stack,1)
        navigation=QWidget(); nav=QVBoxLayout(navigation);nav.setContentsMargins(0,0,0,0);nav.setSpacing(8)
        self.new_btn = self.button(nav, '＋  Cuộc trò chuyện mới', self.new_chat)
        self.history_search=QLineEdit();self.history_search.setPlaceholderText('Tìm kiếm hội thoại');nav.addWidget(self.history_search)
        self.history_search.textChanged.connect(self.filter_history)
        self.button(nav,'▧  Thư viện ảnh / video',lambda:self.tabs.setCurrentIndex(2))
        self.button(nav,'Trò chuyện',self.open_normal_chat)
        self.button(nav,'Chuyên gia',self.open_expert_chat)
        self.work_support_button=self.button(nav,'🗂️  Hỗ trợ Công việc',self.open_work_support)
        self.support_button=self.button(nav,'Quản lý hỗ trợ',self.open_support)
        self.support_button.setIcon(self.style().standardIcon(QStyle.StandardPixmap.SP_DialogHelpButton))
        self.button(nav,'ⓘ  Giới thiệu',self.open_about)
        self.admin_button=self.button(nav,'♙  Quản lý người dùng',self.open_user_admin);self.admin_button.hide()
        recent_header=QHBoxLayout();recent_header.setContentsMargins(0,0,0,0)
        recent_header.addWidget(QLabel('Gần đây'),1)
        self.delete_btn=QPushButton('Xóa');self.delete_btn.setFlat(True)
        self.delete_btn.setFixedWidth(40);self.delete_btn.setToolTip('Xóa cuộc trò chuyện đang chọn')
        self.delete_btn.setStyleSheet('QPushButton{border:0;background:transparent;padding:10px 2px 2px 2px;color:#8f8f8f;font-size:12px;font-weight:600;} QPushButton:hover{color:#f28b82;} QPushButton:disabled{color:#5f6368;}')
        self.delete_btn.clicked.connect(self.delete_chat);recent_header.addWidget(self.delete_btn)
        nav.addLayout(recent_header)
        self.history = QListWidget(); self.history.setObjectName('history'); nav.addWidget(self.history,1)
        self.history.itemClicked.connect(self.select_chat)
        self.history.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.undo_delete_box=QWidget();undo_row=QHBoxLayout(self.undo_delete_box);undo_row.setContentsMargins(0,0,0,0)
        undo_row.addWidget(QLabel('Đã xóa hội thoại'))
        undo_button=QPushButton('Hoàn tác');undo_button.clicked.connect(self.undo_delete_chat);undo_row.addWidget(undo_button)
        nav.addWidget(self.undo_delete_box);self.undo_delete_box.hide()
        self.undo_delete_timer=QTimer(self);self.undo_delete_timer.setSingleShot(True);self.undo_delete_timer.timeout.connect(self.expire_delete_undo)
        self.deleted_chat_backup=None
        self.sidebar_stack.addWidget(navigation)
        settings_nav=QWidget();self.settings_nav_layout=QVBoxLayout(settings_nav)
        self.settings_nav_layout.setContentsMargins(0,0,0,0);self.settings_nav_layout.setSpacing(8)
        self.sidebar_stack.addWidget(settings_nav)
        profile = QHBoxLayout(); profile.setSpacing(6)
        self.profile_button = QPushButton('Chưa đăng nhập')
        self.profile_button.setStyleSheet('text-align:left;font-size:13px;border-radius:12px;padding:8px;')
        self.profile_button.clicked.connect(self.account_menu)
        self.profile_avatar=QLabel();profile.insertWidget(0,self.profile_avatar)
        profile.addWidget(self.profile_button, 1)
        self.settings_button = SettingsButton(self)
        self.settings_button.setToolTip('Cài đặt Chat AI'); self.settings_button.clicked.connect(lambda:self.open_settings_section(None))
        profile.addWidget(self.settings_button); side.addLayout(profile)
        splitter.addWidget(sidebar)
        self.tabs = QTabWidget(); self.tabs.tabBar().hide(); splitter.addWidget(self.tabs)
        splitter.setChildrenCollapsible(False)
        splitter.setSizes([240, 920]); self.setCentralWidget(splitter)
        self.tabs.currentChanged.connect(self.balance_panels)
        self.tabs.currentChanged.connect(lambda index:self.sync_welcome())
        chat = QWidget(); outer = QVBoxLayout(chat)
        outer.setContentsMargins(24, 10, 24, 14)
        row = QHBoxLayout(); row.addWidget(self.logo_label(38)); brand = QLabel('Chat AI'); self.conversation_title=brand; brand.setMaximumWidth(320)
        brand.setStyleSheet('font-size:16px;font-weight:500;'); row.addWidget(brand); row.addStretch()
        self.quick_provider=QComboBox();self.quick_provider.setAccessibleName('Chế độ AI')
        self.quick_provider.setToolTip('Chọn AI trực tuyến hoặc AI trên máy.')
        self.quick_provider.addItem('AI trực tuyến','online');self.quick_provider.addItem('AI trên máy','local');row.addWidget(self.quick_provider)
        self.model = QComboBox(); self.model.setAccessibleName('Chọn AI')
        self.online_ai_names=list(REMOTE_MODELS)
        self.local_ai_names=list(dict.fromkeys([self.cfg['default_model'],self.cfg['code_model'],*CHAT_MODELS]))
        row.addWidget(self.model)
        self.model.setToolTip('Chọn nhà cung cấp AI trực tuyến hoặc mô hình Ollama trên máy.')
        configured_provider=PROVIDER_NAMES.get(self.cfg.get('chat_provider'),self.cfg['default_model'])
        if configured_provider not in REMOTE_MODELS and configured_provider not in self.local_ai_names:
            configured_provider=self.cfg['default_model']
        initial_kind='online' if configured_provider in REMOTE_MODELS else 'local'
        self.last_online_model=configured_provider if configured_provider in REMOTE_MODELS else next(iter(REMOTE_MODELS),NVIDIA_MODEL)
        initial_names=self.online_ai_names if initial_kind=='online' else self.local_ai_names
        self.model.addItems(initial_names)
        if configured_provider not in initial_names:configured_provider=initial_names[0]
        self.model.setCurrentText(configured_provider)
        self.deep_analysis_enabled=False
        self.quick_provider.setCurrentIndex(0 if initial_kind=='online' else 1)
        self.quick_provider.currentIndexChanged.connect(self.quick_ai_changed)
        self.chat_register=self.button(row,'Đăng ký',self.register_dialog)
        self.chat_login=self.button(row,'Đăng nhập',self.login_dialog)
        self.chat_mode = QComboBox(); self.chat_mode.addItems(['Chat nhanh', 'Dùng công cụ / Office', 'Tạo hình ảnh', 'Tạo video', 'Tìm kiếm mạng', 'Chuyên gia'])
        self.chat_mode.setCurrentIndex(0); self.chat_mode.setToolTip('Chọn chat, xử lý tài liệu, tạo ảnh hoặc tạo video từ ảnh AI.')
        self.mode_return_index=None
        self.chat_mode.hide()
        refresh = QPushButton('↻'); refresh.setToolTip('Kết nối lại Ollama')
        refresh.clicked.connect(self.refresh_models); row.addWidget(refresh)
        gpu = QPushButton('Cấu hình máy'); gpu.clicked.connect(lambda:self.open_settings_section('Cấu hình máy và AI')); row.addWidget(gpu)
        outer.addLayout(row)
        center_row = QHBoxLayout(); center_row.addStretch(1)
        center = QWidget(); center.setMaximumWidth(820); layout = QVBoxLayout(center)
        layout.setContentsMargins(6, 4, 6, 4); layout.setSpacing(12)
        center_row.addWidget(center, 8); center_row.addStretch(1); outer.addLayout(center_row, 1)
        self.chat_stack = QStackedWidget(); layout.addWidget(self.chat_stack, 1)
        welcome = QWidget(); wl = QVBoxLayout(welcome); wl.addStretch(2)
        self.welcome_robot=WelcomeRobot();wl.addWidget(self.welcome_robot)
        self.greeting=QLabel('');self.greeting.setMinimumHeight(48)
        self.greeting.setStyleSheet('font-size:30px;font-weight:600;color:#a8c7fa;');wl.addWidget(self.greeting)
        self.welcome_subtitle=QLabel('Tôi có thể giúp gì cho bạn?')
        self.welcome_subtitle.setWordWrap(True);self.welcome_subtitle.setMinimumHeight(52)
        self.welcome_subtitle.setStyleSheet('font-size:18px;color:#c4c7c5;margin-bottom:16px;');wl.addWidget(self.welcome_subtitle)
        self.welcome_opacity=QGraphicsOpacityEffect(self.welcome_subtitle);self.welcome_opacity.setOpacity(0)
        self.welcome_subtitle.setGraphicsEffect(self.welcome_opacity)
        self.welcome_fade=QPropertyAnimation(self.welcome_opacity,b'opacity',self)
        self.welcome_fade.setDuration(400);self.welcome_fade.setStartValue(0.0);self.welcome_fade.setEndValue(1.0)
        self.welcome_fade.setEasingCurve(QEasingCurve.Type.OutCubic)
        self.greeting_timer=QTimer(self);self.greeting_timer.setInterval(60);self.greeting_timer.timeout.connect(self.advance_greeting)
        self.welcome_conversation=None
        suggestions = QHBoxLayout()
        for text, prompt in [('▦  Xử lý Excel', 'Liệt kê các file Excel trong thư mục tài liệu.'),
                             ('▤  Soạn tài liệu', 'Giúp tôi soạn một báo cáo công việc bằng tiếng Việt.'),
                             ('✦  Tạo hình ảnh', 'Tạo ảnh một ngôi nhà nhỏ bên hồ vào buổi sáng.')]:
            btn = QPushButton(text); btn.clicked.connect(lambda checked=False, p=prompt: self.fill_prompt(p))
            suggestions.addWidget(btn)
        wl.addLayout(suggestions); wl.addStretch(3); self.chat_stack.addWidget(welcome)
        self.code_actions={};self.copied_codes={};self.saved_code_files=[]
        self.view = ChatView();self.view.codeActionRequested.connect(self.handle_code_action); self.view.setObjectName('chatView'); self.chat_stack.addWidget(self.view)
        self.status = FadingStatusLabel('Sẵn sàng'); self.status.setStyleSheet('color:#9aa0a6;font-size:12px;'); self.status.setWordWrap(True)
        status_row=QHBoxLayout(); self.reply_logo=self.logo_label(22);self.reply_logo.hide();status_row.addWidget(self.reply_logo)
        status_row.addWidget(self.status,1)
        self.cancel_countdown_button=QPushButton('Hủy yêu cầu');self.cancel_countdown_button.clicked.connect(self.send_or_stop);self.cancel_countdown_button.hide();status_row.addWidget(self.cancel_countdown_button)
        self.windows_stop_button=QPushButton('Dừng app AI');self.windows_stop_button.clicked.connect(self.stop_windows_apps)
        self.windows_stop_button.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.windows_stop_button.setToolTip('Chặn các bước điều khiển ứng dụng tiếp theo; không tắt app hay bỏ qua lưu tài liệu.')
        self.windows_stop_button.setVisible(self.cfg.get('windows_apps_enabled',False));status_row.addWidget(self.windows_stop_button)
        self.reply_dots=QPushButton('● · ·'); self.reply_dots.setFixedWidth(86); self.reply_dots.setVisible(False)
        self.reply_dots.setToolTip('Đến phần AI đang trả lời'); self.reply_dots.setAccessibleName('Cuộn xuống câu trả lời mới nhất')
        self.reply_dots.setStyleSheet('color:#a8c7fa;background:#282a2c;font-size:20px;padding:4px;border-radius:14px;')
        self.reply_dots.clicked.connect(self.jump_to_reply); status_row.addWidget(self.reply_dots); status_row.addStretch(1); layout.addLayout(status_row)
        self.answer_actions=QWidget(); self.answer_actions_layout=QHBoxLayout(self.answer_actions)
        self.answer_actions_layout.setContentsMargins(0,0,0,0)
        self.answer_actions.hide(); layout.addWidget(self.answer_actions)
        self.expert_toolbar=QWidget();er=QHBoxLayout(self.expert_toolbar);er.setContentsMargins(0,0,0,0)
        er.addWidget(QLabel('Chuyên gia chính:'))
        self.expert_choice=QComboBox();self.expert_choice.addItem('Tự động phối hợp',None)
        from assistant.orchestrator import load_experts
        for name,entry in load_experts()['experts'].items():
            if entry.get('kind')=='llm' and entry.get('role')!='planner':
                labels={'chat':'Chat chung','code':'Lập trình','vision':'Đọc ảnh'}
                self.expert_choice.addItem(labels.get(entry.get('role'),name),name)
        er.addWidget(self.expert_choice,1)
        self.expert_details=self.button(er,'Chi tiết xử lý',self.show_expert_details)
        self.expert_toolbar.hide();layout.insertWidget(0,self.expert_toolbar)
        self.expert_reselect=QPushButton('Hiểu sai ý? Chọn lại chuyên gia')
        self.expert_reselect.clicked.connect(self.reselect_expert);self.expert_reselect.hide();layout.addWidget(self.expert_reselect)
        self.reply_timer=QTimer(self); self.reply_timer.setInterval(350); self.reply_timer.timeout.connect(self.animate_reply)

        composer = QWidget(); composer.setObjectName('composer'); cl = QVBoxLayout(composer)
        cl.setContentsMargins(10, 6, 10, 9); cl.setSpacing(3)
        _shadow = QGraphicsDropShadowEffect(composer)
        _shadow.setBlurRadius(20); _shadow.setOffset(0, 4); _shadow.setColor(QColor(0, 0, 0, 60))
        composer.setGraphicsEffect(_shadow)
        self.input = PromptEdit(); self.input.setObjectName('prompt')
        self.input.setPlaceholderText('Hỏi Chat AI…'); self.input.setMinimumHeight(60); self.input.setMaximumHeight(130)
        self.input.setAccessibleName('Tin nhắn. Enter gửi; Shift Enter xuống dòng.')
        self.input.sendRequested.connect(self.send); self.input.imagePasted.connect(self.attach_image);self.input.filesPasted.connect(self.attach_files)
        self.attachment_box=QWidget(); attachment=QHBoxLayout(self.attachment_box)
        self.attachment_preview=ImagePreview(); attachment.addWidget(self.attachment_preview)
        self.button(attachment,'Copy ảnh',self.copy_attachment)
        self.button(attachment,'Bỏ ảnh',self.remove_attachment); attachment.addStretch()
        self.attachment_box.hide(); cl.addWidget(self.attachment_box)
        self.document_box=QWidget();doc_row=QHBoxLayout(self.document_box)
        self.document_label=QLabel();self.document_label.setWordWrap(True);doc_row.addWidget(self.document_label,1)
        self.button(doc_row,'Bỏ tài liệu',self.remove_documents);self.document_box.hide();cl.addWidget(self.document_box)
        cl.addWidget(self.input)
        actions = QHBoxLayout(); cl.addLayout(actions)
        hint = QLabel('Enter để gửi · Shift + Enter xuống dòng'); hint.setStyleSheet('color:#9aa0a6;font-size:11px;background:transparent;')
        self.image_btn = self.button(actions, 'Ảnh', lambda: self.toggle_temporary_mode(2))
        self.video_btn = self.button(actions, 'Video', lambda: self.toggle_temporary_mode(3))
        self.web_btn=self.button(actions,'Tìm web',self.toggle_web); self.web_btn.setCheckable(True)
        self.web_btn.setToolTip('Bật để tra Bing và trả lời kèm nguồn. Chỉ gửi câu hỏi, không gửi ảnh.')
        self.image_btn.setToolTip('Tạo hình ảnh');self.video_btn.setToolTip('Tạo video từ ảnh')
        self.image_btn.setCheckable(True); self.video_btn.setCheckable(True)
        self.deep_btn=self.button(actions,'Phân tích sâu',self.toggle_deep_analysis);self.deep_btn.setCheckable(True)
        self.deep_btn.setToolTip('Bật lượt phân tích và rà soát câu trả lời kỹ hơn. Có thể mất thêm thời gian.')
        actions.addStretch()
        self.resume_btn = self.button(actions, 'Tiếp tục / Duyệt', self.resume)
        self.button(actions, 'Lịch sử', self.full_history)
        self.button(actions, 'Nhật ký', self.show_audit)
        self.send_btn = self.button(actions, '↑', self.send_or_stop); self.send_btn.setObjectName('send')
        self.send_btn.setToolTip('Gửi tin nhắn (Enter)'); self.send_btn.setAccessibleName('Gửi tin nhắn')
        layout.addWidget(composer)
        layout.addWidget(hint)
        self.tabs.addTab(chat, 'Trò chuyện')
        modules = QWidget(); ml = QVBoxLayout(modules)
        self.button(ml, '← Quay lại chat', lambda: self.tabs.setCurrentIndex(0))
        ml.addWidget(QLabel('Thư viện ảnh/video cơ bản tự tải nền khi thiếu. Model AI dung lượng lớn chỉ tải khi bạn chọn. Có thể chat trong lúc tải.'))
        for key, spec in MODULES.items():
            if key=='web':
                note=QLabel('Tìm kiếm Bing tích hợp sẵn. Bấm Tìm kiếm mạng trong khung chat; không cần tải hoặc bật module.');note.setWordWrap(True);ml.addWidget(note);continue
            block = QHBoxLayout(); label = QLabel(spec['label'] + '\n' + spec['note']); label.setWordWrap(True)
            block.addWidget(label, 1)
            self.button(block, 'Bật / Tải', lambda checked=False, k=key: self.download(k))
            self.button(block, 'Tắt', lambda checked=False, k=key: self.disable(k))
            ml.addLayout(block)
            if key=='media':
                variant_row=QHBoxLayout();variant_row.addWidget(QLabel('Model tạo ảnh'))
                self.media_variant_combo=QComboBox()
                for variant,variant_spec in MEDIA_VARIANTS.items():
                    self.media_variant_combo.addItem(variant_spec['label'],variant)
                self.media_variant_combo.setCurrentIndex(self.media_variant_combo.findData(self.manager.media_variant()))
                self.media_variant_note=QLabel(MEDIA_VARIANTS[self.manager.media_variant()]['quality_note'])
                self.media_variant_note.setWordWrap(True)
                def set_media_variant(index):
                    variant=self.media_variant_combo.itemData(index)
                    if variant not in MEDIA_VARIANTS:return
                    try:self.manager.set_media_variant(variant)
                    except Exception as error:
                        QMessageBox.warning(self,'Model tạo ảnh',str(error));return
                    self.media_variant_note.setText(MEDIA_VARIANTS[variant]['quality_note'])
                self.media_variant_combo.currentIndexChanged.connect(set_media_variant)
                variant_row.addWidget(self.media_variant_combo,1);ml.addLayout(variant_row);ml.addWidget(self.media_variant_note)
        modelrow = QHBoxLayout(); catalog = QComboBox()
        for name in [*CHAT_MODELS,*sorted(ALLOWED_MODELS-set(CHAT_MODELS))]:
            label = CHAT_MODELS.get(name, {}).get('label', 'Embedding tài liệu')
            catalog.addItem(f'{name} · {label}', name)
        modelrow.addWidget(catalog, 1)
        self.button(modelrow, 'Tải AI đã chọn', lambda: self.download(catalog.currentData()))
        self.button(modelrow, 'Gỡ AI đã chọn', lambda: self.remove_ai(catalog.currentData()))
        ml.addLayout(modelrow)
        self.download_status=QLabel('Chưa có tác vụ tải.');self.download_status.setWordWrap(True);ml.addWidget(self.download_status)
        self.progress = QProgressBar(); ml.addWidget(self.progress)
        controls=QHBoxLayout()
        self.pause_download_btn=self.button(controls,'Tạm dừng tải',lambda:self.control_download('paused'))
        self.continue_download_btn=self.button(controls,'Tiếp tục tải',lambda:self.control_download('resume'))
        self.cancel_download_btn=self.button(controls,'Hủy tải',lambda:self.control_download('cancelled'))
        for button in (self.pause_download_btn,self.continue_download_btn,self.cancel_download_btn):button.setEnabled(False)
        ml.addLayout(controls)
        self.jobs = QPlainTextEdit(); self.jobs.setReadOnly(True); ml.addWidget(self.jobs, 1)
        self.add_scroll_page(modules, 'Tải mô hình')
        media = QWidget(); gl = QVBoxLayout(media)
        self.button(gl, '← Quay lại chat', lambda: self.tabs.setCurrentIndex(0))
        gl.addWidget(QLabel('Ảnh PNG và video MP4 được tạo trong thư mục whitelist.'))
        self.media_list = QListWidget(); gl.addWidget(self.media_list)
        self.button(gl, 'Làm mới danh sách', self.refresh_media)
        self.button(gl, 'Mở ảnh / video bằng ứng dụng Windows', self.open_media)
        self.add_scroll_page(media, 'Ảnh và Video')
        settings = QWidget(); self.settings_page= settings; settings.setMaximumWidth(600); sl = QVBoxLayout(settings);sl.setAlignment(Qt.AlignmentFlag.AlignTop);sl.setSpacing(18)
        self.settings_heading=QLabel('Cài đặt');self.settings_heading.setStyleSheet('font-size:20px;font-weight:600;margin-bottom:8px;');sl.addWidget(self.settings_heading)
        self.build_settings(sl)
        self.add_scroll_page(settings, 'Office / Cài đặt')
        self.build_settings_navigation()
        self.paint_timer = QTimer(self); self.paint_timer.setSingleShot(True)
        self.paint_timer.timeout.connect(self.draw)
        self.timer = QTimer(self); self.timer.timeout.connect(self.poll); self.timer.start(1500)
        self.apply_theme(self.cfg.get('theme','dark'));self.render(context)
        self.tabs.currentChanged.connect(self.check_settings_navigation)
        self.support_timer=QTimer(self);self.support_timer.timeout.connect(self.poll_support_badge);self.support_timer.start(30000)
        self.model.currentTextChanged.connect(self.model_changed)
        self.chat_mode.currentIndexChanged.connect(self.mode_changed)
        self.mode_changed(self.chat_mode.currentIndex())
        QTimer.singleShot(200, self.refresh_startup_models)
        QTimer.singleShot(100,self.restore_login)

    @staticmethod
    def logo_label(size):
        label = QLabel(); label.setFixedSize(size, size)
        label.setStyleSheet('background:transparent;border:0;padding:0;')
        label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        pixmap = QPixmap(str(ROOT / 'logo_chat_ai.png'))
        if not pixmap.isNull():
            label.setPixmap(pixmap.scaled(size, size, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        else: label.setText('✦')
        return label

    def add_scroll_page(self, widget, title):
        scroll = QScrollArea(); scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        wrapper=QWidget();row=QHBoxLayout(wrapper);row.setContentsMargins(24,24,24,24)
        widget.setMaximumWidth(600 if widget is getattr(self,'settings_page',None) else 860)
        if widget.layout():widget.layout().setAlignment(Qt.AlignmentFlag.AlignTop)
        row.addStretch(1);row.addWidget(widget,1,Qt.AlignmentFlag.AlignTop);row.addStretch(1)
        scroll.setWidget(wrapper);self.tabs.addTab(scroll,title)
        if hasattr(self,'paint_timer'):self.apply_theme(self.preview_theme,self.preview_font_size)

    @staticmethod
    def button(layout, text, callback):
        button = QPushButton(text); button.clicked.connect(callback); layout.addWidget(button); return button

    def compact_app_activity(self,text):
        labels={'windows_list_apps':'Đang tìm ứng dụng','windows_open':'Đang mở ứng dụng','windows_inspect':'Đang đọc giao diện','windows_action':'Đang thao tác ứng dụng','browser_search':'Đang tìm trên Chrome','browser_run':'Đang thao tác Chrome','cad3d_create_open':'Đang tạo mô hình 3D và mở AutoCAD','cad_create_open':'Đang tạo bản vẽ và mở AutoCAD','word_create_open':'Đang tạo tài liệu và mở Word','pdf_source_open':'Đang tải và mở PDF','pdf_read':'Đang đọc PDF'}
        text=labels.get(text.split(': ')[-1],text)
        if not hasattr(self,'automation_panel'):
            from assistant.automation_panel import AutomationPanel
            self.automation_panel=AutomationPanel()
            self.automation_panel.pauseRequested.connect(self.pause_app_activity)
            self.automation_panel.chatRequested.connect(self.reveal_app_chat)
            self.automation_panel.stopRequested.connect(self.end_app_activity)
        if not getattr(self,'app_compact_active',False):
            self.app_compact_active=True
            self.app_was_maximized=self.isMaximized()
            self.automation_panel.begin(text)
            if self.cfg.get('windows_apps_compact',True):self.showMinimized()
            self.automation_panel.show()
        else:
            self.automation_panel.message=text;self.automation_panel.update_label()
            self.automation_panel.show()

    def pause_app_activity(self,paused):
        from assistant.windows_apps import pause_automation
        pause_automation(paused)

    def reveal_app_chat(self):
        if getattr(self,'app_was_maximized',False):self.showMaximized()
        else:self.showNormal()
        self.raise_();self.activateWindow()

    def end_app_activity(self):
        self.stop_windows_apps()
        if self.worker and self.worker.cancellable:self.worker.stop_requested.set()
        self.automation_panel.pause.setEnabled(False)
        self.automation_panel.message='Đang kết thúc sau thao tác hiện tại'
        self.automation_panel.paused=False;self.automation_panel.update_label()

    def finish_app_activity(self):
        if not getattr(self,'app_compact_active',False):return
        from assistant.windows_apps import pause_automation
        pause_automation(False)
        self.automation_panel.end();self.app_compact_active=False
        self.reveal_app_chat()

    def busy(self):
        return self.worker is not None

    def send_or_stop(self):
        if self.worker and self.worker.cancellable:
            from assistant.windows_apps import stop_automation
            try:self.store.audit(self.cid,'automation_stop_requested',{'source':'send_or_stop'})
            except Exception:pass
            stop_automation()
            self.worker.stop_requested.set()
            self.send_btn.setEnabled(False)
            self.status.setText('Đang dừng… chờ thao tác hiện tại kết thúc an toàn.')
        elif not self.busy():self.send()

    def sync_send_button(self, enabled=True):
        stopping=bool(self.worker and self.worker.cancellable)
        self.send_btn.setText('■' if stopping else '↑')
        self.send_btn.setToolTip('Dừng phản hồi' if stopping else 'Gửi tin nhắn (Enter)')
        self.send_btn.setAccessibleName('Dừng phản hồi' if stopping else 'Gửi tin nhắn')
        self.send_btn.setEnabled(not self.worker.stop_requested.is_set() if stopping else enabled and not self.busy())

    def work(self, fn, callback=None, cancellable=False, app_countdown=False):
        if self.busy():
            return
        worker = Worker(fn); self.worker, self.callback = worker, callback
        worker.app_countdown=app_countdown
        worker.cancellable=cancellable;worker.chat_cid=self.cid if cancellable else None
        worker.event.connect(self.on_event)
        worker.finished.connect(self.finished)
        self.new_btn.setEnabled(False); self.delete_btn.setEnabled(False); self.history.setEnabled(False); self.model.setEnabled(False); self.chat_mode.setEnabled(False)
        self.image_btn.setEnabled(False); self.video_btn.setEnabled(False); self.web_btn.setEnabled(False);self.deep_btn.setEnabled(False)
        self.sync_send_button(False); self.resume_btn.setEnabled(False)
        self.apply_settings_button.setEnabled(False)
        self.sync_welcome()
        worker.start()

    def finished(self):
        self.cancel_countdown_button.hide()
        self.finish_app_activity()
        worker, callback = self.worker, self.callback
        if worker.failure and self.sent_prompt and not self.input.toPlainText().strip():
            self.input.setPlainText(self.sent_prompt);self.input.setFocus();self.sent_prompt=None
        if worker.cancelled:
            state=self.store.load(worker.chat_cid)
            visible=[m for m in state['messages'] if m['role'] in ('user','assistant') and m.get('content')]
            if worker.partial and len(visible)<=worker.snapshot_count:
                state['messages'].append({'role':'assistant','content':worker.partial})
            pending=state.get('pending')
            if not (pending and pending.get('decision_started')):
                for call in state.get('queue',[]):
                    state['messages'].append({'role':'tool','tool_name':call['function']['name'],
                        'content':json.dumps({'ok':False,'cancelled':True,'error':'Người dùng đã dừng; công cụ chưa được thực hiện.'},ensure_ascii=False)})
                state.update(queue=[],pending=None)
            state.update(running=False,followups=[],stopped=True)
            marker='Đã dừng phản hồi.'
            if state['messages'] and state['messages'][-1]['role']=='assistant':state['messages'][-1]['content']+='\n\n'+marker
            else:state['messages'].append({'role':'assistant','content':marker})
            self.store.save(worker.chat_cid,state)
            self.store.audit(worker.chat_cid,'chat_stopped',{'partial_preserved':bool(worker.partial)})
        self.reply_active=False; self.reply_timer.stop(); self.reply_dots.hide();self.reply_logo.hide()
        self.worker, self.callback, self.stream = None, None, ''
        self.paint_timer.stop()
        self.image_btn.setEnabled(True); self.video_btn.setEnabled(True); self.web_btn.setEnabled(True);self.deep_btn.setEnabled(True)
        self.new_btn.setEnabled(True); self.delete_btn.setEnabled(True); self.chat_mode.setEnabled(True); self.history.setEnabled(True); self.model.setEnabled(True)
        if getattr(worker,'performance_login',False):
            with measure('login.finish_worker_render'):self.render()
        else:self.render()
        worker.deleteLater()
        if worker.cancelled:
            self.status.setText('Đã dừng phản hồi.');self.sent_prompt=None
        elif worker.failure:
            self.status.setText(worker.failure)
            QMessageBox.warning(self, 'Chat AI', worker.failure)
        elif callback:
            if not getattr(self,'exit_when_idle',False):callback(worker.result)
        if not worker.failure:self.sync_history()
        if getattr(self,'exit_when_idle',False):self.close()

    def on_event(self, event):
        if event['type']=='app_countdown':
            seconds=event['seconds'];self.cancel_countdown_button.setVisible(seconds>0)
            if seconds:self.status.setText(f'AI sẽ thực hiện yêu cầu sau {seconds} giây…')
            return
        if event['type']=='app_activity':
            self.compact_app_activity(event['text']);return
        if event['type']=='status' and getattr(self,'app_compact_active',False):
            self.automation_panel.message=event['text'];self.automation_panel.update_label()
        if event['type']=='auth_failed':
            self.server_session=None;self.personal_memories=[]
            self.cid=self.store.create(persist=False)
            if hasattr(self,'api_key_group'):self.api_key_group.hide()
            self.account_status.setText('Cần xác thực lại tài khoản.');return
        if event['type']=='sent':
            if event.get('cid')!=self.cid:return
            sent=event.get('prompt','')
            if self.input.toPlainText().strip()==sent:self.input.clear()
            self.remove_attachment();self.remove_documents();self.sent_prompt=None
            self.render()
            return
        if event['type'] == 'token':
            if not self.stream: self._flash_chat_view()
            self.stream += event['text']
            if not self.paint_timer.isActive(): self.paint_timer.start(250)
        elif event['type'] == 'snapshot':
            if self.sent_prompt and any(m['role'] == 'user' and m['content'] == self.sent_prompt for m in event['messages'][-1:]):
                if self.input.toPlainText().strip() == self.sent_prompt: self.input.clear()
                self.remove_attachment();self.remove_documents()
                self.sent_prompt = None
            self.chat_messages = event['messages']; self.stream = ''
            if self.sent_prompt and self.chat_messages and self.chat_messages[-1]['role']=='user' and self.chat_messages[-1]['content']==self.sent_prompt:self.remove_attachment()
            if not self.paint_timer.isActive(): self.paint_timer.start(250)
        elif event['type'] == 'status':
            if event['text'].startswith('Đang trả lời'): self.stream = ''
            self.status.setText(event['text'])

    def fill_prompt(self, text):
        if text.startswith(('Liệt kê các file', 'Tạo ảnh')):
            if not CHAT_MODELS.get(self.model.currentText(), {}).get('tools',False):
                fallback = self.cfg['default_model'] if CHAT_MODELS[self.cfg['default_model']]['tools'] else 'qwen2.5:3b'
                self.select_ai(fallback)
            self.chat_mode.setCurrentIndex(1)
        self.tabs.setCurrentIndex(0); self.input.setPlainText(text); self.input.setFocus()
        self.input.moveCursor(QTextCursor.MoveOperation.End)

    def attach_files(self, paths):
        for raw in paths[:5]:
            path=Path(raw)
            if not path.is_file():continue
            max_mb=20 if path.suffix.lower()=='.dxf' else 10
            if path.stat().st_size>max_mb*1024*1024:
                QMessageBox.warning(self,'Đính kèm',f'Tệp tối đa {max_mb} MB: '+path.name);continue
            if path.suffix.lower() in ('.png','.jpg','.jpeg','.webp','.bmp'):
                self.attach_image(QImage(str(path)))
                if len(self.pending_media_paths)<8:
                    try:
                        source=path.resolve(strict=True)
                        roots=[Path(root).resolve() for root in self.cfg.get('roots',[])]
                        if any(source.is_relative_to(root) for root in roots):
                            media_path=source
                        else:
                            # Import the user-selected image into the approved workspace so media tools never read outside the whitelist.
                            import shutil,uuid
                            media_root=roots[0]/'attachments'
                            media_root.mkdir(parents=True,exist_ok=True)
                            media_path=media_root/(uuid.uuid4().hex[:8]+'_'+path.name)
                            shutil.copyfile(source,media_path)
                        if str(media_path) not in self.pending_media_paths:
                            self.pending_media_paths.append(str(media_path))
                    except Exception as error:
                        QMessageBox.warning(self,'Ảnh nguồn video','Ảnh vẫn được đính kèm để AI xem, nhưng không thể chuẩn bị làm nguồn video: '+str(error))
                continue
            if len(self.pending_documents)<5 and str(path) not in self.pending_documents:self.pending_documents.append(str(path))
        self.document_label.setText('📎 '+', '.join(Path(p).name for p in self.pending_documents))
        self.document_box.setVisible(bool(self.pending_documents))

    def remove_documents(self):
        self.pending_documents=[];self.pending_media_paths=[];self.document_label.clear();self.document_box.hide()

    def read_attachments(self, paths, progress=None, online=False):
        items=[]
        for raw in paths:
            p=Path(raw)
            if not p.is_file() or p.stat().st_size>(20 if p.suffix.lower()=='.dxf' else 10)*1024*1024:raise ValueError('Tệp đính kèm không còn hợp lệ: '+p.name)
            try:
                ext=p.suffix.lower()
                if ext in ('.pdf','.docx','.txt','.md'):
                    from assistant.documents import read_local
                    from assistant.online_documents import online_pdf_reader
                    reader=online_pdf_reader(self.cfg,getattr(self,'server_session',None),on_status=progress) if ext=='.pdf' and getattr(self,'server_session',None) else None
                    items.append(read_local(p,pdf_ocr=reader,foxit_ocr=False,progress=progress))
                    continue
                from assistant.code_files import CODE_SUFFIXES
                if ext in CODE_SUFFIXES:content=p.read_text(encoding='utf-8-sig',errors='strict')
                elif ext=='.docx':
                    from docx import Document
                    doc=Document(p);content='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
                elif ext=='.pptx':
                    from pptx import Presentation
                    content='\n'.join(shape.text for slide in Presentation(p).slides for shape in slide.shapes if shape.has_text_frame)
                elif ext=='.xlsx':
                    from assistant.geoslope_inspect import inspect_workbook
                    summary=inspect_workbook(p)
                    if summary['material_tables'] or summary['section_tables']:
                        import json
                        summary['section_tables']=[{'sheet':t['sheet'],'sections':[
                            {'section':s['section'],'row':s['row'],'layers':s['layers']} for s in t['sections']],
                            'truncated':t['truncated']} for t in summary['section_tables']]
                        content='Bảng tổng hợp/chỉ tiêu địa kỹ thuật (địa chỉ ô nguồn; không tính lại công thức):\n'+json.dumps(summary,ensure_ascii=False,default=str)
                    else:
                        import pandas as pd
                        book=pd.ExcelFile(p);content='\n'.join('Sheet '+name+'\n'+book.parse(name,nrows=40).to_csv(index=False) for name in book.sheet_names[:5]);book.close()
                elif ext=='.gsz':
                    from assistant.geoslope_inspect import inspect_gsz
                    import json
                    summary=inspect_gsz(p)
                    content='GeoStudio GSZ đã lưu; chưa chạy Solve mới. Dùng geoslope_inspect để đối chiếu đầy đủ.\n'+json.dumps(
                        {k:v for k,v in summary.items() if k not in {'geometries','stability_xml','coordinates_xml'}},ensure_ascii=False)
                elif ext=='.dxf':
                    # Parse ASCII DXF as inert group-code/value pairs; never execute CAD data.
                    raw=p.read_bytes()
                    if raw.startswith(b'AutoCAD Binary DXF'):
                        content='DXF nhị phân chưa được hỗ trợ; hãy xuất bản vẽ sang DXF ASCII rồi đính kèm lại.'
                    else:
                        lines=raw.decode('utf-8-sig',errors='replace').splitlines()
                        pairs=[]
                        for i in range(0,len(lines)-1,2):
                            try:code=int(lines[i].strip())
                            except ValueError:continue
                            pairs.append((code,lines[i+1].strip()))
                        entities=[];in_entities=False;entity=None
                        for code,value in pairs:
                            if code==2 and value.upper()=='ENTITIES':in_entities=True;continue
                            if code==0 and value.upper()=='ENDSEC':
                                if entity:entities.append(entity);entity=None
                                in_entities=False;continue
                            if not in_entities:continue
                            if code==0:
                                if entity:entities.append(entity)
                                entity={'type':value.upper(),'values':[]}
                            elif entity:entity['values'].append((code,value))
                        if entity:entities.append(entity)
                        from collections import Counter
                        counts=Counter(e['type'] for e in entities)
                        layers=Counter(next((v for c,v in e['values'] if c==8),'0') for e in entities)
                        parts=['DXF ASCII (đọc an toàn, không chạy macro).',f'Tổng đối tượng: {len(entities)}.',
                               'Loại đối tượng: '+', '.join(f'{k}: {v}' for k,v in counts.most_common(40)),
                               'Layer: '+', '.join(f'{k}: {v}' for k,v in layers.most_common(40))]
                        labels=[]
                        for e in entities:
                            if e['type'] in ('TEXT','MTEXT','ATTRIB','ATTDEF','DIMENSION'):
                                vals=[v for c,v in e['values'] if c in (1,3)]
                                if vals:labels.append(' '.join(vals))
                        if labels:parts.append('Nhãn/chú thích:\n'+'\n'.join(labels[:1000]))
                        # Include coordinate-bearing group codes to help the model interpret geometry.
                        geometry=[]
                        for e in entities[:2000]:
                            coords=[f'{c}={v}' for c,v in e['values'] if c in (10,20,30,11,21,31,40,41,42)]
                            if coords:geometry.append(e['type']+' ['+', '.join(coords[:12])+']')
                        if geometry:parts.append('Tọa độ/kích thước (mẫu):\n'+'\n'.join(geometry))
                        content='\n'.join(parts)
                else:
                    # Unknown formats are accepted as attachments, never executed.
                    with p.open('rb') as stream:sample=stream.read(8192)
                    try:
                        sample.decode('utf-8-sig')
                        textual=b'\x00' not in sample and (not sample or sum(b<32 and b not in (9,10,13) for b in sample)/len(sample)<.02)
                    except UnicodeDecodeError:textual=False
                    if textual:
                        with p.open('r',encoding='utf-8-sig',errors='replace') as stream:content=stream.read(12000)
                    else:
                        content='Tệp nhị phân đã đính kèm; chưa có bộ đọc nội dung cho định dạng này. Không mở hoặc chạy tệp. Chỉ biết tên: '+p.name+'; dung lượng: '+str(p.stat().st_size)+' byte.' 
            except ImportError:
                content='Đã nhận tệp '+p.name+', nhưng chưa đọc nội dung vì thiếu thư viện Office/RAG tương ứng. Không suy đoán nội dung.'
            limit=12000 if ext in CODE_SUFFIXES else 16000
            items.append({'file':p.name,'text':content[:limit],'truncated':len(content)>limit})
        from assistant.document_memory import DocumentMemory
        owner=self.server_session['username'] if self.server_session else GUEST_OWNER
        DocumentMemory(self.store).remember(owner,items)
        return items

    def attach_image(self,image):
        if self.busy():
            QMessageBox.information(self,'Ảnh','Đợi AI trả lời xong trước khi đính kèm ảnh mới.'); return
        if image.isNull(): return
        image=image.scaled(1600,1600,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation)
        data=QByteArray(); buffer=QBuffer(data); buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        if not image.save(buffer,'JPEG',85):buffer.close(); return
        buffer.close()
        if data.size()>950000:
            QMessageBox.warning(self,'Ảnh','Ảnh quá lớn sau khi nén. Chọn ảnh nhỏ hơn.'); return
        self.pending_image=bytes(data.toBase64()).decode('ascii')
        self.attachment_preview.image=QImage(image)
        self.attachment_preview.setPixmap(QPixmap.fromImage(image).scaled(120,80,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation))
        self.attachment_box.show()

    def copy_attachment(self):
        if self.attachment_preview.image is not None:QApplication.clipboard().setImage(self.attachment_preview.image)

    def remove_attachment(self):
        self.pending_image=None; self.attachment_preview.image=None; self.attachment_preview.clear(); self.attachment_box.hide()

    def sync_welcome(self):
        if not hasattr(self,'greeting_timer'):return
        active=self.tabs.currentIndex()==0 and self.chat_stack.currentIndex()==0 and not self.busy()
        if not active:
            self.welcome_robot.timer.stop();self.greeting_timer.stop();self.welcome_fade.stop()
            if self.welcome_conversation==self.cid:
                self.greeting.setText('Xin chào bạn!');self.welcome_opacity.setOpacity(1)
            return
        self.welcome_robot.timer.start()
        if self.welcome_conversation!=self.cid:
            self.welcome_conversation=self.cid;self.greeting_position=0
            self.greeting.setText('');self.welcome_opacity.setOpacity(0);self.greeting_timer.start()

    def advance_greeting(self):
        text='Xin chào bạn!';self.greeting_position+=1
        self.greeting.setText(text[:self.greeting_position])
        if self.greeting_position>=len(text):
            self.greeting_timer.stop();self.welcome_fade.start()

    def animate_reply(self):
        self.reply_frame=(self.reply_frame+1)%3
        dots=['● · ·','· ● ·','· · ●'][self.reply_frame]
        # Update text without triggering FadingStatusLabel animation
        QPushButton.setText(self.reply_dots, dots)

    def _smooth_scroll_to_bottom(self):
        bar = self.view.verticalScrollBar()
        anim = QPropertyAnimation(bar, b'value', self.view)
        anim.setDuration(280); anim.setEasingCurve(QEasingCurve.Type.OutCubic)
        anim.setStartValue(bar.value()); anim.setEndValue(bar.maximum())
        anim.start(QPropertyAnimation.DeletionPolicy.DeleteWhenStopped)

    def _flash_chat_view(self):
        if not hasattr(self, '_view_opacity'):
            self._view_opacity = QGraphicsOpacityEffect(self.view)
            self.view.setGraphicsEffect(self._view_opacity)
        eff = self._view_opacity
        anim = QPropertyAnimation(eff, b'opacity', self.view)
        anim.setDuration(300); anim.setEasingCurve(QEasingCurve.Type.OutCubic)
        anim.setKeyValueAt(0, 0.7); anim.setKeyValueAt(0.5, 1.0); anim.setKeyValueAt(1, 1.0)
        anim.start(QPropertyAnimation.DeletionPolicy.DeleteWhenStopped)

    def jump_to_reply(self):
        self.paint_timer.stop(); self.draw()
        self.tabs.setCurrentIndex(0)
        self._smooth_scroll_to_bottom()

    def draw(self):
        messages = self.chat_messages[-40:]
        self.chat_stack.setCurrentIndex(1 if messages or self.stream else 0)
        self.sync_welcome()
        parts = []
        self.view.clipboard_images={}
        self.view.allowed_roots=[Path(root).resolve() for root in self.cfg.get('roots',[])]
        self.code_actions={}
        flags = QTextDocument.MarkdownFeature.MarkdownDialectGitHub | QTextDocument.MarkdownFeature.MarkdownNoHTML
        def markdown(content):
            doc = QTextDocument()
            font=self.view.font();font.setPointSizeF(self.preview_font_size*72.0/max(1,self.view.logicalDpiY()));doc.setDefaultFont(font)
            doc.setMarkdown(content, flags)
            body = re.search(r'<body[^>]*>(.*)</body>', doc.toHtml(), re.S)
            from assistant.themes import normalize_chat_html
            return normalize_chat_html(body.group(1)) if body else html.escape(content)
        def assistant_block(content, cache=True, message_index=None, media=None):
            from assistant.code_files import fences
            has_code=bool(fences(content))
            if cache and not has_code and not media and content in self.html_cache:return self.html_cache[content]
            safe_content=self.render_code_content(content,markdown,message_index) if has_code else markdown(content)
            media_html=''
            for item in media or []:
                try:
                    path=Path(item.get('path','')).resolve(strict=True)
                    if not path.is_file() or not any(path.is_relative_to(root) for root in self.view.allowed_roots):continue
                    if item.get('kind')=='image' and path.suffix.lower() in ('.png','.jpg','.jpeg','.webp'):
                        import hashlib
                        image=QImage(str(path))
                        if image.isNull():continue
                        key='chat-output-image:'+hashlib.sha256(str(path).encode()).hexdigest()
                        self.view.clipboard_images[key]=image
                        # Keep native pixel dimensions for small generated images. QImage.scaled()
                        # also enlarges by default, so a 512px SD-Turbo result was being stretched
                        # to 640px in chat and looked softer than the saved PNG.
                        if image.width()>768 or image.height()>640:
                            preview=image.scaled(768,640,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation)
                        else:
                            preview=image
                        media_html+='<p><img src="'+html.escape(key,quote=True)+'" width="'+str(preview.width())+'" height="'+str(preview.height())+'"></p>'
                    elif item.get('kind')=='video' and path.suffix.lower()=='.mp4':
                        import base64
                        token=base64.urlsafe_b64encode(str(path).encode()).decode().rstrip('=')
                        media_html+='<p><a href="chatai-media:'+token+'">▶ Mở video MP4</a></p>'
                    elif item.get('kind')=='file':
                        import base64
                        ft=base64.urlsafe_b64encode(str(path).encode()).decode().rstrip('=')
                        fdt=base64.urlsafe_b64encode(str(path.parent).encode()).decode().rstrip('=')
                        media_html+=('<p>📄 <a href="chatai-file:'+ft+'">'+html.escape(path.name)+'</a>'
                                     +'&nbsp;&nbsp;<a href="chatai-folder:'+fdt+'">📁 Mở thư mục</a></p>')
                except Exception:continue
            if media_html and not safe_content:safe_content=media_html
            elif media_html:safe_content+='<br>'+media_html
            block = ('<table width="100%" cellspacing="0" cellpadding="10"><tr>'
                    + '<td width="35" valign="top"><img src="' + html.escape(QUrl.fromLocalFile(str(ROOT / 'logo_chat_ai.png')).toString(), quote=True) + '" width="32" height="32"></td>'
                    + '<td>' + safe_content + '</td></tr></table><br>')
            if len(self.html_cache) > 80: self.html_cache.clear()
            if cache and not has_code and not media:self.html_cache[content]=block
            return block
        for msg in messages:
            if msg['role'] == 'user':
                content = html.escape(msg['content']).replace('\n', '<br>')
                from assistant.message_attachments import attachment_html
                content+=attachment_html(msg.get('documents',[]))
                for encoded in msg.get('images',[]):
                    import hashlib
                    image=QImage.fromData(QByteArray.fromBase64(encoded.encode('ascii')))
                    if image.isNull():continue
                    key='chat-image:'+hashlib.sha256(encoded.encode('ascii')).hexdigest()
                    self.view.clipboard_images[key]=image
                    thumb=image.scaled(320,220,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation)
                    content+='<br><img src="'+key+'" width="'+str(thumb.width())+'" height="'+str(thumb.height())+'"><br>'

                parts.append('<table width="100%" cellspacing="0" cellpadding="0"><tr><td width="20%"></td><td>'
                    '<table width="100%" bgcolor="'+bubble_color(self.preview_theme)+'" cellspacing="0" cellpadding="16"><tr><td>'+content+'</td></tr></table></td></tr></table><br>')
            else: parts.append(assistant_block(msg['content'],message_index=msg.get('source_index'),media=msg.get('media')))
        if self.stream: parts.append(assistant_block(self.stream, cache=False))
        if len(self.chat_messages) > 40: parts.insert(0, '<p><i>Hiển thị 40 tin nhắn gần nhất; toàn bộ lịch sử vẫn được lưu.</i></p>')
        self.view.document().setDefaultStyleSheet(chat_style(self.preview_theme,self.preview_font_size))
        bar = self.view.verticalScrollBar(); old_value = bar.value()
        follow = bar.maximum() - old_value < 24 or not self.busy()
        self.view.setHtml('<html><body>'+''.join(parts)+'</body></html>')
        self.view.enforce_text_size(self.preview_font_size)
        def mark_bubbles(frame):
            for child in frame.childFrames():
                if isinstance(child,QTextTable) and child.rows()==1 and child.columns()==1:
                    fmt=child.format();fmt.setProperty(1048581,True);fmt.setBackground(QColor(bubble_color(self.preview_theme)));child.setFormat(fmt)
                mark_bubbles(child)
        mark_bubbles(self.view.document().rootFrame())
        bar.setValue(bar.maximum() if follow else old_value)

    def render(self, initial=None):
        self.sync_support_identity()
        self.admin_button.setVisible(admin_session(self.server_session))
        if hasattr(self,'api_key_group'):self.api_key_group.setVisible(admin_session(self.server_session) and self.tabs.currentIndex()==3)
        if not admin_session(self.server_session) and hasattr(self,'admin_table'):
            self.admin_table.setRowCount(0);self.admin_users=[]
            if self.tabs.currentWidget()==self.admin_scroll:self.tabs.setCurrentIndex(0)
        history_rows=self.store.list(self.server_session['username'] if self.server_session else GUEST_OWNER,include_empty=False,limit=100)
        if not any(cid==self.cid for cid,_ in history_rows):
            title=self.store.conversation_title(self.cid,self.server_session['username'] if self.server_session else GUEST_OWNER)
            if title:history_rows.append((self.cid,title))
        history_key=(self.cid,self.preview_font_size,tuple(history_rows))
        rebuild_history=getattr(self,'history_render_key',None)!=history_key
        self.history_render_key=history_key
        if rebuild_history:self.history.clear()
        for cid, title in history_rows if rebuild_history else []:
            item = QListWidgetItem(title); item.setData(Qt.ItemDataRole.UserRole, cid); self.history.addItem(item)
            if cid == self.cid: self.history.setCurrentItem(item)
            row=QWidget();buttons=QHBoxLayout(row);buttons.setContentsMargins(4,2,4,2);buttons.setSpacing(2)
            choose=QPushButton(row.fontMetrics().elidedText(title,Qt.TextElideMode.ElideRight,165));choose.setToolTip(title)
            row.setMinimumHeight(48)
            choose.setMinimumHeight(36)
            choose.setStyleSheet('text-align:left;background:transparent;border:none;padding:0 4px;')
            choose.clicked.connect(lambda checked=False,i=item:self.select_chat(i));buttons.addWidget(choose,1)
            action=QPushButton('🗑' if cid==self.cid else '⋯');action.setFixedSize(30,30)
            action.setStyleSheet('padding:0;border:none;background:transparent;')
            action.setToolTip('Xóa cuộc trò chuyện' if cid==self.cid else 'Đổi tên hoặc xóa')
            action.setAccessibleName(action.toolTip())
            if cid==self.cid:action.clicked.connect(lambda checked=False,c=cid:self.delete_chat(c))
            else:
                menu=QMenu(action);menu.addAction('Đổi tên',lambda checked=False,c=cid:self.rename_chat(c));menu.addAction('Xóa',lambda checked=False,c=cid:self.delete_chat(c));action.setMenu(menu)
            buttons.addWidget(action);item.setSizeHint(QSize(210, max(64,row.sizeHint().height()+20)));self.history.setItemWidget(item,row)
        self.filter_history(self.history_search.text())
        selected=self.history.currentItem()
        full_title=selected.text() if selected else 'Chat AI'
        self.conversation_title.setText(self.conversation_title.fontMetrics().elidedText(full_title,Qt.TextElideMode.ElideRight,300))
        self.conversation_title.setToolTip(full_title)
        state = initial["state"] if initial else self.store.load(self.cid)
        self.saved_code_files=state.get('generated_files',[])
        self.chat_messages = [{'role': m['role'], 'content': m.get('content',''), 'images': m.get('images',[]), 'documents':m.get('documents',[]),
                               'media':m.get('media',[]),'source_index':i} for i,m in enumerate(state['messages'])
                              if m['role'] in ('user', 'assistant') and (m.get('content') or m.get('media'))]
        authorized=bool(self.server_session)
        self.chat_register.setVisible(not authorized)
        name=(self.server_session.get('fullname') or self.server_session['username']) if authorized else 'Khách · Dùng thử'
        self.profile_button.setText(self.profile_button.fontMetrics().elidedText(name,Qt.TextElideMode.ElideRight,115)); self.profile_button.setToolTip(html.escape(name))
        self.refresh_account_ui()
        self.sync_send_button(not state['running'] and not state['pending'])
        self.input.setEnabled(True)
        self.chat_login.setText('Tài khoản: '+self.server_session['username'] if authorized else 'Đăng nhập')
        self.resume_btn.setVisible(bool(state['running'] or state['pending']))
        self.resume_btn.setEnabled((authorized or self.trial.can_continue(self.cid,state) or state.get('model')==CLOUD_MODEL) and not self.busy() and bool(state['running'] or state['pending']))
        self.status.setText(('Bạn có thể bắt đầu trò chuyện.' if self.trial.remaining() else 'Đăng nhập để tiếp tục trò chuyện.') if not authorized else ('Đang chờ xác nhận' if state['pending'] else ('Lượt bị ngắt: bấm Tiếp tục' if state['running'] else 'Sẵn sàng')))
        self.draw()
        self.refresh_answer_actions(state)

    def refresh_answer_actions(self,state):
        while self.answer_actions_layout.count():
            item=self.answer_actions_layout.takeAt(0)
            if item.widget():item.widget().deleteLater()
        owner=self.server_session['username'] if self.server_session else None
        visible=bool(owner and state.get('account_username')==owner and not state['running'] and not self.busy())
        index=next((i for i in range(len(state['messages'])-1,-1,-1) if state['messages'][i]['role']=='assistant' and not state['messages'][i].get('tool_calls')),None)
        self.answer_actions.setVisible(visible and index is not None)
        if not visible or index is None:return
        for vote,label in ((1,'👍'),(-1,'👎')):
            button=QPushButton(label);button.setFixedWidth(44)
            button.setToolTip('Đánh giá câu trả lời mới nhất')
            button.clicked.connect(lambda checked=False,v=vote,i=index,c=self.cid,o=owner:self.rate_answer(o,c,i,v))
            self.answer_actions_layout.addWidget(button)
        for suggestion in state.get('followups',[])[:3]:
            button=QPushButton(suggestion['label']);button.setToolTip(suggestion['prompt'])
            button.clicked.connect(lambda checked=False,p=suggestion['prompt']:self.fill_prompt(p))
            self.answer_actions_layout.addWidget(button)
        self.answer_actions_layout.addStretch(1)

    def rate_answer(self,owner,cid,index,vote):
        if self.busy() or not self.server_session or self.server_session['username']!=owner:return
        comment=''
        if vote<0:
            comment,accepted=QInputDialog.getText(self,'Góp ý cho Chat AI','Điểm nào cần sửa? (có thể để trống)')
            if not accepted:return
        try:
            from assistant.feedback import FeedbackStore
            FeedbackStore(self.store).put(owner,cid,index,vote,comment)
            self.status.setText('Đã lưu đánh giá. Cảm ơn bạn đã góp ý.')
        except Exception as error:QMessageBox.warning(self,'Đánh giá',str(error))

    def select_chat(self, item):
        if not self.busy():
            target=item.data(Qt.ItemDataRole.UserRole)
            owner=self.server_session['username'] if self.server_session else GUEST_OWNER
            if not owner or self.store.load(target).get('account_username')!=owner:return
            self.remove_attachment();self.remove_documents();self.cid=target;self.stream=''
            self.store.remember_conversation(self.server_session['username'] if self.server_session else GUEST_OWNER,self.cid)
            self.tabs.setCurrentIndex(0);self.render()

    def new_chat(self):
        if not self.busy():
            self.remove_attachment();self.remove_documents();self.cid=self.store.create(persist=False);self.input.clear()
            self.store.remember_conversation(self.server_session['username'] if self.server_session else GUEST_OWNER,self.cid)
            self.tabs.setCurrentIndex(0);self.render()

    def rename_chat(self,cid):
        if self.busy():return
        owner=self.server_session['username'] if self.server_session else GUEST_OWNER
        if self.store.load(cid).get('account_username')!=owner:return
        title=next((title for ident,title in self.store.list(owner) if ident==cid),'')
        text,accepted=QInputDialog.getText(self,'Đổi tên hội thoại','Tên mới:',text=title)
        if accepted and text.strip():
            state=self.store.load(cid);state['custom_title']=text.strip()[:120];self.store.save(cid,state);self.render()

    def expire_delete_undo(self):
        self.deleted_chat_backup=None;self.undo_delete_box.hide()

    def undo_delete_chat(self):
        if not self.deleted_chat_backup:return
        backup,owner=self.deleted_chat_backup
        current=self.server_session['username'] if self.server_session else GUEST_OWNER
        if current!=owner:return
        try:
            cid=self.store.restore_deleted(backup,owner)
            self.undo_delete_timer.stop();self.expire_delete_undo()
            if not self.busy():self.cid=cid
            self.render();self.status.setText('Đã khôi phục hội thoại.')
        except Exception as exc:QMessageBox.warning(self,'Hoàn tác',str(exc))

    def delete_chat(self,cid=None):
        if self.busy():return
        if not isinstance(cid,str):cid=self.cid
        owner=self.server_session['username'] if self.server_session else GUEST_OWNER
        if self.store.load(cid).get('account_username')!=owner:return
        if not any(ident==cid for ident,_ in self.store.list(owner)):return
        try:
            with execution_lock(ROOT/'data/agent.lock'):
                backup=self.store.delete(cid,ROOT/'data/backups')
                if cid==self.cid:
                    rows=self.store.list(owner,include_empty=False)
                    self.cid=rows[0][0] if rows else self.store.create(persist=False)
            self.deleted_chat_backup=(backup,owner);self.undo_delete_box.show();self.undo_delete_timer.start(5000)
            self.stream='';self.html_cache.clear();self.input.clear();self.render()
            self.status.setText('Đã xóa hội thoại — có thể Hoàn tác trong 5 giây. Các tài liệu đã tạo được giữ nguyên.')
        except Exception as exc:QMessageBox.warning(self,'Xóa hội thoại',str(exc))

    def toggle_web(self):
        if self.busy():return
        if self.chat_mode.currentIndex()==5:
            self.expert_web_enabled=self.web_btn.isChecked()
            self.status.setText('Đã bật tìm kiếm mạng cho Chuyên gia.' if self.expert_web_enabled else 'Đã tắt tìm kiếm mạng cho Chuyên gia.')
            return
        if self.chat_mode.currentIndex()==4:
            restore=self.mode_return_index if self.mode_return_index is not None else 0
            self.mode_return_index=None
            self.chat_mode.setCurrentIndex(restore);self.status.setText('Đã tắt tìm kiếm mạng.');return
        self.mode_return_index=self.chat_mode.currentIndex()
        self.chat_mode.setCurrentIndex(4)
        self.status.setText('Tìm kiếm Bing đã sẵn sàng; không cần tải thư viện hoặc module mạng.')
        self.input.setFocus()

    def toggle_temporary_mode(self,mode):
        """Toggle image/video mode and restore the mode that was active before it."""
        if self.busy():return
        current=self.chat_mode.currentIndex()
        if current==mode:
            restore=self.mode_return_index if self.mode_return_index is not None else 0
            self.mode_return_index=None
            self.chat_mode.setCurrentIndex(restore)
            self.input.setFocus()
            return
        self.mode_return_index=current
        self.chat_mode.setCurrentIndex(mode)
        self.input.setFocus()

    def toggle_deep_analysis(self,checked):
        self.deep_analysis_enabled=bool(checked)
        self.status.setText('Đã bật phân tích sâu cho tin nhắn tiếp theo; AI sẽ rà soát lại trước khi trả lời.' if checked
                            else 'Đã tắt phân tích sâu; chat thường sẽ phản hồi nhanh hơn.')

    def open_normal_chat(self):
        if self.busy():return
        self.tabs.setCurrentIndex(0);self.chat_mode.setCurrentIndex(0)
        normal=getattr(self,'normal_model_before_experts',None)
        if normal:self.select_ai(normal)

    def open_expert_chat(self):
        if self.busy():return
        if self.chat_mode.currentIndex()!=5:self.normal_model_before_experts=self.model.currentText()
        self.tabs.setCurrentIndex(0);self.chat_mode.setCurrentIndex(5)
        if self.model.currentText() in REMOTE_MODELS:
            preferred=next((x for x in (self.cfg['default_model'],'qwen2.5:7b','qwen2.5:3b') if x in self.models),self.cfg['default_model'])
            self.select_ai(preferred)
            self.status.setText('Chuyên gia dùng AI local; AI chưa tải sẽ hỏi trước khi tải.')

    # ── Hỗ trợ Công việc ──────────────────────────────────────────────────────

    def open_work_support(self):
        """Mở tab Hỗ trợ Công việc (tab động, tạo lại mỗi lần)."""
        self.remove_dynamic_page('work_support_page_index')
        page = QWidget(); layout = QVBoxLayout(page); layout.setContentsMargins(0, 0, 0, 0); layout.setSpacing(16)
        self.button(layout, '← Quay lại chat', lambda: self.tabs.setCurrentIndex(0))
        title_lbl = QLabel('🗂️  Hỗ trợ Công việc')
        title_lbl.setStyleSheet('font-size:22px;font-weight:600;margin-bottom:4px;')
        layout.addWidget(title_lbl)

        # ── 1. Thư mục công việc ────────────────────────────────────────────
        folder_box = QGroupBox('📁 Thư mục công việc')
        fb = QVBoxLayout(folder_box)
        roots = self.cfg.get('whitelist', [])
        self._work_folder_list = QListWidget(); self._work_folder_list.setMaximumHeight(90)
        for r in roots:
            self._work_folder_list.addItem(r)
        fb.addWidget(self._work_folder_list)
        fb_row = QHBoxLayout()
        def _add_work_folder():
            from PySide6.QtWidgets import QFileDialog
            path = QFileDialog.getExistingDirectory(page, 'Chọn thư mục công việc')
            if not path: return
            if path not in self.cfg.get('whitelist', []):
                self.cfg.setdefault('whitelist', []).append(path)
                from assistant.config import save_config
                save_config(self.cfg)
                self._work_folder_list.addItem(path)
        def _scan_templates():
            folders = [self._work_folder_list.item(i).text() for i in range(self._work_folder_list.count())]
            if not folders:
                QMessageBox.information(page, 'Quét mẫu', 'Thêm ít nhất một thư mục trước.')
                return
            self._work_scan_templates(folders, page)
        self.button(fb_row, 'Thêm thư mục…', _add_work_folder)
        self.button(fb_row, '🔍 Quét mẫu & Form biểu', _scan_templates)
        fb.addLayout(fb_row)
        layout.addWidget(folder_box)

        # ── 2. Kho mẫu ──────────────────────────────────────────────────────
        template_box = QGroupBox('📄 Mẫu thuyết minh & Form biểu')
        tb = QVBoxLayout(template_box)
        self._template_status = QLabel('Nhấn "Quét mẫu" để nạp tài liệu mẫu từ thư mục công việc.')
        self._template_status.setWordWrap(True)
        self._template_status.setStyleSheet('color:#9aa0a6;font-size:12px;')
        tb.addWidget(self._template_status)
        self._template_list = QListWidget(); self._template_list.setMaximumHeight(130)
        self._template_list.setToolTip('Bấm đúp để gửi mẫu này vào chat')
        self._template_list.itemDoubleClicked.connect(self._send_template_to_chat)
        tb.addWidget(self._template_list)
        tb_row = QHBoxLayout()
        self.button(tb_row, 'Gửi mẫu vào chat', lambda: self._send_template_to_chat(self._template_list.currentItem()))
        self.button(tb_row, 'Xóa khỏi danh sách', self._remove_template)
        tb.addLayout(tb_row)
        layout.addWidget(template_box)
        self._refresh_template_list()

        # ── 3. Tạo nhanh cho dự án mới ──────────────────────────────────────
        project_box = QGroupBox('✏️  Dự án mới — Tạo tài liệu nhanh')
        pb = QVBoxLayout(project_box)
        pf = QFormLayout()
        self._work_project_name = QLineEdit(); self._work_project_name.setPlaceholderText('Ví dụ: Nhà phố Quận 7 - A1')
        self._work_project_name.setMaxLength(120)
        pf.addRow('Tên dự án:', self._work_project_name)
        self._work_project_type = QComboBox()
        for t in ['Móng đơn', 'Móng băng', 'Móng bè', 'Móng cọc khoan nhồi', 'Móng cọc ép',
                  'Tường chắn đất', 'Mái dốc / taluy', 'Nền đường', 'Mặt cắt địa chất', 'Khác']:
            self._work_project_type.addItem(t)
        pf.addRow('Loại công trình:', self._work_project_type)
        self._work_project_scale = QComboBox()
        for s in ['1:50', '1:100', '1:200', '1:500', '1:1000', 'Không cần']:
            self._work_project_scale.addItem(s)
        self._work_project_scale.setCurrentText('1:100')
        pf.addRow('Tỉ lệ bản vẽ:', self._work_project_scale)
        pb.addLayout(pf)
        btn_grid = QHBoxLayout()
        self.button(btn_grid, '📝 Viết thuyết minh', lambda: self._quick_action('thuyet_minh'))
        self.button(btn_grid, '📐 Vẽ mặt cắt CAD', lambda: self._quick_action('mat_cat_cad'))
        pb.addLayout(btn_grid)
        btn_grid2 = QHBoxLayout()
        self.button(btn_grid2, '📊 Bảng tính Excel', lambda: self._quick_action('bang_tinh'))
        self.button(btn_grid2, '📋 Báo cáo Word', lambda: self._quick_action('bao_cao_word'))
        pb.addLayout(btn_grid2)
        btn_grid3 = QHBoxLayout()
        self.button(btn_grid3, '🔩 Tính lún (SoilFim)', lambda: self._quick_action('tinh_lun'))
        self.button(btn_grid3, '📈 Tính ổn định mái', lambda: self._quick_action('on_dinh_mai'))
        pb.addLayout(btn_grid3)
        layout.addWidget(project_box)

        # ── 4. Công cụ nhanh ────────────────────────────────────────────────
        tools_box = QGroupBox('🔧 Công cụ nhanh — không cần nhập dự án')
        qb = QVBoxLayout(tools_box)
        quick_items = [
            ('🗺️  Vẽ mặt cắt địa chất nhiều lớp', 'Vẽ mặt cắt địa chất gồm nhiều lớp đất. Hỏi tôi số lớp và thông số từng lớp.'),
            ('🏗️  Tạo bản vẽ móng AutoCAD', 'Tạo bản vẽ AutoCAD cho móng công trình. Hỏi tôi loại móng và kích thước.'),
            ('📐 Tính toán sức chịu tải cọc', 'Tính sức chịu tải cọc theo phương pháp tĩnh. Hỏi tôi thông số đất và cọc.'),
            ('📊 Lập bảng tổng hợp số liệu địa chất', 'Lập bảng tổng hợp số liệu địa chất từ báo cáo khảo sát. Hỏi tôi số liệu hố khoan.'),
            ('📝 Viết thuyết minh từ mẫu đã học', 'Dùng mẫu thuyết minh đã lưu trong thư viện để viết thuyết minh tính toán mới. Hỏi tôi loại công trình.'),
            ('🔢 Kiểm tra nội lực / tổ hợp tải trọng', 'Kiểm tra nội lực và tổ hợp tải trọng cho cấu kiện. Hỏi tôi loại kết cấu và số liệu.'),
            ('🌊 Tính lún cố kết theo thời gian', 'Tính lún cố kết theo thời gian cho nền đất. Hỏi tôi thông số lớp đất và tải trọng.'),
            ('🔍 Tìm tiêu chuẩn TCVN / QCVN liên quan', 'Tìm và giải thích các tiêu chuẩn TCVN / QCVN liên quan đến yêu cầu kỹ thuật. Hỏi tôi lĩnh vực cần tra.'),
        ]
        for label, prompt in quick_items:
            btn = QPushButton(label)
            btn.clicked.connect(lambda chk=False, p=prompt: self._send_quick_prompt(p))
            btn.setStyleSheet('text-align:left;padding:6px 10px;')
            qb.addWidget(btn)
        layout.addWidget(tools_box)
        layout.addStretch(1)

        self.add_scroll_page(page, '🗂️ Công việc')
        self.work_support_page_index = self.tabs.count() - 1
        self.tabs.setCurrentIndex(self.work_support_page_index)

    def _refresh_template_list(self):
        """Điền danh sách mẫu đã lưu trong RAG/library vào widget."""
        if not hasattr(self, '_template_list'): return
        self._template_list.clear()
        try:
            from assistant.rag import RagSearch
            rag = RagSearch(self.store, self.cfg)
            hits = rag.rag_search('thuyết minh tính toán mẫu form biểu')
            seen = set()
            for h in hits.get('sources', [])[:20]:
                src = h.get('source') or h.get('file') or ''
                if src and src not in seen:
                    seen.add(src)
                    from pathlib import Path as _P
                    item = QListWidgetItem('📄 ' + _P(src).name)
                    item.setData(Qt.ItemDataRole.UserRole, src)
                    item.setToolTip(src)
                    self._template_list.addItem(item)
            if seen:
                self._template_status.setText(f'Đã nạp {len(seen)} mẫu từ thư viện RAG.')
        except Exception:
            pass

    def _work_scan_templates(self, folders, parent_widget):
        """Quét thư mục, nạp tài liệu mẫu vào RAG và library."""
        from pathlib import Path
        exts = {'.docx', '.xlsx', '.pdf', '.txt', '.md'}
        files = []
        for folder in folders:
            p = Path(folder)
            if p.is_dir():
                for f in p.rglob('*'):
                    if f.suffix.lower() in exts and f.is_file() and f.stat().st_size < 32*1024*1024:
                        files.append(f)
        if not files:
            QMessageBox.information(parent_widget, 'Quét mẫu', 'Không tìm thấy file Word/Excel/PDF/TXT trong thư mục đã chọn.')
            return
        count = len(files)
        reply = QMessageBox.question(parent_widget, 'Quét mẫu',
            f'Tìm thấy {count} file.\nNạp vào thư viện để AI học mẫu?\n\n' +
            '\n'.join(str(f.name) for f in files[:8]) + (f'\n... và {count-8} file khác' if count > 8 else ''),
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.Cancel)
        if reply != QMessageBox.StandardButton.Yes: return
        scanned = {'ok': 0, 'fail': 0}
        def task(emit):
            from assistant.rag import RagSearch
            rag = RagSearch(self.store, self.cfg)
            for i, f in enumerate(files):
                emit({'type': 'status', 'text': f'Đang nạp ({i+1}/{count}): {f.name}…'})
                try:
                    rag.rag_index(str(f))
                    scanned['ok'] += 1
                except Exception:
                    scanned['fail'] += 1
            return scanned
        def done(result):
            msg = f"Đã nạp {result['ok']} mẫu vào thư viện."
            if result['fail']: msg += f" {result['fail']} file lỗi (bỏ qua)."
            self._template_status.setText(msg)
            self._refresh_template_list()
            QMessageBox.information(parent_widget, 'Quét mẫu', msg)
        self.work({'type': 'status', 'text': 'Đang quét thư mục…'} if False else None, None)
        self.work(task, done)

    def _send_template_to_chat(self, item):
        """Gửi mẫu đã chọn vào chat để AI phân tích và học."""
        if not item: return
        src = item.data(Qt.ItemDataRole.UserRole) or ''
        name = item.text().lstrip('📄 ')
        self.tabs.setCurrentIndex(0)
        self.input.setPlainText(
            f'Hãy đọc và phân tích cấu trúc mẫu tài liệu "{name}". '
            f'Ghi nhớ định dạng, các mục tiêu đề, bảng biểu và chỗ điền số liệu để '
            f'dùng làm mẫu cho dự án mới. Đường dẫn: {src}'
        )
        self.input.setFocus()

    def _remove_template(self):
        item = self._template_list.currentItem()
        if not item: return
        self._template_list.takeItem(self._template_list.row(item))

    def _quick_action(self, action: str):
        """Tạo prompt từ thông tin dự án đã nhập và gửi vào chat."""
        name = self._work_project_name.text().strip()
        kind = self._work_project_type.currentText()
        scale = self._work_project_scale.currentText()
        project_ctx = f' cho dự án "{name}"' if name else ''
        prompts = {
            'thuyet_minh': (
                f'Viết thuyết minh tính toán {kind.lower()}{project_ctx}. '
                f'Dùng mẫu thuyết minh đã lưu trong thư viện nếu có. '
                f'Hỏi tôi các số liệu cần thiết (địa chất, tải trọng, kích thước).'
            ),
            'mat_cat_cad': (
                f'Vẽ mặt cắt {kind.lower()}{project_ctx} bằng AutoCAD, tỉ lệ {scale}. '
                f'Hỏi tôi thông số lớp đất, kích thước kết cấu và cao độ nền.'
            ),
            'bang_tinh': (
                f'Tạo bảng tính Excel{project_ctx} cho {kind.lower()}. '
                f'Dùng form biểu mẫu đã lưu nếu có. '
                f'Hỏi tôi số liệu đầu vào cần điền.'
            ),
            'bao_cao_word': (
                f'Tạo báo cáo kỹ thuật Word{project_ctx} về {kind.lower()}. '
                f'Dùng mẫu báo cáo đã lưu trong thư viện nếu có. '
                f'Hỏi tôi nội dung cần đưa vào.'
            ),
            'tinh_lun': (
                f'Tính lún cố kết{project_ctx} cho {kind.lower()}. '
                f'Hỏi tôi: số lớp đất, chiều dày, e₀, Cc, Cs, áp lực tiền cố kết, tải trọng và diện tích gia tải.'
            ),
            'on_dinh_mai': (
                f'Kiểm tra ổn định mái dốc{project_ctx}. '
                f'Hỏi tôi: góc dốc, chiều cao, thông số cường độ (c, φ), mực nước ngầm và phương pháp tính (Bishop, Fellenius).'
            ),
        }
        prompt = prompts.get(action, f'Hỗ trợ {action}{project_ctx}.')
        self._send_quick_prompt(prompt)

    def _send_quick_prompt(self, prompt: str):
        """Chuyển sang chat và điền prompt sẵn."""
        self.tabs.setCurrentIndex(0)
        self.input.setPlainText(prompt)
        self.input.setFocus()

    # ── Kết thúc Hỗ trợ Công việc ─────────────────────────────────────────

    def show_expert_details(self):
        state=self.store.load(self.cid);details=state.get('orchestration')
        if not details:
            QMessageBox.information(self,'Chi tiết xử lý','Gửi yêu cầu trong mục Chuyên gia để xem kế hoạch và kết quả từng bước.');return
        dialog=QDialog(self);dialog.setWindowTitle('Chi tiết điều phối');dialog.resize(760,580)
        layout=QVBoxLayout(dialog);view=QPlainTextEdit();view.setReadOnly(True)
        view.setPlainText(json.dumps(details,ensure_ascii=False,indent=2,default=str));layout.addWidget(view)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Close);buttons.rejected.connect(dialog.reject);layout.addWidget(buttons);dialog.exec()

    def reselect_expert(self):
        if self.busy():return
        labels=[self.expert_choice.itemText(i) for i in range(self.expert_choice.count())]
        label,accepted=QInputDialog.getItem(self,'Chọn lại chuyên gia','Ưu tiên cho yêu cầu gửi tiếp theo:',labels,self.expert_choice.currentIndex(),False)
        if not accepted:return
        self.expert_choice.setCurrentIndex(labels.index(label))
        state=self.store.load(self.cid);last=next((m for m in reversed(state['messages']) if m['role']=='user'),None)
        if last:
            self.input.setPlainText(last['content'])
            if last.get('images'):self.pending_image=last['images'][0]
            self.pending_documents=[d.get('path') or d.get('source') for d in last.get('documents',[]) if (d.get('path') or d.get('source')) and Path(d.get('path') or d.get('source')).is_file()]
        self.input.setFocus();self.status.setText('Đã chọn lại chuyên gia. Bấm Gửi để xử lý lại yêu cầu.')

    def mode_changed(self, index):
        self.image_btn.setChecked(index == 2); self.video_btn.setChecked(index == 3); self.web_btn.setChecked(index==4 or (index==5 and getattr(self,'expert_web_enabled',False)))
        placeholders = ['Hỏi Chat AI…', 'Yêu cầu xử lý tài liệu, file hoặc tra cứu…',
                        'Mô tả hình ảnh bạn muốn tạo…', 'Mô tả video bạn muốn tạo từ ảnh AI…', 'Nhập câu hỏi để tìm kiếm mạng (tối đa500 ký tự)…']
        placeholders.append('Gửi văn bản, ảnh hoặc file; các chuyên gia phối hợp tự động…')
        self.input.setPlaceholderText(placeholders[index])
        self.expert_toolbar.setVisible(index==5);self.expert_reselect.setVisible(index==5)
        if index==5:self.status.setText('Chuyên gia phối hợp theo yêu cầu. Web chỉ chạy khi bạn bật nút Tìm kiếm mạng.')
        if index == 1:
            if not CHAT_MODELS.get(self.model.currentText(), {}).get('tools',False):
                compatible = [name for name, spec in CHAT_MODELS.items() if spec['tools'] and name in self.models]
                fallback = compatible[0] if compatible else 'qwen2.5:3b'
                self.select_ai(fallback)
            self.status.setText('Office và Python dùng module đã bật. Tra web chỉ khi bật nút Tìm kiếm mạng.')
        elif index==2:self.status.setText('Tạo ảnh bằng mô-đun ảnh chuyên dụng; AI chat không cần tải riêng.')
        elif index==3:self.status.setText('Tạo video bằng mô-đun video; ghép ảnh có sẵn dùng thư viện Python cơ bản.')

    def apply_account_model(self,name):
        if name not in REMOTE_MODELS and name not in CHAT_MODELS:return
        self.select_ai(name)

    def refresh_ai_choices(self):
        """Refresh the quick selector from the current online/local catalogs."""
        self.online_ai_names=list(REMOTE_MODELS)
        self.local_ai_names=list(dict.fromkeys([self.cfg['default_model'],self.cfg['code_model'],*CHAT_MODELS]))

    def quick_ai_changed(self,index):
        choice=self.quick_provider.itemData(index)
        if choice not in ('online','local'):return
        self.refresh_ai_choices()
        names=self.online_ai_names if choice=='online' else self.local_ai_names
        if not names:return
        current=self.model.currentText()
        if choice=='online':
            preferred=current if current in names else self.last_online_model if self.last_online_model in names else names[0]
        else:
            preferred=current if current in names else self.cfg['default_model'] if self.cfg['default_model'] in names else names[0]
        self.set_ai_choices(choice,preferred,notify=True)

    def set_ai_choices(self,choice,preferred,notify=True):
        """Show only the choices for the currently selected AI connection type."""
        self.refresh_ai_choices()
        names=self.online_ai_names if choice=='online' else self.local_ai_names
        if not names:return
        selected=preferred if preferred in names else names[0]
        self.model.blockSignals(True)
        self.model.clear();self.model.addItems(names);self.model.setCurrentText(selected)
        self.model.blockSignals(False)
        self.quick_provider.blockSignals(True)
        self.quick_provider.setCurrentIndex(self.quick_provider.findData(choice))
        self.quick_provider.blockSignals(False)
        if notify:self.model_changed(selected)

    def select_ai(self,name):
        """Select an AI and switch the filtered dropdown when its provider type changes."""
        self.refresh_ai_choices()
        choice='online' if name in REMOTE_MODELS else 'local'
        names=self.online_ai_names if choice=='online' else self.local_ai_names
        visible=[self.model.itemText(i) for i in range(self.model.count())]
        if visible!=names:
            self.set_ai_choices(choice,name,notify=True)
        elif self.model.currentText()!=name:
            self.model.setCurrentText(name)

    def model_changed(self,name):
        provider=REMOTE_MODELS.get(name,'local')
        if hasattr(self,'quick_provider'):
            self.quick_provider.blockSignals(True)
            index=self.quick_provider.findData('online' if name in REMOTE_MODELS else 'local')
            if index>=0:self.quick_provider.setCurrentIndex(index)
            self.quick_provider.blockSignals(False)
        if name in REMOTE_MODELS:self.last_online_model=name
        if hasattr(self,'settings_provider'):self.settings_provider.setCurrentText(name if name in REMOTE_MODELS else 'AI trên máy')
        self.cfg['chat_provider']=provider
        if name in CHAT_MODELS:self.cfg['default_model']=name
        from assistant.model_preferences import save_model
        save_model(self.store,getattr(self,'server_session',None),name)
        if name in REMOTE_MODELS:self.model.setToolTip('AI trực tuyến dùng key chung trên server.');return
        if name in CHAT_MODELS:
            self.model.setToolTip(CHAT_MODELS[name]['label'])
            if not self.models:QTimer.singleShot(200,self.refresh_startup_models)

    def check_gpu(self):
        def task(emit):
            client = LocalOllamaClient(host=self.cfg['ollama_host'], timeout=8)
            lines = []
            for model in client.ps().models:
                size, vram = model.size or 0, model.size_vram or 0
                fraction = min(100, round(100 * vram / size)) if size else 0
                lines.append(f'{model.model}: VRAM {vram/1024**3:.2f} GiB / tổng {size/1024**3:.2f} GiB ({fraction}% theo bộ nhớ).')
            return '\n'.join(lines) or 'Chưa có model đang nạp. Gửi một tin nhắn rồi kiểm tra lại.'
        self.work(task, lambda text: QMessageBox.information(self, 'Model / GPU', text))

    def refresh_startup_models(self):
        # Online login must not compete with importing the local AI stack.
        if self.model.currentText() in REMOTE_MODELS:return
        if self.busy():
            QTimer.singleShot(500,self.refresh_startup_models);return
        self.refresh_models()

    def refresh_models(self):
        if self.model_probe_running:return
        import threading
        self.model_probe_running=True
        host=self.cfg['ollama_host']
        def probe():
            try:
                client=LocalOllamaClient(host=host,timeout=8)
                result={'models':{m.model for m in client.list().models}}
            except Exception:
                result={'error':'Chưa kết nối được Ollama. Bạn vẫn có thể đăng nhập và mở Cài đặt.'}
            try:self.model_probe_finished.emit(result)
            except RuntimeError:pass  # Window was closed while the daemon was probing.
        threading.Thread(target=probe,daemon=True).start()

    def models_probed(self,result):
        self.model_probe_running=False
        if 'models' in result:
            self.models=result['models']
            message='Ollama đã kết nối. '+('Chọn model và gửi tin nhắn.' if self.server_session else 'Bạn có thể bắt đầu trò chuyện.')
        else:message=result['error']
        if self.model.currentText() in REMOTE_MODELS:message='Đã chọn Cloudflare AI. Nội dung chat sẽ được gửi lên server.'
        if not self.busy():self.status.setText(message)

    def start_direct_media(self,prompt,mode):
        """Create media directly; a chat LLM is not required to call the media engine."""
        if mode==2 and self.pending_image:
            QMessageBox.information(self,'Tạo ảnh','Tạo ảnh từ mô tả chưa nhận ảnh nguồn để chỉnh sửa. Hãy bỏ ảnh đính kèm hoặc dùng AI đọc ảnh trước.')
            return
        image_paths=list(self.pending_media_paths) if mode==3 else []
        module='media_basic' if mode==3 and image_paths else 'media'
        if not self.manager.ready(module):
            if module=='media' and (media_model_dir(ROOT,self.manager.media_variant())/'ready.json').is_file():
                # The large model is already in the user's Drive folder. If only
                # this venv's Python libraries are missing, prepare them in the
                # background instead of presenting a misleading model download prompt.
                try:
                    self.manager.request('media',automatic=True)
                    self.tabs.setCurrentIndex(1)
                    self.status.setText('Đã có model ảnh trên Google Drive; đang chuẩn bị thư viện còn thiếu, không tải lại model.')
                    self.poll()
                except Exception as error:
                    QMessageBox.warning(self,'Chuẩn bị tạo ảnh',str(error))
                return
            self.download(module)
            return
        if mode==2:
            image_model=MEDIA_VARIANTS[self.manager.media_variant()]['display']
            action='image_generate';description='Tạo ảnh PNG bằng AI ảnh local '+image_model+'.'
            args={'prompt':prompt}
        elif image_paths:
            action='video_from_images';description='Ghép các ảnh đã chọn thành video MP4.'
            args={'image_paths':image_paths,'seconds_per_image':3}
        else:
            image_model=MEDIA_VARIANTS[self.manager.media_variant()]['display']
            action='video_generate';description='Tạo ảnh bằng '+image_model+' rồi ghép thành video MP4.'
            args={'prompts':[prompt]}
        confirm=QMessageBox.question(self,'Xác nhận tạo nội dung',
            description+'\n\nMô tả: '+prompt[:1200]+'\n\nỨng dụng sẽ lưu tệp mới trong workspace/outputs. Tiếp tục?',
            QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No)
        if confirm!=QMessageBox.StandardButton.Yes:return
        if not self.server_session and self.trial.remaining()==0:
            self.login_dialog();return

        cid=self.cid;session=dict(self.server_session) if self.server_session else None
        selected_media_model=self.model.currentText()
        owner=session['username'] if session else GUEST_OWNER
        self.sent_prompt=prompt;self.input.clear();self.remove_attachment();self.remove_documents()
        self.reply_active=True;self.reply_frame=0;self.reply_logo.show();self.reply_dots.show();self.reply_timer.start()
        self.status.setText('Đang chuẩn bị bộ tạo ảnh/video…')

        def task(emit):
            from assistant.accounts import request_account
            from assistant.capabilities import Capabilities
            from assistant.collaboration import collect_artifacts
            from assistant.excel import ExcelTools
            import time as _time
            with execution_lock(ROOT/'data/agent.lock'):
                if session:
                    _cache = getattr(self, '_auth_check_cache', {})
                    _key = (session['endpoint'], session['username'], session['key'])
                    if _time.monotonic() - _cache.get(_key, 0) > 60:
                        try:
                            request_account(session['endpoint'],'/api/models',{'username':session['username'],'key':session['key']},timeout=8)
                            _cache[_key] = _time.monotonic()
                            self._auth_check_cache = _cache
                        except Exception:raise RuntimeError('Phiên đăng nhập hết hạn. Hãy đăng nhập lại rồi thử tạo nội dung.') from None
                state=self.store.load(cid)
                if state.get('account_username') not in (None,owner):
                    raise RuntimeError('Hội thoại thuộc tài khoản khác. Hãy tạo cuộc trò chuyện mới.')
                state.setdefault('messages',[]).append({'role':'user','content':prompt})
                state.update(running=True,queue=[],pending=None,media_done=False,ui_mode=mode,
                             model=self.model.currentText(),account_username=owner)
                if not session:state['guest_trial_token']=__import__('uuid').uuid4().hex
                self.store.save(cid,state)
                if not session:self.trial.consume(cid,state)
                self.store.remember_conversation(owner,cid)
                def snapshot():
                    return [{'role':m['role'],'content':m.get('content',''),'images':m.get('images',[]), 'documents':m.get('documents',[]),
                             'media':m.get('media',[]),'source_index':i}
                            for i,m in enumerate(state['messages'])
                            if m['role'] in ('user','assistant') and (m.get('content') or m.get('media'))]
                emit({'type':'sent','cid':cid,'prompt':prompt})
                emit({'type':'snapshot','messages':snapshot()})
                emit({'type':'status','text':'Đang tạo ảnh/video bằng mô-đun chuyên dụng…'})
                media_args=dict(args)
                media_translation_error=None
                if action in ('image_generate','video_generate'):
                    try:
                        from assistant.media import to_english_prompt
                        emit({'type':'status','text':'Đang dịch prompt sang tiếng Anh bằng Qwen2.5 7B…'})
                        translated=to_english_prompt(
                            self.client,prompt,model='qwen2.5:7b',num_ctx=self.cfg.get('num_ctx',4096),
                            on_status=lambda text:emit({'type':'status','text':text}))
                        if action=='image_generate':media_args['prompt']=translated
                        else:media_args['prompts']=[translated]
                    except Exception as error:
                        media_translation_error='Chưa dịch được prompt sang tiếng Anh bằng Qwen2.5:7b; đã dừng tạo để tránh gửi prompt sai. ('+str(error)[:240]+')'
                audit=lambda name,details:self.store.audit(cid,name,details)
                excel=ExcelTools(self.cfg['roots'],ROOT/'data/backups',audit)
                caps=Capabilities(self.manager,excel,self.client,self.cfg,ROOT,audit,owner=owner,
                    storage_root=Path(self.store.path).parent,allow_web=False,
                    on_status=lambda text:emit({'type':'status','text':text}))
                expected='image' if mode==2 else 'video'
                try:
                    if media_translation_error:raise RuntimeError(media_translation_error)
                    plan=caps.prepare(action,media_args)
                    result=caps.commit(plan)
                    artifacts=collect_artifacts(result,self.cfg.get('roots',[]),action)
                    media=[item for item in artifacts if item.get('kind')==expected]
                    if not media:raise RuntimeError('Bộ tạo chưa trả về tệp '+('ảnh PNG' if mode==2 else 'video MP4')+'.')
                    state=self.store.load(cid)
                    state['messages'].append({'role':'assistant','content':'','media':media})
                    state.update(running=False,queue=[],pending=None,media_done=True,followups=[])
                    self.store.save(cid,state)
                    self.store.audit(cid,'media_delivered',{'kind':expected,'count':len(media),'direct':True})
                except Exception as exc:
                    state=self.store.load(cid)
                    state['messages'].append({'role':'assistant','content':'Chưa tạo được nội dung media: '+str(exc)[:500]})
                    state.update(running=False,queue=[],pending=None,media_done=False)
                    self.store.save(cid,state)
                    self.store.audit(cid,'media_failed',{'action':action,'error':str(exc)[:500],'direct':True})
                emit({'type':'snapshot','messages':snapshot()})
                return state

        self.work(task,self.after_chat,cancellable=False)

    def send(self):
        media_mode=self.chat_mode.currentIndex() in (2,3)
        if (self.model.currentText() not in REMOTE_MODELS or media_mode) and not self.server_session and self.trial.remaining()==0:
            self.login_dialog();return
        prompt = self.input.toPlainText().strip()
        if not prompt and self.pending_image:prompt='Hãy phân tích hình ảnh này bằng tiếng Việt.'
        if not prompt and self.pending_documents:prompt='Hãy tóm tắt tài liệu đính kèm bằng tiếng Việt.'
        if self.pending_documents and self.chat_mode.currentIndex() not in (3,4,5):self.chat_mode.setCurrentIndex(0)
        if not prompt or self.busy() or not self.send_btn.isEnabled(): return
        if self.chat_mode.currentIndex() in (2,3):
            self.start_direct_media(prompt,self.chat_mode.currentIndex())
            return
        from assistant.collaboration import collaboration_intent
        cooperation=collaboration_intent(prompt,mode=self.chat_mode.currentIndex()) if self.chat_mode.currentIndex()==5 else {'media_tool':None}
        if self.chat_mode.currentIndex()==5 and not cooperation.get('media_tool'):
            import re
            if not re.search(r'đừng|không tạo|chưa tạo|đề xuất|ý tưởng|trao đổi|nói trước',prompt,re.I):
                if re.search(r'(tạo|sinh)\s+(?:một\s+)?(?:video|clip)',prompt,re.I):cooperation['media_tool']='video_generate'
                elif re.search(r'(tạo|vẽ|sinh)\s+(?:một\s+)?(?:ảnh|hình|logo)',prompt,re.I):cooperation['media_tool']='image_generate'
        model = self.model.currentText()
        if model in REMOTE_MODELS and (cooperation['media_tool'] or self.chat_mode.currentIndex()==5):
            QMessageBox.information(self,'Phối hợp AI','Mục Chuyên gia phối hợp các AI trên máy. Hãy chọn AI local để dùng mục này.');self.choose_other_model();return
        if model in REMOTE_MODELS:
            from assistant.online_automation import use_automation
            if use_automation(prompt,self.cfg,self.store.load(self.cid),self.chat_mode.currentIndex()==1):
                self.online_windows_task(prompt=prompt);return
            if self.chat_mode.currentIndex() in (1,2,3):
                QMessageBox.information(self,'Chọn chế độ','Chế độ công cụ Office cần AI trên máy. Tệp Office, PDF, DXF và ảnh có thể gửi cùng AI trực tuyến.');return
            if len(prompt)>6000:
                QMessageBox.information(self,'Tin nhắn quá dài','Tin nhắn tối đa 6000 ký tự.');return
            self.sent_prompt=prompt;self.input.clear();self.cloud_chat_task(prompt);return
        if self.pending_image and self.chat_mode.currentIndex() not in (3,5):
            if not CHAT_MODELS[model].get('vision'):
                if QMessageBox.question(self,'Đọc ảnh','Ảnh cần model vision. Chuyển sang Gemma3 4B để đọc ảnh?') != QMessageBox.StandardButton.Yes:return
                self.select_ai('gemma3:4b'); model='gemma3:4b'
            if self.chat_mode.currentIndex() not in (4,5):self.chat_mode.setCurrentIndex(0)
        if not CHAT_MODELS[model]['tools'] and self.chat_mode.currentIndex() not in (0,4,5):
            self.chat_mode.setCurrentIndex(0)
            QMessageBox.information(self, 'Chế độ model', 'DeepSeek dùng Chat nhanh trong bản này. Xử lý Office/file chọn Qwen rồi chọn Dùng công cụ / Office.')
        if model not in self.models:
            self.download(model); return
        if self.chat_mode.currentIndex()==4:
            if len(prompt)>500:
                QMessageBox.warning(self,'Tìm kiếm mạng','Rút gọn câu hỏi tra mạng còn tối đa500 ký tự.'); return
        if cooperation['media_tool'] and not CHAT_MODELS[model]['tools']:
            compatible=[name for name in self.models if name in CHAT_MODELS and CHAT_MODELS[name]['tools']]
            preferred=next((name for name in ('qwen2.5:3b','qwen2.5:7b',self.cfg['default_model']) if name in compatible),None)
            if not preferred:
                QMessageBox.information(self,'Phối hợp AI','Cần AI hỗ trợ gọi công cụ để điều phối tạo ảnh. Hãy tải Qwen 3B hoặc Qwen 7B trước.');return
            self.select_ai(preferred);model=preferred
        import re
        mode_now=self.chat_mode.currentIndex()
        basic_video=mode_now==3 and (bool(self.pending_documents or self.pending_image) or bool(re.search(r'(ghép|slideshow|từ ảnh|ảnh có sẵn|từ các ảnh|từ file)',prompt,re.I)))
        needs_image_model=mode_now==2 or (mode_now==3 and not basic_video) or cooperation.get('media_tool')=='image_generate' or (cooperation.get('media_tool')=='video_generate' and not basic_video)
        if needs_image_model and not self.manager.enabled('media'):
            self.download('media'); return
        if len(prompt) > 6000:
            QMessageBox.information(self, 'Tin nhắn quá dài', 'Tin nhắn tối đa 6000 ký tự. Nội dung của bạn vẫn được giữ lại.'); return
        previous=self.store.load(self.cid).get('account_username')
        owner=self.server_session['username'] if self.server_session else GUEST_OWNER
        if previous and previous!=owner:
            self.cid=self.store.create(persist=False);self.render()
        self.store.remember_conversation(owner,self.cid)
        self.sent_prompt = prompt; self.input.clear(); self.status.setText('Đang chuẩn bị yêu cầu…')
        self.chat_task(prompt=prompt)

    def online_windows_task(self,prompt=None,allowed=None,expected=None,recover=False):
        from assistant.cloud import ServerApiClient
        if self.busy():return
        if not self.server_session:
            self.login_dialog();return
        state=self.store.load(self.cid)
        model=self.model.currentText() if prompt is not None else state['model']
        provider=REMOTE_MODELS.get(model)
        if provider=='cloudflare' or not provider:
            QMessageBox.information(self,'Điều khiển ứng dụng','Chọn DeepSeek API, NVIDIA hoặc Gemini. Cloudflare chưa hỗ trợ lập kế hoạch điều khiển app.');return
        if prompt is not None and len(prompt)>6000:
            QMessageBox.information(self,'Tin nhắn quá dài','Tin nhắn tối đa 6000 ký tự.');return
        if prompt is not None:self.begin_app_request()
        session=dict(self.server_session);cid=self.cid;cfg=dict(self.cfg)
        attachment_paths=list(self.pending_documents) if prompt is not None else []
        attachment_image=self.pending_image if prompt is not None else None
        if state.get('account_username') not in (None,session['username']):
            self.cid=self.store.create(persist=False);cid=self.cid
        if prompt is not None:
            self.sent_prompt=prompt;self.input.clear()
        self.answer_actions.hide();self.reply_active=True;self.reply_frame=0
        self.reply_logo.show();self.reply_dots.show();self.reply_timer.start()
        self.status.setText('AI trực tuyến đang chuẩn bị điều khiển app trên máy Windows…')
        def task(emit):
            from assistant.accounts import request_account
            from assistant.browser import BrowserTools
            from assistant.windows_apps import WindowsApps
            from assistant.online_automation import OnlineAutomation
            from assistant.files import FileTools
            from assistant.pdf_source import PDFSource
            from assistant.online_documents import online_pdf_reader
            from assistant.online_documents import online_pdf_reader
            from assistant.word_app import WordApp
            from assistant.cad_app import CadApp
            from assistant.cad3d_app import Cad3DApp
            from assistant.cloud import cancellable_request
            cancellable_request(lambda:request_account(session['endpoint'],'/api/models',{'username':session['username'],'key':session['key']},timeout=8),self.worker.stop_requested)
            with execution_lock(ROOT/'data/agent.lock'):
                state=self.store.load(cid)
                if state.get('account_username') not in (None,session['username']):raise RuntimeError('Hội thoại thuộc tài khoản khác.')
                audit=lambda action,details:self.store.audit(cid,action,details)
                from assistant.automation_setup import ensure_dependencies
                ensure_dependencies(cfg,ROOT/'config.json',on_status=lambda text:emit({'type':'status','text':text}),audit=audit)
                client=ServerApiClient(session,provider,on_status=lambda text:emit({'type':'status','text':text}),cancel_event=self.worker.stop_requested)
                windows=WindowsApps(cfg,audit,owner=session['username'],policy_path=ROOT/'config.json')
                files=FileTools(cfg['roots'],ROOT/'data/backups',audit)
                from assistant.cdm_layout import CdmLayoutApp
                from assistant.cad_tracdoc import CadTracDocApp
                from assistant.plaxis_app import PlaxisApp
                from assistant.plaxis_remote import PlaxisRemoteApp
                from assistant.geoslope_inspect import GeoslopeInspect
                _plaxis_app=PlaxisApp(files,audit)
                agent=OnlineAutomation(client,cfg,self.store,cid,windows,
                    BrowserTools(cfg,audit,policy_path=ROOT/'config.json',on_status=lambda text:emit({'type':'status','text':text})),
                    PDFSource(windows,files,audit,pdf_ocr=online_pdf_reader(cfg,session,cancel_event=self.worker.stop_requested,on_status=lambda text:emit({'type':'status','text':text}))),WordApp(windows,files),CadApp(windows,files,audit),Cad3DApp(windows,files,audit),
                    CdmLayoutApp(windows,files,audit),CadTracDocApp(windows,files,audit),
                    plaxis_app=_plaxis_app,plaxis_remote=PlaxisRemoteApp(_plaxis_app,on_status=lambda text:emit({'type':'status','text':text})),
                    geoslope_inspector=GeoslopeInspect(files))
                if prompt is not None:
                    attachments=[]
                    for raw in attachment_paths:
                        source=Path(raw)
                        if not source.is_file() or source.stat().st_size>(20 if source.suffix.lower()=='.dxf' else 10)*1024*1024:raise ValueError('Tệp đính kèm không còn hợp lệ: '+source.name)
                        if not files.roots:raise PermissionError('Thêm thư mục được phép để xử lý tài liệu đính kèm.')
                        import uuid
                        destination=files.path(str(files.roots[0]/('attachment-'+uuid.uuid4().hex+source.suffix)),exists=False)
                        with destination.open('xb') as output:output.write(source.read_bytes())
                        attachments.append({'name':source.name,'path':str(destination)})
                        audit('automation_attachment_staged',{'name':source.name,'path':str(destination)})
                    state['automation_attachments']=attachments
                    from assistant.automation_start import start_automation
                    start_automation(agent,state,prompt,model,session['username'],image=attachment_image)
                    state['messages'][-1]['documents']=attachments
                    agent.save(state)
                    self.store.remember_conversation(session['username'],cid)
                    emit({'type':'sent','cid':cid,'prompt':prompt})
                elif recover:
                    if not state.get('pending') or not state['pending'].get('decision_started'):raise RuntimeError('Không có thao tác cần phục hồi.')
                    state['messages'].append({'role':'assistant','content':'Thao tác trước bị ngắt; kết quả chưa rõ. Không tự thực hiện lại.'})
                    state.update(running=False,pending=None,queue=[]);agent.save(state)
                elif allowed is not None:
                    if allowed:emit({'type':'app_activity','text':'Đang thực hiện: '+state['pending']['plan']['action']})
                    agent.approve(state,allowed,expected)
                def snapshot():
                    return [{'role':m['role'],'content':m.get('content',''),'images':m.get('images',[]), 'documents':m.get('documents',[]),'source_index':i}
                            for i,m in enumerate(state['messages']) if m['role'] in ('user','assistant') and m.get('content')]
                emit({'type':'snapshot','messages':snapshot()})
                try:
                    for event in agent.run(state):
                        if self.worker.stop_requested.is_set():
                            state['running']=False;agent.save(state);break
                        emit(event)
                except Exception:
                    state['running']=False;agent.save(state);raise
                emit({'type':'snapshot','messages':snapshot()})
                return state
        self.work(task,self.after_chat,cancellable=True,app_countdown=prompt is not None)

    def chat_task(self, prompt=None, allowed=None, expected=None, recover=False):
        if prompt is None and self.store.load(self.cid).get('online_automation'):
            self.online_windows_task(allowed=allowed,expected=expected,recover=recover);return
        if not self.server_session:
            if (prompt is not None and self.trial.remaining()==0) or (prompt is None and not self.trial.can_continue(self.cid,self.store.load(self.cid))):
                self.login_dialog();return
        self.answer_actions.hide()
        self.reply_active=True; self.reply_frame=0; self.reply_logo.show();self.reply_dots.show(); self.reply_timer.start()
        cid, model = self.cid, self.model.currentText()
        requested_mode = self.chat_mode.currentIndex()
        images=[self.pending_image] if prompt is not None and self.pending_image and requested_mode!=3 else None
        media_paths=list(self.pending_media_paths) if prompt is not None and requested_mode==3 else []
        session=dict(self.server_session) if self.server_session else None
        cached_memories=list(self.personal_memories)
        attached_paths=list(self.pending_documents) if prompt is not None else []
        from assistant.collaboration import collaboration_intent
        requested_tools = requested_mode in (1,2,3,5) or bool(CHAT_MODELS.get(model,{}).get("tools"))
        requested_web=requested_mode==4 or (requested_mode==5 and self.web_btn.isChecked())
        requested_deep_analysis=bool(self.deep_analysis_enabled) if prompt is not None else None
        requested_expert=self.expert_choice.currentData() if requested_mode==5 else None
        def task(emit):
            from assistant.agent import Agent
            from assistant.accounts import request_account
            import time as _time
            if session:
                _cache = getattr(self, '_auth_check_cache', {})
                _key = (session['endpoint'], session['username'], session['key'])
                if _time.monotonic() - _cache.get(_key, 0) > 60:
                    try:
                        request_account(session['endpoint'],'/api/models',{'username':session['username'],'key':session['key']},timeout=8)
                        _cache[_key] = _time.monotonic()
                        self._auth_check_cache = _cache
                    except Exception:
                        _cache.pop(_key, None)
                        emit({'type':'auth_failed'});raise RuntimeError('Không xác thực được tài khoản với server. Đăng nhập lại hoặc kiểm tra mạng.') from None
            with execution_lock(ROOT / 'data/agent.lock'):
                state = self.store.load(cid)
                owner=session['username'] if session else GUEST_OWNER
                if state.get('account_username') and state['account_username']!=owner:
                    raise RuntimeError('Hội thoại thuộc tài khoản khác. Hãy tạo cuộc trò chuyện mới.')
                audit = lambda action, details: self.store.audit(cid, action, details)
                if self.cfg.get('windows_apps_enabled'):
                    from assistant.automation_setup import ensure_dependencies
                    ensure_dependencies(self.cfg,ROOT/'config.json',on_status=lambda text:emit({'type':'status','text':text}),audit=audit)
                use_tools = requested_tools if prompt is not None else state.get('tools_enabled', True)
                excel, caps = None, None
                if use_tools:
                    from assistant.excel import ExcelTools
                    from assistant.capabilities import Capabilities
                    excel = ExcelTools(self.cfg['roots'], ROOT / 'data/backups', audit)
                    caps = Capabilities(self.manager, excel, self.client, self.cfg, ROOT, audit, owner=owner,
                        storage_root=Path(self.store.path).parent,
                        allow_web=requested_web if prompt is not None else state.get('ui_mode') in (4,5) and bool(state.get('web_search_requested')),
                        on_status=lambda message:self.worker.event.emit({'type':'status','text':message}))
                mode = requested_mode if prompt is not None else state.get('ui_mode', 1 if use_tools else 0)
                if mode in (2, 3):
                    if mode==2:
                        caps.schemas = [schema for schema in caps.schemas if schema['function']['name']=='image_generate']
                        if not caps.schemas:raise RuntimeError('Mô-đun Tạo ảnh AI nâng cao chưa sẵn sàng.')
                    else:
                        caps.schemas = [schema for schema in caps.schemas if schema['function']['name'] in {'video_from_images','video_generate'}]
                        if not caps.schemas:raise RuntimeError('Chưa có công cụ video cơ bản hoặc nâng cao.')
                agent = Agent(self.client, excel, self.cfg, self.store, cid, caps, tools_enabled=use_tools)
                agent.web_enabled=True  # Built-in Bing; actual permission remains gated by the web button.
                if session and self.manager.ready('rag'):
                    if caps:agent.document_search=caps.rag
                    else:
                        from assistant.files import FileTools
                        from assistant.rag import RagTools
                        agent.document_search=RagTools(FileTools(self.cfg['roots'],ROOT/'data/backups',audit),self.client,ROOT,audit,owner=session['username'],storage_root=Path(self.store.path).parent,vision_model=self.cfg.get('vision_model','gemma3:4b'))
                if session:
                    from assistant.memory import PersonalMemory
                    agent.memory=PersonalMemory(self.store,self.client,session['username'],self.manager.ready('rag'))
                agent.personal_memories=[{'title':m['title'],'text':m['text'][:400]} for m in cached_memories[:12]] if session else []
                if prompt is not None:
                    web_results=None
                    if not state.get('account_username'):state['account_username']=owner
                    state['video_source_paths']=media_paths
                    if session:agent.start(state, prompt, model, images=images,expert_mode=requested_mode==5,expert_override=requested_expert)
                    else:self.trial.start(agent,state,prompt,model,images=images,expert_mode=requested_mode==5,expert_override=requested_expert)
                    from assistant.message_attachments import attachment_records
                    state['messages'][-1]['documents']=attachment_records(attached_paths)
                    agent.save(state)
                    state['web_results']=web_results
                    state.pop('attached_documents',None)
                    if attached_paths:
                        if any(Path(path).suffix.lower()=='.pdf' for path in attached_paths):
                            emit({'type':'app_activity','text':'Đang đọc PDF; AI đọc ảnh trước, OCR dự phòng'})
                        emit({'type':'status','text':'Đang đọc tài liệu đính kèm…'})
                        state['attached_documents']=self.read_attachments(attached_paths,
                            progress=lambda message:emit({'type':'status','text':message}))
                        from assistant.code_files import source_record
                        user_index=len(state['messages'])-1
                        for attachment_number,path in enumerate(attached_paths):
                            record=source_record(path,with_text=True)
                            if record:
                                text=record.pop('_text')
                                state['attached_documents'][attachment_number]={'file':record['name'],'source':str(path),'text':text,'truncated':record['truncated'],'full_text':not record['truncated'],'format':'text'}
                                record['user_index']=user_index
                                state.setdefault('attachment_records',[]).append(record)
                        audit('attachment_read',{'files':[Path(p).name for p in attached_paths]})
                    agent.personal_memories=[]
                    state.pop('personal_memories',None)
                    if session:
                        from assistant.accounts import request_account
                        memory_synced=False
                        try:
                            memories=request_account(session['endpoint'],'/api/memory/personal/list',{'username':session['username'],'key':session['key']},timeout=8)['items']
                            memory_synced=True
                        except Exception:
                            memories=cached_memories
                            emit({'type':'status','text':'Chưa đồng bộ bộ nhớ; dùng bản đã nạp cho tài khoản này.'})
                        if memory_synced:agent.memory.sync_server(memories)
                        # Chọn ký ức trong Agent.run một lần, sau phân loại.
                        state['account_username']=session['username']
                    state['web_search_requested'] = requested_web
                    state['deep_analysis']=requested_deep_analysis
                    state['ui_mode'] = mode; state['media_done'] = False; agent.save(state)
                elif recover: agent.recover_uncertain(state)
                elif allowed is not None:
                    if state['pending'] != expected: raise RuntimeError('Preview đã thay đổi. Xin duyệt lại.')
                    action=state['pending']['plan'].get('action','')
                    if allowed and (action.startswith(('windows_','browser_','pdf_')) or action in {'word_create_open','cad_create_open','cad3d_create_open'}):
                        emit({'type':'app_activity','text':'Đang thực hiện: '+action})
                    agent.approve(state, allowed)
                def snapshot():
                    return [{'role': m['role'], 'content': m.get('content',''), 'images': m.get('images',[]), 'documents':m.get('documents',[]),
                             'media':m.get('media',[]),'source_index':i} for i,m in enumerate(state['messages'])
                            if m['role'] in ('user', 'assistant') and (m.get('content') or m.get('media'))]
                if prompt is not None:emit({'type':'sent','cid':cid,'prompt':prompt})
                emit({'type': 'snapshot', 'messages': snapshot()})
                events=agent.run(state)
                try:
                    for event in events:
                        if event['type'] == 'status' and event['text'].startswith('Đang trả lời'):
                            emit({'type': 'snapshot', 'messages': snapshot()})
                        emit(event)
                finally:
                    events.close()
                return self.store.load(cid)
        self.work(task, self.after_chat, cancellable=True,app_countdown=prompt is not None)

    def after_chat(self, state):
        if state['pending']:
            self.approve_pending(state['pending'])
        else:
            perf = state.get('performance', {})
            count, duration = perf.get('eval_count') or 0, perf.get('eval_duration') or 0
            load = (perf.get('load_duration') or 0) / 1e9
            text = f'Hoàn tất · {count/(duration/1e9):.1f} token/s · nạp model {load:.1f}s' if count and duration else 'Hoàn tất'
            if not self.server_session:
                remaining = self.trial.remaining()
                text += ' · Đăng nhập để hỏi tiếp' if not remaining and self.model.currentText()!=CLOUD_MODEL else ''
            self.status.setText(text)

    def resume(self):
        state = self.store.load(self.cid)
        if state.get('online_automation') and state.get('running'):
            if state.get('pending') and not (self.cfg.get('windows_apps_auto_execute') and not state['pending'].get('decision_started')):self.approve_pending(state['pending'])
            else:self.online_windows_task()
            return
        if state.get('model') in REMOTE_MODELS and state['running']:
            state['running']=False;self.store.save(self.cid,state);self.render()
            self.status.setText('Luồng AI trên server đã bị ngắt. Nội dung đã lưu được giữ; bạn có thể gửi câu hỏi mới.');return
        action=(state.get('pending') or {}).get('plan',{}).get('action','')
        automatic_app=self.cfg.get('windows_apps_auto_execute') and (action.startswith(('windows_','browser_','pdf_')) or action in {'word_create_open','cad_create_open','cad3d_create_open'}) and not (state.get('pending') or {}).get('decision_started')
        if state['pending'] and not automatic_app: self.approve_pending(state['pending'])
        elif state['running']: self.chat_task()

    def approve_pending(self, pending):
        action=pending['plan'].get('action','')
        app_action=action in {'windows_list_apps','windows_open','windows_inspect','windows_action','browser_search','browser_run','pdf_source_open','pdf_read','word_create_open','cad_create_open','cad3d_create_open'}
        if app_action and self.cfg.get('windows_apps_enabled') and self.cfg.get('windows_apps_auto_execute'):
            if pending.get('decision_started'):
                self.status.setText('Thao tác trước chưa rõ kết quả; không chạy lại. Đang ghi nhận trạng thái.')
                self.chat_task(recover=True)
            else:self.chat_task(allowed=True,expected=pending)
            return
        if pending.get('decision_started'):
            if QMessageBox.question(self, 'Phục hồi', 'Lượt trước bị ngắt trong khi ghi. Đánh dấu chưa rõ kết quả và tiếp tục? Không chạy lại thao tác.') == QMessageBox.StandardButton.Yes:
                self.chat_task(recover=True)
            return
        dialog = QDialog(self); dialog.setWindowTitle('Xác nhận thao tác'); dialog.resize(800, 620)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel('Kiểm tra đường dẫn, nội dung và tác động trước khi đồng ý.'))
        remember=None
        if app_action:
            remember=QCheckBox('Ghi nhớ quyền điều khiển app, không hỏi lại sau mỗi bước hoặc lỗi')
            remember.setChecked(True);layout.addWidget(remember)
        preview = QPlainTextEdit(); preview.setReadOnly(True)
        preview.setPlainText(json.dumps(pending['plan'], ensure_ascii=False, indent=2, default=str)); layout.addWidget(preview)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Yes | QDialogButtonBox.StandardButton.No)
        buttons.button(QDialogButtonBox.StandardButton.Yes).setText('Đồng ý thực hiện')
        buttons.button(QDialogButtonBox.StandardButton.No).setText('Từ chối')
        buttons.accepted.connect(dialog.accept); buttons.rejected.connect(dialog.reject); layout.addWidget(buttons)
        allowed = dialog.exec() == QDialog.DialogCode.Accepted
        if allowed and remember is not None and remember.isChecked():
            from assistant.config import remember_app_permission
            remember_app_permission()
            self.cfg['windows_apps_auto_execute']=True
            if hasattr(self,'windows_auto_execute_check'):
                self.windows_auto_execute_check.setChecked(True)
        self.chat_task(allowed=allowed, expected=pending)

    def download(self, target):
        if target=='web':
            self.tabs.setCurrentIndex(0);self.chat_mode.setCurrentIndex(4);self.status.setText('Tìm kiếm Bing sẵn sàng, không cần tải module.');return
        note = MODULES[target]['note'] if target in MODULES else 'Tải model Ollama về máy; cần Internet và dung lượng đĩa.'
        title={'media':'AI tạo ảnh nâng cao · '+MEDIA_VARIANTS[self.manager.media_variant()]['display'],'media_basic':'Thư viện ảnh / video cơ bản'}.get(target,target)
        if target=='media':
            selected=MEDIA_VARIANTS[self.manager.media_variant()]
            note+='\nModel chọn: '+selected['display']+'\n'+selected['quality_note']+'\nDung lượng trống khuyến nghị: '+str(selected['disk_gb'])+' GiB.'
        if target in CHAT_MODELS:
            spec=CHAT_MODELS[target]
            note+='\n'+spec['label']+'\nCần ít nhất '+str(spec['disk_gb'])+' GiB trống trên ổ dự án; kiểm tra thêm ổ chứa AI của Ollama.'
            if spec.get('large'):note+='\n\nLưu ý cấu hình: '+spec['warning']
        if target in ALLOWED_MODELS or (target in MODULES and MODULES[target]['models']):
            note+='\nNếu chưa có Ollama, ứng dụng sẽ tự tải, cài và khởi động Ollama trước khi tải AI.'
        if QMessageBox.question(self, 'Bật / tải ' + title, note + '\n\nBạn có muốn tải và bật dưới nền?') != QMessageBox.StandardButton.Yes: return
        try:
            self.manager.request(target); self.tabs.setCurrentIndex(1); self.poll()
        except Exception as exc: QMessageBox.warning(self, 'Tải module', str(exc))

    def remove_ai(self, name):
        if self.busy() or self.manager.busy:
            QMessageBox.information(self,'Gỡ AI','Đợi tác vụ tải/cài hiện tại kết thúc trước khi gỡ AI.');return
        if QMessageBox.question(self,'Gỡ AI',
            'Gỡ '+name+' khỏi Ollama?\nBạn sẽ cần tải lại để dùng AI này. Lịch sử chat và tài liệu vẫn được giữ nguyên.',
            QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)!=QMessageBox.StandardButton.Yes:return
        try:
            self.manager.remove_model(name);self.jobs_loaded=False;self.tabs.setCurrentIndex(1);self.poll()
        except Exception as error:QMessageBox.warning(self,'Gỡ AI',str(error))

    def control_download(self, action):
        try:
            if action=='resume':self.manager.resume_download()
            else:
                self.manager.stop_download(action)
                self.download_status.setText('Đang ngắt tải tại điểm an toàn…')
            self.jobs_loaded=False;self.poll()
        except Exception as error:QMessageBox.warning(self,'Tải AI',str(error))

    def disable(self, key):
        if self.busy() or getattr(self,'update_worker',None) is not None:
            QMessageBox.information(self, 'Chat AI', 'Đợi tác vụ hiện tại kết thúc trước khi đổi module.'); return
        self.manager.set_enabled(key, False); self.status.setText('Đã tắt ' + key)

    def build_settings(self, layout):
        group = QGroupBox('Cấu hình máy và AI'); group_layout = QVBoxLayout(group)
        choice_form=QFormLayout();group_layout.addLayout(choice_form)
        local_group=QGroupBox('AI trên máy · Ollama');form=QFormLayout(local_group)
        local_note=QLabel('Các model và cấu hình phần cứng dưới đây chỉ dành cho AI chạy trên máy.');local_note.setWordWrap(True);form.addRow(local_note)
        group_layout.addWidget(local_group)
        self.settings_fields = {}
        from assistant.hardware_profile import LABELS
        self.hardware_info={}
        self.machine_profile=QComboBox()
        for key,label in LABELS.items():self.machine_profile.addItem(label,key)
        self.machine_profile.setCurrentIndex(self.machine_profile.findData(self.cfg.get('machine_profile','medium')))
        form.addRow('Cấu hình máy',self.machine_profile)
        self.machine_auto=QCheckBox('Tự chọn AI đã tải theo cấu hình máy');self.machine_auto.setChecked(self.cfg.get('machine_auto_ai',True));form.addRow(self.machine_auto)
        self.hardware_summary=QLabel('Nhận diện máy để xem RAM, CPU, GPU và VRAM.');self.hardware_summary.setWordWrap(True);form.addRow(self.hardware_summary)
        self.hardware_advice=QLabel();self.hardware_advice.setWordWrap(True);form.addRow(self.hardware_advice)
        self.button(form,'Nhận diện máy',self.detect_machine)
        self.button(form,'Tải AI đề xuất',self.download_recommended_ai)
        self.button(form,'Trạng thái GPU / Ollama',self.check_gpu)
        self.settings_provider=QComboBox();self.settings_provider.addItems([*REMOTE_MODELS,'AI trên máy']);self.settings_provider.setCurrentText(PROVIDER_NAMES.get(self.cfg.get('chat_provider'),'AI trên máy'));choice_form.addRow('AI mặc định',self.settings_provider)
        for key, label in [('default_model','AI trò chuyện'),('code_model','AI lập trình'),('vision_model','AI đọc ảnh')]:
            field = QComboBox(); field.addItems([n for n,v in CHAT_MODELS.items() if key!='vision_model' or v.get('vision')]); field.setCurrentText(self.cfg.get(key,'gemma3:4b'))
            self.settings_fields[key] = field; form.addRow(label,field)
        for key,label,low,high,default in [('num_ctx','Ngữ cảnh Ollama',1024,32768,4096),('num_predict','Token trả lời trên máy',128,8192,2048)]:
            field = QSpinBox(); field.setRange(low,high); field.setValue(self.cfg.get(key,default))
            self.settings_fields[key]=field; form.addRow(label,field)
        field = QDoubleSpinBox(); field.setRange(0,2); field.setSingleStep(.1); field.setValue(self.cfg['temperature'])
        self.settings_fields['temperature']=field; form.addRow('Độ sáng tạo trên máy',field)
        online_group=QGroupBox('AI trực tuyến · API');online_form=QFormLayout(online_group)
        online_note=QLabel('Không cần tải model hoặc cài Ollama. Flash dành cho chat nhanh; Pro/Suy luận dành cho phân tích phức tạp. Cần Internet và key trên server.');online_note.setWordWrap(True);online_form.addRow(online_note)
        field=QSpinBox();field.setRange(128,4096);field.setValue(self.cfg.get('api_num_predict',4096))
        field.setToolTip('Giới hạn câu trả lời API, độc lập với cấu hình máy và nút cấu hình nhẹ. Server hiện giới hạn tối đa 4096 token.')
        self.settings_fields['api_num_predict']=field;online_form.addRow('Token trả lời trực tuyến',field)
        field=QDoubleSpinBox();field.setRange(0,1);field.setSingleStep(.1);field.setValue(self.cfg.get('api_temperature',.2))
        self.settings_fields['api_temperature']=field;online_form.addRow('Độ sáng tạo trực tuyến',field)
        online_form.addRow(QLabel('Ngữ cảnh API do dịch vụ và dữ liệu gửi quyết định; không dùng ô Ngữ cảnh Ollama.'))
        group_layout.addWidget(online_group)
        network_group=QGroupBox('Công cụ trực tuyến');network_form=QFormLayout(network_group)
        self.online_tools_check=QCheckBox('Cho phép dùng công cụ API trực tuyến');self.online_tools_check.setChecked(self.cfg.get('online_tools_enabled',False))
        self.online_upload_check=QCheckBox('Cho phép gửi ảnh trang tài liệu tới dịch vụ AI');self.online_upload_check.setChecked(self.cfg.get('online_document_upload',False))
        network_form.addRow(self.online_tools_check);network_form.addRow(self.online_upload_check)
        field=QComboBox();field.addItems(['deepseek_flash']);field.setCurrentText(self.cfg.get('online_document_provider','deepseek_flash'))
        self.settings_fields['online_document_provider']=field;network_form.addRow('API đọc trang PDF',field)
        field=QSpinBox();field.setRange(0,1000000);field.setSpecialValueText('Không giới hạn');field.setValue(self.cfg.get('online_document_pages',0))
        self.settings_fields['online_document_pages']=field;network_form.addRow('Số trang nhận dạng tối đa',field)
        note=QLabel('Đọc chữ trong PDF trước, chỉ gửi ảnh trang thiếu chữ hoặc lỗi mã hóa khi cả hai quyền được bật. Quyền được lưu một lần. Chỉ dùng DeepSeek Flash với key DeepSeek trên server; không chạy Foxit OCR. Đặt số trang bằng 0 để đọc lần lượt không giới hạn tổng số trang. Có thể phát sinh phí API. Trang chưa đọc sẽ được báo rõ; công cụ trên máy vẫn cần cài đặt.');note.setWordWrap(True);network_form.addRow(note)
        group_layout.addWidget(network_group)
        tools_group=QGroupBox('Công cụ trên máy');tools_form=QFormLayout(tools_group)
        self.ai_tools_auto_check=QCheckBox('AI tự thực hiện các công cụ trong phạm vi đã cấp quyền, không hỏi lại từng bước');self.ai_tools_auto_check.setChecked(self.cfg.get('ai_tools_auto_execute',False))
        tools_form.addRow(self.ai_tools_auto_check)
        self.auto_python_check=QCheckBox('AI tự viết/chạy Python tra cứu trong Docker');self.auto_python_check.setChecked(self.cfg.get('auto_python',True))
        tools_form.addRow(self.auto_python_check)
        self.button(tools_form,'Dùng công cụ AI tự động trong chat',lambda:(self.tabs.setCurrentIndex(0),self.chat_mode.setCurrentIndex(1),self.input.setFocus()))
        field=QSpinBox();field.setRange(1,20);field.setValue(self.cfg.get('max_rounds',8))
        self.settings_fields['max_rounds']=field;tools_form.addRow('Số vòng Agent trên máy',field)
        tools_note=QLabel('Điều khiển Word, Excel, CAD và chạy Python vẫn cần công cụ cài trên máy, kể cả khi AI lập kế hoạch qua API.');tools_note.setWordWrap(True);tools_form.addRow(tools_note)
        group_layout.addWidget(tools_group)
        layout.addWidget(group)
        self.api_key_group=QGroupBox('Key AI trực tuyến · Quản trị viên');api_form=QFormLayout(self.api_key_group);self.api_key_fields={};self.api_model_fields={}
        api_form.addRow(QLabel('Danh sách chọn nhanh có Cloudflare, NVIDIA, DeepSeek, Gemini và Groq. Cloudflare dùng binding Workers; các dịch vụ còn lại dùng key quản trị viên.'))
        for provider,label in [('nvidia','NVIDIA'),('deepseek','DeepSeek'),('gemini','Gemini'),('groq','Groq')]:
            field=QLineEdit();field.setEchoMode(QLineEdit.EchoMode.Password);field.setPlaceholderText('Nhập key mới; để trống để giữ key trên server');self.api_key_fields[provider]=field;api_form.addRow(label,field)
            model_field=QLineEdit(self.cfg.get(provider+'_model',''));self.api_model_fields[provider]=model_field;api_form.addRow('Mã AI '+label,model_field)
            if provider in ('nvidia','deepseek','groq'):
                self.button(api_form,'Chọn AI từ '+label,lambda checked=False,p=provider:self.load_provider_catalog(True,p))
                field.editingFinished.connect(lambda p=provider:QTimer.singleShot(200,lambda:self.load_provider_catalog(False,p)))
            self.button(api_form,'Kiểm tra '+label,lambda checked=False,p=provider:self.check_online_provider(p))
        self.button(api_form,'Thêm AI',self.add_ai_dialog)
        self.button(api_form,'Cấu hình 3 AI DeepSeek dùng chung key',self.configure_deepseek_presets)
        self.button(api_form,'Lưu key API',self.save_settings)
        self.button(api_form,'Xem trạng thái key trên server',self.provider_key_status)
        self.api_key_group.setVisible(admin_session(self.server_session));layout.addWidget(self.api_key_group)
        self.machine_profile.currentIndexChanged.connect(self.refresh_machine_choices)
        self.machine_auto.toggled.connect(self.refresh_machine_choices)
        self.refresh_machine_choices()
        self.button(form,'Cấu hình nhẹ: Qwen 3B / 2048 context',self.light_settings)
        appearance=QGroupBox('Giao diện');appearance_form=QFormLayout(appearance)
        self.settings_theme=QComboBox();self.settings_theme.addItem('Sáng','light');self.settings_theme.addItem('Tối','dark');self.settings_theme.setCurrentIndex(self.settings_theme.findData(self.cfg.get('theme','dark')))
        appearance_form.addRow('Chế độ hiển thị',self.settings_theme)
        font=QSpinBox();font.setRange(12,22);font.setValue(self.cfg.get('font_size',13));self.settings_fields['font_size']=font;appearance_form.addRow('Cỡ chữ chat',font)
        appearance_form.addRow(QLabel('Xem trước ngay. Bấm Lưu cài đặt để ghi nhớ cho lần mở sau.'))
        self.settings_theme.currentIndexChanged.connect(self.preview_appearance);font.valueChanged.connect(self.preview_appearance)
        layout.addWidget(appearance)
        files = QGroupBox('Office và thư mục được phép'); box = QVBoxLayout(files)
        box.addWidget(QLabel('Excel, Word, PowerPoint. Mỗi dòng là một thư mục whitelist.'))
        self.settings_roots = QPlainTextEdit('\n'.join(self.cfg['whitelist'])); self.settings_roots.setMaximumHeight(100); box.addWidget(self.settings_roots)
        box.addWidget(QLabel('Ghi/sửa/xóa/chạy lệnh cần xác nhận; sửa file có backup.'))
        self.button(box,'Dùng công cụ Office trong chat',lambda:(self.tabs.setCurrentIndex(0),self.chat_mode.setCurrentIndex(1),self.input.setFocus()))
        layout.addWidget(files)
        automation = QGroupBox('Công cụ AI · Điều khiển ứng dụng Windows'); app_box = QVBoxLayout(automation)
        self.windows_apps_check = QCheckBox('Cho AI mở và điều khiển các ứng dụng được phép')
        self.windows_apps_check.setChecked(self.cfg.get('windows_apps_enabled',False));app_box.addWidget(self.windows_apps_check)
        self.automation_auto_install_check=QCheckBox('Tự cài thư viện cần thiết đã được kiểm tra')
        self.automation_auto_install_check.setChecked(self.cfg.get('automation_auto_install',False));app_box.addWidget(self.automation_auto_install_check)
        self.windows_all_apps_check=QCheckBox('Cho phép mở mọi ứng dụng đã cài')
        self.windows_all_apps_check.setChecked(self.cfg.get('windows_apps_all_installed',False));app_box.addWidget(self.windows_all_apps_check)
        self.windows_compact_check=QCheckBox('Tự thu gọn chat khi AI điều khiển ứng dụng')
        self.windows_compact_check.setChecked(self.cfg.get('windows_apps_compact',True));app_box.addWidget(self.windows_compact_check)
        self.windows_auto_execute_check=QCheckBox('Tự thực hiện yêu cầu điều khiển app, không hỏi lại từng bước')
        self.windows_auto_execute_check.setChecked(self.cfg.get('windows_apps_auto_execute',False));app_box.addWidget(self.windows_auto_execute_check)
        permission_note=QLabel('Tự cài chỉ áp dụng pywinauto, psutil, comtypes, Playwright, pypdf, python-docx, ezdxf và mapbox-earcut trong Python riêng của ChatAI. Mở mọi app dùng danh sách đăng ký Windows; không tự cấp quyền quản trị. Quyền tự thực hiện áp dụng các công cụ điều khiển app khi được bật.')
        permission_note.setWordWrap(True);app_box.addWidget(permission_note)
        self.browser_background_check=QCheckBox('Chrome chạy nền (không hiện cửa sổ)')
        self.browser_background_check.setChecked(self.cfg.get('browser_background',False));app_box.addWidget(self.browser_background_check)
        note=QLabel('Nếu bật tự thực hiện, chỉ cấp quyền một lần; nút Dừng ngắt các bước tiếp theo. Cửa sổ có thể hiện; chỉ hỗ trợ app UI Automation. Không tự sao lưu dữ liệu của app bên ngoài.')
        note.setWordWrap(True);app_box.addWidget(note)
        self.windows_apps_paths = QPlainTextEdit('\n'.join(self.cfg.get('windows_apps_allowed',[])))
        self.windows_apps_paths.setPlaceholderText('Mỗi dòng một đường dẫn EXE được phép. Dùng nút Thêm ứng dụng để chọn.')
        self.windows_apps_paths.setMaximumHeight(100);app_box.addWidget(self.windows_apps_paths)
        self.button(app_box,'Thêm ứng dụng EXE…',self.add_windows_app)
        self.button(app_box,'Dừng điều khiển app',self.stop_windows_apps)
        self.button(app_box,'Tiếp tục điều khiển app',self.resume_windows_apps)
        from assistant.windows_apps import readiness
        dependency=QLabel(readiness(self.cfg))
        dependency.setWordWrap(True);app_box.addWidget(dependency)
        self.windows_readiness_label=dependency
        self.windows_apps_check.clicked.connect(lambda enabled:self.stop_windows_apps() if not enabled else None)
        layout.addWidget(automation)
        history = QGroupBox('Lịch sử và dữ liệu'); box = QVBoxLayout(history)
        self.button(box,'Xem toàn bộ cuộc trò chuyện',self.full_history)
        self.button(box,'Xóa cuộc trò chuyện đang chọn',self.delete_chat)
        self.button(box,'Mở thư mục backup',lambda:self.open_path(ROOT/'data/backups'))
        self.button(box,'Xem nhật ký thao tác',self.show_audit)
        layout.addWidget(history)
        appearance_form.addRow(QLabel('Logo robot · Enter gửi · Shift+Enter xuống dòng'))
        account = QGroupBox('Thông tin tài khoản'); box = QVBoxLayout(account)
        # Keep connection configuration in the backend; it is not a visible account field.
        self.settings_server = QLineEdit(self.cfg.get('server_url',''),self);self.settings_server.hide()
        self.account_status=QLabel('Chưa đăng nhập.');box.addWidget(self.account_status)
        self.account_name_label=QLabel();self.account_name_label.setWordWrap(True);box.addWidget(self.account_name_label)
        self.account_username_label=QLabel();box.addWidget(self.account_username_label)
        self.account_email_label=QLabel();self.account_email_label.setWordWrap(True);box.addWidget(self.account_email_label)
        self.account_phone_label=QLabel();box.addWidget(self.account_phone_label)
        for label in (self.account_name_label,self.account_username_label,self.account_email_label,self.account_phone_label):label.setTextFormat(Qt.TextFormat.PlainText)
        self.account_register_button=self.button(box,'Đăng ký',self.register_dialog)
        self.account_login_button=self.button(box,'Đăng nhập',self.login_dialog)
        self.account_remember=QCheckBox('Ghi nhớ đăng nhập');self.account_remember.setChecked(True);box.addWidget(self.account_remember)
        self.account_edit_button=self.button(box,'Chỉnh sửa thông tin cá nhân',self.edit_personal_info)
        self.account_password_button=self.button(box,'Đổi mật khẩu',self.change_password_dialog)
        self.account_logout_button=self.button(box,'Đăng xuất',self.logout_account)
        layout.addWidget(account)
        updates=QGroupBox('Cập nhật Chat AI'); box=QVBoxLayout(updates)
        box.addWidget(QLabel('Phiên bản hiện tại: 2.6.6 · GitHub vuanh97nd/ChatAI'))
        self.update_status=QLabel('Chưa kiểm tra cập nhật.'); self.update_status.setWordWrap(True); box.addWidget(self.update_status)
        self.update_notes=QPlainTextEdit(); self.update_notes.setReadOnly(True); self.update_notes.setMaximumHeight(130); box.addWidget(self.update_notes)
        self.update_check=self.button(box,'Kiểm tra cập nhật',self.check_updates)
        self.update_get=self.button(box,'Tải bản cập nhật',self.download_update); self.update_get.setEnabled(False)
        self.button(box,'Mở trang phát hành GitHub',lambda:QDesktopServices.openUrl(QUrl('https://github.com/vuanh97nd/ChatAI/releases')))
        layout.addWidget(updates)
        advanced=QGroupBox('Nâng cao');advanced_layout=QVBoxLayout(advanced)
        self.button(advanced_layout,'Mở config.json',lambda:self.open_path(ROOT/'config.json'))
        self.button(advanced_layout,'Mở thư mục log và backup',lambda:self.open_path(ROOT/'data'))
        layout.addWidget(advanced)
        self.settings_sections={'Giao diện':appearance,'Nâng cao':advanced,'Cấu hình máy và AI':group,'Office và thư mục':files,'Điều khiển ứng dụng':automation,'Lịch sử và dữ liệu':history,'Tài khoản':account,'Cập nhật Chat AI':updates}
        self.apply_settings_button=self.button(layout,'Lưu cài đặt',self.save_settings)
        self.apply_settings_button.setEnabled(False)
        for field in self.settings_fields.values():
            (field.currentTextChanged if isinstance(field,QComboBox) else field.valueChanged).connect(self.account_settings_changed)
        self.settings_theme.currentIndexChanged.connect(self.account_settings_changed)
        self.settings_provider.currentIndexChanged.connect(self.account_settings_changed)
        self.auto_python_check.toggled.connect(self.account_settings_changed)
        self.online_tools_check.toggled.connect(self.account_settings_changed)
        self.online_upload_check.toggled.connect(self.account_settings_changed)
        self.settings_roots.textChanged.connect(self.account_settings_changed)
        self.windows_apps_check.toggled.connect(self.account_settings_changed)
        self.windows_compact_check.toggled.connect(self.account_settings_changed)
        self.windows_auto_execute_check.toggled.connect(self.account_settings_changed)
        self.automation_auto_install_check.toggled.connect(self.account_settings_changed)
        self.windows_all_apps_check.toggled.connect(self.account_settings_changed)
        self.browser_background_check.toggled.connect(self.account_settings_changed)
        self.windows_apps_paths.textChanged.connect(self.account_settings_changed)
        self.settings_server.textChanged.connect(self.account_settings_changed)
        self.apply_font()

    def update_task(self, fn, callback):
        if getattr(self,'update_worker',None) is not None: return
        worker=Worker(fn); self.update_worker=worker
        self.update_check.setEnabled(False); self.update_get.setEnabled(False)
        worker.event.connect(lambda event:self.update_status.setText(event.get('text','')))
        def done():
            self.update_worker=None; self.update_check.setEnabled(True); worker.deleteLater()
            if worker.failure:
                self.update_status.setText(worker.failure)
                self.update_get.setEnabled(bool(getattr(self,'latest_release',None) and self.latest_release.get('newer') and self.latest_release.get('asset')))
                QMessageBox.warning(self,'Cập nhật Chat AI',worker.failure)
            else: callback(worker.result)
        worker.finished.connect(done); worker.start()

    def check_updates(self):
        def task(emit):
            from assistant.updater import check_latest
            return check_latest()
        def done(release):
            self.latest_release=release
            self.update_status.setText(('Có bản mới: ' if release['newer'] else 'Phiên bản hiện tại đã mới nhất: ')+release['version'])
            self.update_notes.setPlainText(release['notes'] or 'Chưa có ghi chú phát hành.')
            self.update_get.setEnabled(bool(release['newer'] and release['asset']))
        self.update_status.setText('Đang kiểm tra GitHub Releases…'); self.update_task(task,done)

    def download_update(self):
        release=getattr(self,'latest_release',None)
        if not release or not release['newer'] or not release['asset']: return
        asset=release['asset']
        if QMessageBox.question(self,'Tải cập nhật',f'Tải {asset["name"]} ({asset["size"]/1e6:.1f} MB) từ GitHub vuanh97nd/ChatAI?\\nỨng dụng không tự chạy bộ cài.')!=QMessageBox.StandardButton.Yes:return
        def task(emit):
            import os
            from assistant.updater import download_asset
            directory=Path(os.environ.get('LOCALAPPDATA',str(ROOT/'data')))/'ChatAI'/'updates'
            result=download_asset(asset,directory,lambda text:emit({'text':text}))
            self.store.audit('updater','update_downloaded',{'version':release['version'],'sha256':result['sha256'],'verified':result['verified']})
            return result
        def done(result):
            self.update_status.setText('Đã tải. Đóng Chat AI rồi chạy bộ cài trong thư mục vừa mở.')
            self.open_path(Path(result['path']).parent)
        self.update_task(task,done)

    def sync_history(self):
        if not self.server_session or self.history_sync_loading or self.busy():return
        self.history_sync_loading=True
        session=dict(self.server_session)
        from threading import Thread
        def task():
            try:
                from assistant.history_sync import HistorySync
                changed=HistorySync(self.store,session).cycle()
                result={'session':session,'changed':changed}
            except Exception as error:result={'session':session,'error':str(error)}
            try:self.history_sync_finished.emit(result)
            except RuntimeError:pass
        Thread(target=task,daemon=True,name='ChatAI-history-sync').start()

    def apply_history_sync(self,result):
        self.history_sync_loading=False
        session=result['session']
        if not self.server_session or any(self.server_session.get(k)!=session.get(k) for k in ('username','endpoint')):return
        if result.get('error'):
            self.account_status.setToolTip('Lịch sử vẫn được lưu trên máy; đang chờ đồng bộ server: '+result['error'])
            return
        self.account_status.setToolTip('Nội dung hội thoại đã đồng bộ lên server. Ảnh và tệp gốc được giữ trên máy.')
        if result.get('changed') and not self.busy():
            state=self.store.load(self.cid)
            if not state.get('messages'):
                saved=self.store.last_conversation(session['username'])
                if saved:self.cid=saved
            self.html_cache.clear();self.render()

    def restore_login(self):
        if self.server_session or self.restore_login_loading:return
        if self.busy():QTimer.singleShot(500,self.restore_login);return
        self.restore_login_loading=True
        from threading import Thread
        def load():
            try:
                from assistant.accounts import load_login
                with measure('login.restore_credentials'):session=load_login()
                value={'session':session}
            except Exception:value={'error':True}
            try:self.login_restore_finished.emit(value)
            except RuntimeError:pass # The window closed while DPAPI/file I/O was pending.
        Thread(target=load,daemon=True).start()

    def apply_restored_login(self,value):
        self.restore_login_loading=False
        if self.server_session:return
        if value.get('error'):
            self.account_status.setText('Không đọc được đăng nhập đã lưu. Vui lòng đăng nhập lại.');return
        session=value.get('session')
        if not session:return
        # Áp session đã cache ngay lập tức — không chờ mạng
        self._apply_restored_session(session)
        # Xác minh token ở nền, cập nhật nếu cần
        self._verify_restored_session(session)

    def _apply_restored_session(self,session):
        from assistant.model_preferences import load_model
        self.server_session=dict(session)
        self.server_session.setdefault('fullname',session['username'])
        self.server_session.setdefault('role','user')
        self.personal_memories=[]
        self.install_custom_ai(self.cfg.get('custom_ai',[]))
        self.apply_account_model(load_model(self.store,session))
        if hasattr(self,'api_key_group'):self.api_key_group.setVisible(admin_session(self.server_session))
        self.settings_server.setText(session['endpoint'])
        self.chat_login.setText('Tài khoản: '+session['username'])
        self.account_status.setText('Đã đăng nhập: '+session['username'])
        self.trial.adopt(session['username'])
        current=self.store.load(self.cid)
        if current.get('account_username') not in (None,session['username']):self.cid=self.store.create(persist=False)
        saved=self.store.last_conversation(session['username'])
        if saved and not self.store.load(self.cid).get('messages'):self.cid=saved
        self.store.remember_conversation(session['username'],self.cid)
        self.render()

    def _verify_restored_session(self,session):
        from threading import Thread
        def verify():
            try:
                from assistant.accounts import request_account,save_login
                result=request_account(session['endpoint'],'/api/login',
                    {'username':session['username'],'key':session['key'],'device_id':self.device_id})
                updated=dict(session)
                if result.get('session_token'):updated['key']=result['session_token']
                for k in ('fullname','email','phone','avatar'):updated[k]=result.get(k,session.get(k,''))
                updated['fullname']=result.get('fullname') or session['username']
                updated['role']=result.get('role','user')
                try:save_login(updated)
                except Exception:pass
                self.session_verified.emit({'ok':True,'session':updated,'result':result})
            except Exception as exc:
                self.session_verified.emit({'ok':False,'error':str(exc),'username':session['username']})
        Thread(target=verify,daemon=True,name='ChatAI-session-verify').start()

    def _on_session_verified(self,payload):
        if not payload.get('ok'):
            error=payload.get('error','')
            # Chỉ đăng xuất khi token thực sự hết hạn (401/403), không phải lỗi mạng
            if any(code in error for code in ('HTTP 401','HTTP 403')):
                self.server_session=None;self.personal_memories=[]
                self.chat_login.setText('Đăng nhập')
                self.account_status.setText('Phiên đăng nhập hết hạn. Vui lòng đăng nhập lại.')
                self.render()
            # Lỗi mạng: giữ session, sẽ xác minh lần sau
            return
        updated=payload['session']
        if not self.server_session or self.server_session.get('username')!=updated.get('username'):return
        for k in ('key','fullname','email','phone','avatar','role'):
            if k in updated:self.server_session[k]=updated[k]
        result=payload['result']
        result['_custom_ai']=self.cfg.get('custom_ai',[])
        self.install_custom_ai(result.get('_custom_ai',[]))
        self.refresh_account_ui()
        self.load_account_enrichment(dict(self.server_session))
        self.sync_history()

    def login_dialog(self):
        if self.busy():return
        endpoint=self.settings_server.text().strip().rstrip('/')
        if not endpoint:
            self.tabs.setCurrentIndex(3);self.settings_server.setFocus()
            QMessageBox.information(self,'Chat AI','Nhập URL server HTTPS trong Cài đặt trước khi đăng nhập hoặc đăng ký.');return
        dialog=QDialog(self);dialog.setWindowTitle('Đăng nhập Chat AI');form=QFormLayout(dialog)
        username=QLineEdit();password=QLineEdit();password.setEchoMode(QLineEdit.EchoMode.Password)
        remember=QCheckBox('Ghi nhớ đăng nhập');remember.setChecked(self.account_remember.isChecked())
        form.addRow('Tên đăng nhập:',username);form.addRow('Mật khẩu:',password);form.addRow(remember)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);form.addRow(buttons)
        if dialog.exec()!=QDialog.DialogCode.Accepted:return
        if not username.text().strip() or not password.text().strip():return
        self.account_remember.setChecked(remember.isChecked())
        self.perform_login({'endpoint':endpoint,'username':username.text().strip(),'key':password.text().strip()},remember.isChecked())

    def perform_login(self,session,remember=True,restore=False):
        login_wall_started=time.monotonic()
        self.account_status.setText('Đang đăng nhập…')
        def task(emit):
            from assistant.accounts import request_account,save_login,forget_login
            started=time.monotonic()
            record('login.worker_setup',started-login_wall_started)
            with measure('login.https_request'):
                result=request_account(session['endpoint'],'/api/login',{'username':session['username'],'key':session['key'],'device_id':self.device_id})
            authenticated=time.monotonic()
            if result.get('session_token'):session['key']=result['session_token']
            # The login endpoint already returns identity and role; optional profile
            # enrichment must not hold up authentication or the chat controls.
            for key in ('fullname','email','phone','avatar'):result.setdefault(key,session.get(key,''))
            stored=dict(session)
            for key in ('fullname','email','phone','avatar'):stored[key]=result.get(key,'')
            stored['fullname']=result.get('fullname') or session['username'];stored['role']=result.get('role','user')
            warning=''
            try:
                with measure('login.save_credentials'):
                    if remember:save_login(stored)
                    else:forget_login()
            except Exception:
                warning=' · Đăng nhập thành công nhưng chưa lưu được ghi nhớ trên Windows'
            from assistant.model_preferences import load_model
            result['_custom_ai']=self.cfg.get('custom_ai',[])
            with measure('login.model_preference'):result['_saved_ai']=load_model(self.store,session)
            with measure('login.adopt_guest_history'):self.trial.adopt(session['username'])
            result['_login_timings']={'server_seconds':round(authenticated-started,3),'local_seconds':round(time.monotonic()-authenticated,3)}
            return result,[],warning
        def done(result):
            ui_started=time.monotonic()
            self.server_session=dict(session);self.server_session['fullname']=result[0].get('fullname') or session['username'];self.server_session['role']=result[0].get('role','user');self.personal_memories=result[1]
            for key in ('fullname','email','phone','avatar'):self.server_session[key]=result[0].get(key,'')
            self.install_custom_ai(result[0].get('_custom_ai',[]))
            if restore:self.apply_account_model(result[0]['_saved_ai'])
            else:
                from assistant.model_preferences import save_model
                save_model(self.store,self.server_session,self.model.currentText())
            if hasattr(self,'api_key_group'):self.api_key_group.setVisible(admin_session(self.server_session))
            self.settings_server.setText(session['endpoint'])
            self.chat_login.setText('Tài khoản: '+session['username'])
            timings=result[0].get('_login_timings',{})
            self.store.audit(self.cid,'login_timing',timings)
            seconds=timings.get('server_seconds',0)+timings.get('local_seconds',0)
            self.account_status.setText('Đã đăng nhập: '+session['username']+f' · {seconds:.1f} giây'+result[2])
            self.account_status.setToolTip(f"Chờ server: {timings.get('server_seconds',0):.1f} giây; xử lý trên máy: {timings.get('local_seconds',0):.1f} giây")
            self.load_account_enrichment(dict(self.server_session))
            current=self.store.load(self.cid)
            if current.get('account_username') not in (None,session['username']):self.cid=self.store.create(persist=False)
            saved=self.store.last_conversation(session['username'])
            if saved and not self.store.load(self.cid).get('messages'):self.cid=saved
            self.store.remember_conversation(session['username'],self.cid)
            self.render()
            self.sync_history()
            record('login.apply_interface',time.monotonic()-ui_started)
            record('login.total',time.monotonic()-login_wall_started)
        self.work(task,done)
        if self.worker:self.worker.performance_login=True

    def load_account_enrichment(self,session):
        from threading import Thread
        from assistant.accounts import request_account
        def fetch(section,path):
            try:
                result=request_account(session['endpoint'],path,{'username':session['username'],'key':session['key']},timeout=8)
                self.account_enrichment_finished.emit(session,{'section':section,'result':result})
            except Exception:
                # Optional enrichment failure never signs out a valid login.
                pass
        for section,path in (('profile','/api/account/profile/get'),('catalog','/api/provider/catalog')):
            Thread(target=fetch,args=(section,path),daemon=True).start()

    def apply_account_enrichment(self,session,value):
        current=self.server_session
        if not current or any(current.get(k)!=session.get(k) for k in ('endpoint','username','key')):return
        if value['section']=='profile':
            profile=value['result'].get('profile',{})
            for key in ('fullname','email','phone','avatar'):
                if key in profile:current[key]=profile[key]
            self.refresh_account_ui()
        elif value['section']=='catalog':
            from assistant.cloud import register_custom_ai
            if self.busy():
                QTimer.singleShot(500,lambda:self.apply_account_enrichment(session,value));return
            self.install_custom_ai(register_custom_ai(value['result'].get('entries',[])))

    def change_password_dialog(self):
        if self.busy():return
        if not self.server_session:self.login_dialog();return
        if self.server_session['username'].lower()=='admin':
            QMessageBox.information(self,'Đổi mật khẩu','Mật khẩu quản trị được đổi bằng secret ADMIN_KEY trên Cloudflare.');return
        session=dict(self.server_session)
        dialog=QDialog(self);dialog.setWindowTitle('Đổi mật khẩu');dialog.resize(420,260)
        form=QFormLayout(dialog)
        old,new,confirm=QLineEdit(),QLineEdit(),QLineEdit()
        for field in (old,new,confirm):field.setEchoMode(QLineEdit.EchoMode.Password);field.setMaxLength(128)
        form.addRow('Mật khẩu hiện tại',old);form.addRow('Mật khẩu mới',new);form.addRow('Nhập lại mật khẩu mới',confirm)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText('Đổi mật khẩu')
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText('Hủy')
        buttons.rejected.connect(dialog.reject);form.addRow(buttons)
        def accept():
            if not old.text().strip() or not 8<=len(new.text().strip())<=128:
                QMessageBox.warning(dialog,'Đổi mật khẩu','Nhập mật khẩu hiện tại và mật khẩu mới từ 8 đến 128 ký tự.');return
            if new.text()!=confirm.text():
                QMessageBox.warning(dialog,'Đổi mật khẩu','Hai mật khẩu mới chưa khớp.');return
            if old.text().strip()==new.text().strip():
                QMessageBox.warning(dialog,'Đổi mật khẩu','Mật khẩu mới cần khác mật khẩu hiện tại.');return
            dialog.accept()
        buttons.accepted.connect(accept)
        if dialog.exec()!=QDialog.DialogCode.Accepted:return
        current_password=old.text().strip()
        changed=dict(session);changed['key']=new.text().strip()
        def task(emit):
            from assistant.accounts import request_account,load_login,save_login,forget_login
            try:
                saved=load_login()
                remember=bool(saved and saved.get('username')==session['username'] and saved.get('endpoint')==session['endpoint'])
            except Exception:remember=False
            response=request_account(session['endpoint'],'/api/change_password',{'username':session['username'],'old_key':current_password,'new_key':changed['key']})
            if response.get('session_token'):changed['key']=response['session_token']
            warning=''
            try:
                if remember:save_login(changed)
                else:forget_login()
            except Exception:
                warning='\nChưa cập nhật được ghi nhớ đăng nhập. Lần mở sau hãy nhập mật khẩu mới.'
            return warning
        def done(warning):
            self.server_session=changed
            QMessageBox.information(self,'Đổi mật khẩu','Đổi mật khẩu thành công.'+warning)
        self.work(task,done)

    def logout_account(self):
        if self.busy():return
        from assistant.accounts import forget_login
        forget_login();self.server_session=None;self.personal_memories=[]
        if hasattr(self,'api_key_group'):self.api_key_group.hide()
        self.account_status.setText('Chưa đăng nhập.');self.cid=self.store.create(persist=False);self.render()

    def memory_dialog(self):
        if self.busy():return
        session=getattr(self,'server_session',None)
        if not session:QMessageBox.information(self,'Bộ nhớ','Đăng nhập tài khoản server trước.');return
        def task(emit):
            from assistant.accounts import request_account
            from assistant.memory import PersonalMemory
            try:remote=request_account(session['endpoint'],'/api/memory/personal/list',{'username':session['username'],'key':session['key']},timeout=8)['items']
            except Exception:
                remote=[m for m in self.personal_memories if m.get('location')!='local']
                emit({'type':'status','text':'Server chưa sẵn sàng; hiển thị ký ức đã lưu trên máy.'})
            return remote+PersonalMemory(self.store,self.client,session['username']).list_local()
        def show(items):
            self.personal_memories=items
            dialog=QWidget(); layout=QVBoxLayout(dialog)
            self.button(layout,'← Quay lại chat',lambda:self.tabs.setCurrentIndex(0))
            layout.addWidget(QLabel('Bộ nhớ riêng: '+(session.get('fullname') or session['username'])))
            listing=QListWidget();layout.addWidget(listing)
            for value in items:
                item=QListWidgetItem(('[Máy này] ' if value.get('location')=='local' else '[Server] ')+value['title']+'\n'+value['text'][:140]);item.setData(Qt.ItemDataRole.UserRole,value);listing.addItem(item)
            form=QFormLayout();title=QLineEdit();text=QPlainTextEdit();form.addRow('Tiêu đề',title);form.addRow('Nội dung',text);layout.addLayout(form)
            def selected(item):
                value=item.data(Qt.ItemDataRole.UserRole);title.setText(value['title']);text.setPlainText(value['text'])
            listing.itemClicked.connect(selected);action={}
            def choose(delete=False):
                item=listing.currentItem();old=item.data(Qt.ItemDataRole.UserRole) if item else None
                if delete and not old:return
                if not delete and (not title.text().strip() or not text.toPlainText().strip()):return
                if QMessageBox.question(dialog,'Bộ nhớ','Xóa ghi nhớ này?' if delete else ('Lưu nội dung này vào bộ nhớ trên máy?' if old and old.get('location')=='local' else 'Lưu nội dung này vào bộ nhớ riêng trên server?'))!=QMessageBox.StandardButton.Yes:return
                body={'username':session['username'],'key':session['key'],'confirm':True}
                if old:body['id']=old['id']
                if not delete:body.update(title=title.text().strip(),text=text.toPlainText().strip())
                action.update(path='/api/memory/personal/delete' if delete else '/api/memory/personal/put',body=body,local=bool(old and old.get('location')=='local'),delete=delete)
                self.work(commit,done)
            row=QHBoxLayout();self.button(row,'Mục mới',lambda:(listing.clearSelection(),listing.setCurrentRow(-1),title.clear(),text.clear()))
            self.button(row,'Lưu ghi nhớ',lambda:choose());self.button(row,'Xóa mục chọn',lambda:choose(True));layout.addLayout(row)
            def commit(emit):
                from assistant.accounts import request_account
                if action['local']:
                    from assistant.memory import PersonalMemory
                    memory=PersonalMemory(self.store,self.client,session['username'])
                    body=action['body'];ident=body['id'].removeprefix('local:')
                    if action['delete']:
                        memory.delete(ident);return {'id':body['id']}
                    memory.put(body['text'],body['title'],ident)
                    return {'item':{'id':body['id'],'title':body['title'],'text':body['text'],'location':'local'}}
                return request_account(session['endpoint'],action['path'],action['body'])
            def done(result):
                ident=result.get('id') or result['item']['id']
                self.personal_memories=[x for x in self.personal_memories if x['id']!=ident]
                if result.get('item'):self.personal_memories.insert(0,result['item'])
                self.memory_dialog()
            self.remove_dynamic_page('memory_page_index')
            self.add_scroll_page(dialog,'Bộ nhớ cá nhân')
            self.memory_page_index=self.tabs.count()-1
            self.tabs.setCurrentIndex(self.memory_page_index)
        self.work(task,show)

    def register_dialog(self):
        if self.busy(): return
        from urllib.parse import urlparse
        endpoint=self.settings_server.text().strip().rstrip('/')
        url=urlparse(endpoint)
        if url.scheme!='https' or not url.hostname or url.username or url.password or url.query or url.fragment or url.path not in ('','/'):
            QMessageBox.warning(self,'Chat AI','Nhập URL gốc HTTPS của server Chat AI.'); return
        dialog=QDialog(self); dialog.setWindowTitle('Tạo tài khoản Chat AI'); dialog.resize(440,280)
        form=QFormLayout(dialog); fullname=QLineEdit(); username=QLineEdit(); password=QLineEdit()
        password.setEchoMode(QLineEdit.EchoMode.Password)
        fullname.setMaxLength(120); username.setMaxLength(40); password.setMaxLength(128)
        form.addRow('Họ và tên:',fullname); form.addRow('Tên đăng nhập:',username); form.addRow('Mật khẩu:',password)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText('Tạo tài khoản')
        buttons.rejected.connect(dialog.reject)
        data={}
        def accept():
            name,user,pwd=fullname.text().strip(),username.text().strip(),password.text().strip()
            if not name or not re.fullmatch(r'[A-Za-z0-9_.-]{3,40}',user) or user.lower()=='admin' or not 8<=len(pwd)<=128:
                QMessageBox.warning(dialog,'Đăng ký','Nhập họ tên, tên đăng nhập 3–40 ký tự và mật khẩu ít nhất8 ký tự.'); return
            data.update(fullname=name,username=user,password=pwd); dialog.accept()
        buttons.accepted.connect(accept); form.addRow(buttons)
        if dialog.exec()!=QDialog.DialogCode.Accepted: return
        def task(emit):
            from assistant.accounts import request_account
            return request_account(endpoint,'/api/register',data,timeout=30)
        def done(result):
            QMessageBox.information(self,'Chat AI','Đã tạo tài khoản: '+data['username']+'\nKhông lưu mật khẩu vào ứng dụng. Chat local vẫn dùng Ollama.')
        self.work(task,done)

    def apply_font(self):
        self.view.setStyleSheet('font-size:%dpx;' % self.preview_font_size)

    def preview_appearance(self,value=None):
        if not hasattr(self,'paint_timer'):return
        self.apply_theme(self.settings_theme.currentData(),self.settings_fields['font_size'].value())

    def apply_theme(self,theme,font_size=None):
        # Fade the native window, not a QGraphicsEffect on the central widget.
        # Parent effects nest with composer shadows/child fades and Qt can skip
        # painting their subtrees. Keep the transition without that nesting.
        previous=getattr(self,'_theme_animation',None)
        if previous is not None:
            previous.stop();previous.deleteLater()
        self._theme_animation=None
        def apply_colors():
            self.preview_theme=theme
            self.preview_font_size=self.cfg.get('font_size',13) if font_size is None else font_size
            self.setStyleSheet(style_sheet(theme))
            for widget in self.findChildren(QWidget):
                original=widget.property('chatBaseStyle')
                if original is None:
                    original=widget.styleSheet()
                    if not re.search(r'#[0-9a-fA-F]{6}',original):continue
                    widget.setProperty('chatBaseStyle',original)
                widget.setStyleSheet(recolor(original,theme))
            if hasattr(self,'settings_button'):self.settings_button.set_theme(theme)
            if hasattr(self,'profile_avatar') and hasattr(self,'account_register_button'):self.refresh_account_ui()
            self.apply_font();self.html_cache.clear()
            if hasattr(self,'paint_timer'):self.draw()
        if not self.isVisible():
            apply_colors();return
        fade_out=QPropertyAnimation(self,b'windowOpacity',self)
        fade_out.setDuration(80);fade_out.setStartValue(self.windowOpacity());fade_out.setEndValue(.85)
        def fade_in():
            apply_colors();fade_out.deleteLater()
            animation=QPropertyAnimation(self,b'windowOpacity',self)
            animation.setDuration(200);animation.setStartValue(self.windowOpacity());animation.setEndValue(1.0)
            animation.setEasingCurve(QEasingCurve.Type.OutCubic)
            def complete():
                if self._theme_animation is animation:self._theme_animation=None
                animation.deleteLater()
            animation.finished.connect(complete)
            self._theme_animation=animation;animation.start()
        fade_out.finished.connect(fade_in)
        self._theme_animation=fade_out;fade_out.start()

    def light_settings(self):
        self.settings_fields['default_model'].setCurrentText('qwen2.5:3b')
        self.settings_fields['code_model'].setCurrentText('qwen2.5-coder:3b')
        self.settings_fields['num_ctx'].setValue(2048); self.settings_fields['num_predict'].setValue(768)

    def install_custom_ai(self,entries):
        from assistant.cloud import register_custom_ai
        entries=register_custom_ai(entries)
        self.cfg['custom_ai']=entries
        for item in entries:
            name=PROVIDER_NAMES[item['id']]
            if self.model.findText(name)<0:self.model.addItem(name)
            if self.settings_provider.findText(name)<0:self.settings_provider.addItem(name)
        from assistant.config import save_config
        self.cfg=save_config(self.cfg)

    def configure_deepseek_presets(self):
        if not admin_session(self.server_session) or self.busy():return
        from assistant.accounts import request_account
        session=dict(self.server_session)
        def done(value):
            replacements={item['id'] for item in value['entries']}
            entries=[item for item in self.cfg.get('custom_ai',[]) if item['id'] not in replacements]
            self.install_custom_ai([*entries,*value['entries']])
            self.select_ai('DeepSeek Flash')
            QMessageBox.information(self,'DeepSeek',value.get('message','Đã lưu cấu hình.'))
        self.work(lambda emit:request_account(session['endpoint'],'/api/admin/providers/deepseek-presets',dict(username=session['username'],key=session['key']),timeout=30),done)

    def add_ai_dialog(self):
        if not admin_session(self.server_session) or self.busy():return
        dialog=QDialog(self);dialog.setWindowTitle('Thêm AI trực tuyến');form=QFormLayout(dialog)
        label=QLineEdit();model=QLineEdit();key=QLineEdit();key.setEchoMode(QLineEdit.EchoMode.Password)
        key.setPlaceholderText('Để trống để dùng key chung đã lưu trên server')
        provider=QComboBox()
        for title,value in [('NVIDIA','nvidia'),('DeepSeek','deepseek'),('Gemini','gemini'),('Groq','groq')]:provider.addItem(title,value)
        form.addRow('Tên hiển thị',label);form.addRow('Dịch vụ',provider);form.addRow('Mã AI',model);form.addRow('API key',key)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Save|QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);form.addRow(buttons)
        if dialog.exec()!=QDialog.DialogCode.Accepted:return
        if not label.text().strip() or not model.text().strip():QMessageBox.warning(self,'Thêm AI','Nhập tên hiển thị và mã AI.');return
        session=dict(self.server_session)
        body=dict(username=session['username'],key=session['key'],label=label.text().strip(),model=model.text().strip(),provider=provider.currentData(),api_key=key.text().strip())
        from assistant.accounts import request_account
        def done(value):
            entries=[*self.cfg.get('custom_ai',[]),value['entry']];self.install_custom_ai(entries)
            QMessageBox.information(self,'Chat AI','Đã thêm AI trên server. Các tài khoản sẽ thấy AI sau khi đăng nhập lại; key không hiển thị cho người dùng.')
        self.work(lambda emit:request_account(session['endpoint'],'/api/admin/providers/add',body,timeout=30),done)

    def provider_key_status(self):
        if not admin_session(self.server_session) or self.busy():return
        from assistant.accounts import request_account
        session=dict(self.server_session)
        def done(value):
            states=value.get('providers',{})
            lines=[]
            for provider,label in [('nvidia','NVIDIA'),('deepseek','DeepSeek'),('gemini','Gemini'),('groq','Groq')]:
                item=states.get(provider,{})
                lines.append(label+(': Đã có key trên server' if item.get('configured') else ': Chưa có key'))
                if item.get('model'):self.api_model_fields[provider].setText(item['model'])
            QMessageBox.information(self,'API trên server','\n'.join(lines)+'\nĐể trống ô key khi giữ key đã lưu. Không cần nhập key trong Cloudflare.')
        self.work(lambda emit:request_account(session['endpoint'],'/api/admin/providers/status',dict(username=session['username'],key=session['key']),timeout=30),done)

    def load_provider_catalog(self,choose=False,provider='nvidia'):
        if not admin_session(self.server_session) or self.busy():return
        key=self.api_key_fields[provider].text().strip()
        if not choose and not key:return
        import hashlib
        fingerprint=provider+':'+hashlib.sha256(key.encode()).hexdigest()
        if not choose and fingerprint==getattr(self,'last_catalog_key',None):return
        session=dict(self.server_session)
        from assistant.accounts import request_account
        def done(value):
            models=value.get('models',[])
            if not models:QMessageBox.information(self,PROVIDER_NAMES[provider],'Dịch vụ không trả danh sách AI.');return
            current=self.api_model_fields[provider].text().strip()
            if choose:
                selected,ok=QInputDialog.getItem(self,'AI đang có tại '+PROVIDER_NAMES[provider],'Chọn AI trò chuyện (không chọn embedding hoặc reranker):',models,models.index(current) if current in models else 0,False)
                if ok:self.api_model_fields[provider].setText(selected)
            elif current not in models:
                marker={'nvidia':'nemotron','deepseek':'deepseek','groq':'llama'}.get(provider,provider)
                candidates=[m for m in models if marker in m.lower() and not any(t in m.lower() for t in ('embed','rerank','guard','safety','parse','reward'))]
                if candidates:
                    selected=next((m for m in candidates if ('super' if provider=='nvidia' else 'flash') in m.lower()),candidates[0])
                    self.api_model_fields[provider].setText(selected)
                    self.status.setText('Đã chọn mã AI có trong danh sách dịch vụ: '+selected+'. Bấm Lưu key API để lưu.')
                else:self.status.setText('Đã lấy danh sách dịch vụ. Bấm Chọn AI từ dịch vụ để chọn model.')
            self.last_catalog_key=fingerprint
        self.work(lambda emit:request_account(session['endpoint'],'/api/admin/providers/models',dict(username=session['username'],key=session['key'],provider=provider,api_key=key),timeout=30),done)

    def check_online_provider(self,provider):
        if not admin_session(self.server_session):return
        from assistant.accounts import request_account
        session=dict(self.server_session);key=self.api_key_fields[provider].text().strip();selected_model=self.api_model_fields[provider].text().strip()
        self.work(lambda emit:request_account(session['endpoint'],'/api/admin/providers/test',dict(username=session['username'],key=session['key'],provider=provider,api_key=key,model=selected_model),timeout=125),lambda value:QMessageBox.information(self,'Kết nối AI',value.get('message','Kết nối thành công.')))

    def refresh_machine_choices(self,*args):
        from assistant.hardware_profile import recommendations
        a=recommendations(self.machine_profile.currentData(),self.hardware_info)
        self.hardware_advice.setText(a['note']+'\nĐề xuất: '+a['chat'][0]+' / '+a['code'][0]+' / '+a['vision'][0])
        for key in ('default_model','code_model','vision_model'):
            if key in self.settings_fields:self.settings_fields[key].setEnabled(not self.machine_auto.isChecked())

    def apply_machine_choices(self,cfg):
        if not cfg.get('machine_auto_ai'):return
        from assistant.hardware_profile import recommendations,available_choice
        a=recommendations(cfg['machine_profile'],self.hardware_info)
        for role,key in [('chat','default_model'),('code','code_model'),('vision','vision_model')]:
            chosen=available_choice(a[role],getattr(self,'models',set()))
            if chosen:cfg[key]=chosen
        for key in ('num_ctx','num_predict'):cfg[key]=a[key]

    def detect_machine(self):
        from assistant.hardware_profile import detect_hardware,inferred_level,LABELS
        def done(info):
            self.hardware_info=info;level=inferred_level(info)
            self.hardware_summary.setText(f"CPU: {info['cpu']} · RAM: {info['ram_gb'] or '?'} GB\nGPU: {info['gpu']} · VRAM: {info['vram_gb'] or '?'} GB\nGợi ý: {LABELS[level]}")
            if not self.cfg.get('machine_profile_selected'):self.machine_profile.setCurrentIndex(self.machine_profile.findData(level))
            self.refresh_machine_choices()
        self.work(lambda emit:detect_hardware(),done)

    def download_recommended_ai(self):
        from assistant.hardware_profile import recommendations,available_choice
        a=recommendations(self.machine_profile.currentData(),self.hardware_info)
        missing=[a[r][0] for r in ('chat','code','vision') if not available_choice([a[r][0]],getattr(self,'models',set()))]
        if not missing:QMessageBox.information(self,'Chat AI','Các AI đề xuất đã được tải.');return
        name,ok=QInputDialog.getItem(self,'Tải AI đề xuất','Chọn AI muốn tải:',missing,0,False)
        if ok:self.download(name)

    def add_windows_app(self):
        path,_=QFileDialog.getOpenFileName(self,'Ứng dụng được phép','','Ứng dụng Windows (*.exe)')
        if path:
            paths=[p.strip() for p in self.windows_apps_paths.toPlainText().splitlines() if p.strip()]
            if path not in paths:paths.append(path)
            self.windows_apps_paths.setPlainText('\n'.join(paths))

    def begin_app_request(self):
        # A new explicit request resumes a stopped, completed task, not a live one.
        if self.busy() or not self.cfg.get('windows_apps_enabled'):return
        from assistant.windows_apps import resume_automation
        resume_automation()

    def stop_windows_apps(self):
        sender=self.sender() if hasattr(self,'sender') else None
        source=sender.text() if hasattr(sender,'text') else 'end_app_activity'
        try:self.store.audit(self.cid,'automation_stop_requested',{'source':source})
        except Exception:pass
        from assistant.windows_apps import stop_automation
        stop_automation()
        if self.worker and self.worker.cancellable:self.worker.stop_requested.set()
        if hasattr(self,'windows_readiness_label'):
            self.windows_readiness_label.setText('Đã dừng điều khiển ứng dụng; bấm Tiếp tục để cho phép lại.')
        self.status.setText('Đang dừng lượt hiện tại. Thao tác Windows đang thực hiện có thể cần hoàn tất.')

    def resume_windows_apps(self):
        if self.busy():
            self.status.setText('Đợi lượt đang dừng kết thúc rồi gửi yêu cầu mới.');return
        from assistant.windows_apps import resume_automation
        if not self.cfg.get('windows_apps_enabled'):
            QMessageBox.information(self,'Điều khiển app','Bật quyền và Lưu cài đặt trước.');return
        resume_automation();self.status.setText('Đã cho phép lại điều khiển app theo quyền đã lưu.')
        if hasattr(self,'windows_readiness_label'):
            from assistant.windows_apps import readiness
            self.windows_readiness_label.setText(readiness(self.cfg))

    def proposed_settings(self):
        proposed = {k:v for k,v in self.cfg.items() if k != 'roots'}
        for key,field in self.settings_fields.items():
            proposed[key] = field.currentText() if isinstance(field,QComboBox) else field.value()
        proposed['machine_profile']=self.machine_profile.currentData()
        proposed['machine_auto_ai']=self.machine_auto.isChecked()
        proposed['machine_profile_selected']=self.cfg.get('machine_profile_selected',False) or self.machine_profile.currentData()!=self.cfg.get('machine_profile','medium')
        self.apply_machine_choices(proposed)
        for provider,field in self.api_model_fields.items():proposed[provider+'_model']=field.text().strip()
        proposed['theme']=self.settings_theme.currentData()
        proposed['auto_python']=self.auto_python_check.isChecked()
        if hasattr(self,'ai_tools_auto_check'):proposed['ai_tools_auto_execute']=self.ai_tools_auto_check.isChecked()
        if hasattr(self,'online_tools_check'):
            proposed['online_tools_enabled']=self.online_tools_check.isChecked()
            proposed['online_document_upload']=self.online_upload_check.isChecked()
        proposed['chat_provider']=REMOTE_MODELS.get(self.settings_provider.currentText(),'local')
        proposed['server_url']=self.settings_server.text().strip()
        proposed['whitelist']=[x.strip() for x in self.settings_roots.toPlainText().splitlines() if x.strip()]
        if hasattr(self,'windows_apps_check'):
            proposed['windows_apps_enabled']=self.windows_apps_check.isChecked()
            proposed['windows_apps_compact']=self.windows_compact_check.isChecked()
            proposed['windows_apps_auto_execute']=self.windows_auto_execute_check.isChecked()
            proposed['automation_auto_install']=self.automation_auto_install_check.isChecked()
            proposed['windows_apps_all_installed']=self.windows_all_apps_check.isChecked()
            proposed['browser_background']=self.browser_background_check.isChecked()
            proposed['windows_apps_allowed']=[p.strip() for p in self.windows_apps_paths.toPlainText().splitlines() if p.strip()]
        return proposed

    def settings_dirty(self):
        return any(f.text().strip() for f in self.api_key_fields.values()) or self.proposed_settings()!={k:v for k,v in self.cfg.items() if k!='roots'}

    def discard_settings(self):
        for key,field in self.settings_fields.items():
            value=self.cfg.get(key,'gemma3:4b' if key=='vision_model' else None)
            if isinstance(field,QComboBox):field.setCurrentText(value)
            else:field.setValue(value)
        self.settings_theme.setCurrentIndex(self.settings_theme.findData(self.cfg.get('theme','dark')))
        self.apply_theme(self.cfg.get('theme','dark'),self.cfg.get('font_size',13))
        self.auto_python_check.setChecked(self.cfg.get('auto_python',True))
        if hasattr(self,'online_tools_check'):
            self.online_tools_check.setChecked(self.cfg.get('online_tools_enabled',False))
            self.online_upload_check.setChecked(self.cfg.get('online_document_upload',False))
        self.settings_provider.setCurrentText(PROVIDER_NAMES.get(self.cfg.get('chat_provider'),'AI trên máy'))
        self.settings_server.setText(self.cfg.get('server_url',''))
        self.settings_roots.setPlainText('\n'.join(self.cfg['whitelist']))
        if hasattr(self,'windows_apps_check'):
            self.windows_apps_check.setChecked(self.cfg.get('windows_apps_enabled',False))
            self.windows_compact_check.setChecked(self.cfg.get('windows_apps_compact',True))
            self.windows_auto_execute_check.setChecked(self.cfg.get('windows_apps_auto_execute',False))
            self.automation_auto_install_check.setChecked(self.cfg.get('automation_auto_install',False))
            self.windows_all_apps_check.setChecked(self.cfg.get('windows_apps_all_installed',False))
            self.browser_background_check.setChecked(self.cfg.get('browser_background',False))
            self.windows_apps_paths.setPlainText('\n'.join(self.cfg.get('windows_apps_allowed',[])))
        self.machine_profile.setCurrentIndex(self.machine_profile.findData(self.cfg.get('machine_profile','medium')))
        self.machine_auto.setChecked(self.cfg.get('machine_auto_ai',True))
        for provider,field in self.api_model_fields.items():field.setText(self.cfg.get(provider+'_model',''))
        for field in self.api_key_fields.values():field.clear()

    def check_settings_navigation(self,index):
        if self.settings_navigation_guard:return
        previous=self.settings_last_tab
        if previous==3 and index!=3 and self.settings_dirty():
            self.settings_navigation_guard=True;self.tabs.setCurrentIndex(3);self.settings_navigation_guard=False
            if self.busy():return
            dialog=QMessageBox(self);dialog.setWindowTitle('Cài đặt chưa lưu');dialog.setText('Bạn có muốn lưu thay đổi không?')
            apply=dialog.addButton('Lưu',QMessageBox.ButtonRole.AcceptRole)
            discard=dialog.addButton('Bỏ thay đổi',QMessageBox.ButtonRole.DestructiveRole)
            dialog.addButton('Ở lại',QMessageBox.ButtonRole.RejectRole);dialog.exec()
            if dialog.clickedButton() is apply:self.save_settings(target_index=index);return
            if dialog.clickedButton() is not discard:return
            self.discard_settings();self.settings_navigation_guard=True;self.tabs.setCurrentIndex(index);self.settings_navigation_guard=False
        self.settings_last_tab=index

    def save_settings(self,checked=False,target_index=None):
        if self.busy():
            QMessageBox.information(self,'Chat AI','Đợi tác vụ hiện tại kết thúc trước khi lưu.');return
        proposed=self.proposed_settings();proposed['machine_profile_selected']=True;old_endpoint=self.cfg.get('server_url')
        keys={k:f.text().strip() for k,f in self.api_key_fields.items() if f.text().strip()}
        provider_models={k:f.text().strip() for k,f in self.api_model_fields.items() if f.text().strip() and f.text().strip()!=self.cfg.get(k+'_model','')}
        session_snapshot=dict(self.server_session) if self.server_session else None
        def task(emit):
            from assistant.config import save_config
            cfg=save_config(proposed)
            session=session_snapshot
            if session and old_endpoint==cfg.get('server_url'):
                from assistant.model_preferences import save_model
                save_model(self.store,session,PROVIDER_NAMES.get(cfg['chat_provider'],cfg['default_model']))
                if admin_session(session) and (keys or provider_models):
                    from assistant.accounts import request_account
                    request_account(session['endpoint'],'/api/admin/providers/save',dict(username=session['username'],key=session['key'],providers=keys,models=provider_models))
            return cfg
        def done(cfg):
            self.cfg=cfg
            if hasattr(self,'windows_stop_button'):self.windows_stop_button.setVisible(cfg.get('windows_apps_enabled',False))
            if hasattr(self,'windows_readiness_label'):
                from assistant.windows_apps import readiness
                self.windows_readiness_label.setText(readiness(cfg))
            if old_endpoint!=cfg.get('server_url'):
                self.server_session=None;self.personal_memories=[]
                from assistant.accounts import forget_login
                forget_login();self.cid=self.store.create(persist=False)
            self.select_ai(PROVIDER_NAMES.get(cfg['chat_provider'],cfg['default_model']))
            for f in self.api_key_fields.values():f.clear()
            self.apply_theme(cfg['theme'],cfg['font_size']);self.html_cache.clear();self.render();self.status.setText('Đã lưu cài đặt.')
            QMessageBox.information(self,'Chat AI','Đã lưu cài đặt.')
            if target_index is not None:
                self.settings_navigation_guard=True;self.tabs.setCurrentIndex(target_index);self.settings_navigation_guard=False;self.settings_last_tab=target_index
        self.work(task,done)

    def poll(self):
        active_now = self.manager.busy
        if self.jobs_loaded and not active_now and not self.jobs_was_busy: return
        self.jobs_loaded, self.jobs_was_busy = True, active_now
        jobs = self.initial_jobs if self.initial_jobs is not None else self.manager.jobs()
        self.initial_jobs = None
        can_stop=active_now and self.manager.active_target in ALLOWED_MODELS and not self.manager.stop_event.is_set()
        self.pause_download_btn.setEnabled(can_stop);self.cancel_download_btn.setEnabled(can_stop)
        self.continue_download_btn.setEnabled(not active_now and any(j['status']=='paused' for j in jobs))
        self.jobs.setPlainText('\n\n'.join(f"{j['target']} · {j['status']}\n{j['message']}" for j in jobs))
        active = next((j for j in jobs if j['status'] in ('queued', 'running')), None)
        self.download_status.setText(('Đang ngắt tải tại điểm an toàn…' if self.manager.stop_event.is_set() else active['message']) if active else ('Không có tác vụ đang tải. Xem trạng thái hoàn tất/lỗi bên dưới.'))
        if active and active['progress'] is None:
            self.progress.setRange(0,100);self.progress.setValue(0);self.progress.setFormat('Đang chuẩn bị / cài thư viện — chưa có tổng dung lượng')
        else:
            self.progress.setRange(0, 100);self.progress.setFormat('%p%')
            self.progress.setValue(int(100 * ((active or {}).get('progress') or 0)))
        seen=set()
        for j in jobs:
            target=j['target']
            if target in seen:continue
            seen.add(target)
            if target not in ALLOWED_MODELS:continue
            if j['status']=='removed':self.models.discard(target)
            elif j['status']=='succeeded':self.models.add(target)

    def full_history(self):
        dialog = QDialog(self); dialog.setWindowTitle('Toàn bộ cuộc trò chuyện'); dialog.resize(820, 620)
        layout = QVBoxLayout(dialog); view = QPlainTextEdit(); view.setReadOnly(True)
        messages = self.store.load(self.cid)['messages']
        view.setPlainText('\n\n'.join(('Bạn:\n' if m['role'] == 'user' else 'Chat AI:\n') + m['content']
            for m in messages if m['role'] in ('user', 'assistant') and m.get('content')))
        layout.addWidget(view); dialog.exec()

    def remove_dynamic_page(self,attribute):
        old_index=getattr(self,attribute,None)
        if old_index is None:return
        page=self.tabs.widget(old_index)
        if page is None:
            setattr(self,attribute,None);return
        # Update all cached indices before removeTab emits currentChanged.
        for name in ('memory_page_index','profile_page_index','work_support_page_index','help_page_index','admin_page_index'):
            value=getattr(self,name,None)
            if value==old_index:setattr(self,name,None)
            elif value is not None and value>old_index:setattr(self,name,value-1)
        if getattr(self,'settings_last_tab',0)>old_index:self.settings_last_tab-=1
        elif getattr(self,'settings_last_tab',0)==old_index:self.settings_last_tab=0
        self.tabs.removeTab(old_index)
        page.deleteLater()

    def balance_panels(self, index=None):
        if hasattr(self,'settings_button'):self.settings_button.setChecked(self.tabs.currentIndex()==3)
        # Keep navigation stable while the right-hand page changes.
        if hasattr(self, 'sidebar'):
            collapsed=getattr(self,'sidebar_collapsed',False)
            self.sidebar.setFixedWidth(92 if collapsed else 260)
            _dynamic_pages={getattr(self,'memory_page_index',None),getattr(self,'profile_page_index',None)}-{None}
            self.sidebar_stack.setCurrentIndex(1 if self.tabs.currentIndex() in (1,3) or self.tabs.currentIndex() in _dynamic_pages else 0)
            self.main_splitter.setStretchFactor(0,0)
            self.main_splitter.setStretchFactor(1,1)

    def toggle_sidebar(self):
        self.sidebar_collapsed=not getattr(self,'sidebar_collapsed',False)
        collapsed=self.sidebar_collapsed
        self.sidebar_brand_title.setVisible(not collapsed)
        self.sidebar_stack.setVisible(not collapsed)
        self.profile_button.setVisible(not collapsed)
        self.profile_avatar.setVisible(not collapsed)
        self.sidebar.setFixedWidth(92 if collapsed else 260)

    def filter_history(self, query):
        query=query.strip().casefold()
        for i in range(self.history.count()):
            item=self.history.item(i);item.setHidden(bool(query and query not in item.text().casefold()))

    def build_settings_navigation(self):
        layout=self.settings_nav_layout
        self.button(layout,'← Quay lại chat',lambda:self.tabs.setCurrentIndex(0))
        heading=QLabel('CÀI ĐẶT');heading.setStyleSheet('color:#9aa0a6;font-size:12px;padding:12px 4px;');layout.addWidget(heading)
        self.settings_nav_buttons={}
        for label,section in [('Chung',None),('Cấu hình máy và AI','Cấu hình máy và AI'),('Giao diện','Giao diện'),('Office và công cụ','Office và thư mục'),('Điều khiển ứng dụng','Điều khiển ứng dụng'),('Tài khoản và đăng nhập','Tài khoản'),('Lịch sử và dữ liệu','Lịch sử và dữ liệu'),('Cập nhật','Cập nhật Chat AI'),('Nâng cao','Nâng cao')]:
            button=self.button(layout,label,lambda checked=False,n=section:self.open_settings_section(n))
            button.setCheckable(True);self.settings_nav_buttons[label]=(button,section)
        self.button(layout,'Bộ nhớ cá nhân',self.memory_dialog)
        self.button(layout,'Tải mô hình',lambda:self.tabs.setCurrentIndex(1))
        layout.addStretch(1)

    def copy_performance_report(self):
        from assistant.performance import report
        QApplication.clipboard().setText(report())
        self.status.setText('Đã sao chép thời gian khởi động/đăng nhập. Bạn có thể dán để kiểm tra.')

    def account_menu(self):
        menu=QMenu(self)
        if self.server_session:
            menu.addAction(self.server_session.get('fullname') or self.server_session['username']).setEnabled(False)
            if admin_session(self.server_session):menu.addAction('Quản lý người dùng',self.open_user_admin)
            menu.addAction('Hồ sơ',self.open_profile)
            menu.addAction('Chỉnh sửa thông tin cá nhân',self.edit_personal_info)
            menu.addAction('Cá nhân hóa',lambda:self.open_settings_section('Cấu hình máy và AI'))
            menu.addAction('Tài khoản',lambda:self.open_settings_section('Tài khoản'))
            menu.addAction('Bộ nhớ cá nhân',self.memory_dialog)
            menu.addAction('Đổi mật khẩu',self.change_password_dialog)
            menu.addSeparator();menu.addAction('Đăng xuất',self.logout_account)
        else:
            menu.addAction('Đăng nhập',self.login_dialog)
            menu.addAction('Tạo tài khoản',self.register_dialog)
        menu.exec(self.profile_button.mapToGlobal(self.profile_button.rect().topLeft()))

    def settings_menu(self):
        menu=QMenu(self)
        for name in self.settings_sections:
            menu.addAction(name,lambda checked=False,n=name:self.open_settings_section(n))
        menu.addSeparator()
        menu.addAction('Bộ nhớ cá nhân',self.memory_dialog)
        menu.addAction('Tải mô hình',lambda:self.tabs.setCurrentIndex(1))
        menu.addAction('Ảnh và Video',lambda:self.tabs.setCurrentIndex(2))
        menu.addAction('Hỗ trợ',self.open_support)
        menu.addAction('Sao chép thời gian khởi động/đăng nhập',self.copy_performance_report)
        menu.addSeparator();menu.addAction('Tất cả cài đặt',lambda:self.open_settings_section(None))
        menu.exec(self.settings_button.mapToGlobal(self.settings_button.rect().topLeft()))

    def open_about(self):
        dialog=QDialog(self)
        dialog.setWindowTitle('Giới thiệu Chat AI')
        dialog.setMinimumWidth(420)
        layout=QVBoxLayout(dialog)
        logo=QLabel();pixmap=QPixmap(str(ROOT/'logo_chat_ai.png'))
        if not pixmap.isNull():
            logo.setPixmap(pixmap.scaled(88,88,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation))
            logo.setAlignment(Qt.AlignmentFlag.AlignHCenter)
            layout.addWidget(logo)
        title=QLabel('Chat AI')
        title.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        title.setStyleSheet('font-size:20px;font-weight:600;')
        layout.addWidget(title)
        details=QLabel(
            'Desktop 2.6.6\n'
            'Ứng dụng AI hỗ trợ trò chuyện, xử lý tài liệu và sáng tạo nội dung.\n\n'
            'Tác giả: Vũ Ngọc Ánh\n'
            'Email: vuanh97nd@gmail.com\n'
            'Mã nguồn và hỗ trợ: github.com/vuanh97nd/ChatAI'
        )
        details.setWordWrap(True)
        details.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        layout.addWidget(details)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        dialog.exec()

    def open_settings_section(self, name):
        if getattr(self,'sidebar_collapsed',False):self.toggle_sidebar()
        self.settings_heading.setText(name or 'Cài đặt chung')
        for button,section in self.settings_nav_buttons.values():button.setChecked(section==name)
        for title,widget in self.settings_sections.items():
            widget.setVisible(name is None or title==name)
        self.tabs.setCurrentIndex(3)
        self.settings_button.setChecked(True)
        self.api_key_group.setVisible(admin_session(self.server_session) and name in (None,'Cấu hình máy và AI'))
        scroll=self.tabs.widget(3)
        scroll.verticalScrollBar().setValue(0)

    def show_audit(self):
        dialog = QDialog(self); dialog.setWindowTitle('Nhật ký thao tác'); dialog.resize(820, 600)
        layout = QVBoxLayout(dialog); view = QPlainTextEdit(); view.setReadOnly(True)
        view.setPlainText(json.dumps(self.store.recent_audit(self.cid), ensure_ascii=False, indent=2)); layout.addWidget(view); dialog.exec()

    def refresh_media(self):
        if self.busy(): return
        roots = list(self.cfg['roots'])
        def scan(emit):
            records = []
            # Đầu ra của media nằm ở outputs. Không quét toàn bộ cây Drive trên UI.
            for root in roots:
                for folder in (root, root / 'outputs'):
                    if not folder.is_dir(): continue
                    for index, p in enumerate(folder.iterdir()):
                        if index >= 1000: break
                        if p.suffix.lower() in ('.png', '.mp4') and p.is_file() and p.resolve().is_relative_to(root):
                            records.append((p.name, str(p)))
                            if len(records) >= 100: return records
            return records
        def done(records):
            self.media_list.clear()
            for name, path in records:
                item = QListWidgetItem(name); item.setData(Qt.ItemDataRole.UserRole, path)
                self.media_list.addItem(item)
            self.status.setText(f'Đã tìm thấy {len(records)} ảnh/video. Bạn có thể gửi yêu cầu tiếp.')
        self.work(scan, done)

    def open_media(self):
        item = self.media_list.currentItem()
        if item: self.open_path(Path(item.data(Qt.ItemDataRole.UserRole)))

    def open_path(self, path):
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))

    def closeEvent(self, event):
        if getattr(self,'exit_when_idle',False) and not self.busy():
            if hasattr(self,'automation_panel'):self.automation_panel.end()
            event.accept();return
        if self.worker and self.worker.cancellable and getattr(self,'update_worker',None) is None:
            self.exit_when_idle=True
            self.send_or_stop()
            self.status.setText('Đang dừng tác vụ để thoát; không đóng ứng dụng bên ngoài.')
            event.ignore();return
        if self.busy() or getattr(self,'update_worker',None) is not None:
            QMessageBox.information(self, 'Chat AI', 'Đang xử lý câu trả lời hoặc cập nhật ứng dụng. Vui lòng đợi hoàn tất rồi đóng.'); event.ignore()
        else:
            if self.settings_dirty():
                dialog=QMessageBox(self);dialog.setWindowTitle('Cài đặt chưa lưu');dialog.setText('Bạn có muốn lưu thay đổi không?')
                apply=dialog.addButton('Lưu',QMessageBox.ButtonRole.AcceptRole);discard=dialog.addButton('Bỏ thay đổi',QMessageBox.ButtonRole.DestructiveRole);dialog.addButton('Hủy',QMessageBox.ButtonRole.RejectRole);dialog.exec()
                if dialog.clickedButton() is apply:
                    self.save_settings();event.ignore();return
                if dialog.clickedButton() is not discard:event.ignore();return
                self.discard_settings()
            dialog = QMessageBox(self)
            dialog.setWindowTitle('Thoát Chat AI')
            dialog.setIcon(QMessageBox.Icon.Question)
            dialog.setText('Bạn có muốn thoát Chat AI không?')
            if self.manager.busy:
                dialog.setInformativeText('Module đang tải. Bạn vẫn có thể thoát; khi mở lại, bấm tải lại để tiếp tục phần mô hình đã lưu nếu tác vụ bị ngắt.')
            exit_button = dialog.addButton('Thoát', QMessageBox.ButtonRole.AcceptRole)
            cancel_button = dialog.addButton('Hủy', QMessageBox.ButtonRole.RejectRole)
            dialog.setDefaultButton(cancel_button)
            dialog.setEscapeButton(cancel_button)
            dialog.exec()
            if dialog.clickedButton() is exit_button:
                if self.server_session:
                    self.store.remember_conversation(self.server_session['username'],self.cid)
                event.accept()
            else:
                event.ignore()


def main():
    from app import main as launch
    return launch()
