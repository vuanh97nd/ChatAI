from pathlib import Path
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/huong-dan-word'
BOOKS=[]
def book(name,title,subtitle,sections): BOOKS.append((name,title,subtitle,sections))
def s(title,*items):return title,items
# Items: plain paragraphs, sequential steps, or comparison tables.
def steps(*items):return ('steps',items)
def table(headers,*rows):return ('table',headers,rows)

book('01_Bat_dau_nhanh.docx','BẮT ĐẦU VỚI CHAT AI','Hướng dẫn dành cho người dùng mới',[
s('1. ChatAI giúp bạn làm gì?',
'ChatAI hỗ trợ hỏi đáp, soạn nội dung, xử lý tài liệu và giao việc cho máy tính. Bản Windows có các công cụ làm việc với file và phần mềm đã được cấp quyền. Bản Android giúp chat, nhập bằng giọng nói, quản lý ví và giao việc cho Windows đã ghép nối.',
'Android không tự chạy PLAXIS hoặc GeoStudio trên điện thoại. Các phần mềm kỹ thuật và giấy phép vẫn phải có trên máy Windows.',
table(['Nhu cầu','Nên dùng'],['Hỏi đáp, soạn nội dung khi di chuyển','Android hoặc Windows'],['Đọc và xử lý PDF, Excel, Word, DXF trên máy','Windows'],['Giao việc từ điện thoại, xem tiến trình','Android ghép nối Windows'],['Quản lý tài khoản, token, thông báo','Admin trên Windows hoặc Android theo phạm vi hỗ trợ'])),
s('2. Chuẩn bị trước khi dùng',
steps('Chuẩn bị tài khoản ChatAI do đơn vị vận hành cung cấp hoặc đăng ký trong ứng dụng.','Bảo đảm có Internet để đăng nhập và dùng AI trực tuyến.','Cài bản Windows hoặc APK Android phù hợp. Nếu chỉ dùng AI trực tuyến, không cần tải mô hình AI trên máy.','Nếu xử lý phần mềm kỹ thuật, chuẩn bị phần mềm đã cài, giấy phép, tài liệu và thư mục dữ liệu được phép truy cập.')),
s('3. Đăng ký và đăng nhập',
steps('Chọn Tạo tài khoản. Nhập họ tên, tên đăng nhập, email và mật khẩu. Email là thông tin bắt buộc khi đăng ký.','Kiểm tra thông tin trước khi gửi. Android đăng ký thành công sẽ tự đăng nhập.','Lần sau đăng nhập bằng tên đăng nhập hoặc email đủ điều kiện được server chấp nhận. Nếu email chưa được xác minh hoặc bị từ chối, dùng tên đăng nhập và liên hệ Admin.','Chỉ bật ghi nhớ đăng nhập trên thiết bị riêng. Đăng xuất khi bàn giao thiết bị.'),
'Yêu cầu điền email không đồng nghĩa đã bật OTP xác thực đăng ký. Gửi email xác minh hoặc đặt lại mật khẩu phụ thuộc cấu hình email của Admin.'),
s('4. Gửi yêu cầu đầu tiên',
steps('Chọn NVIDIA trực tuyến nếu muốn bắt đầu với lựa chọn mặc định.','Tạo cuộc trò chuyện mới và nhập một yêu cầu rõ ràng.','Đọc phản hồi, kiểm tra dữ kiện và yêu cầu sửa hoặc bổ sung ngay trong cùng cuộc trò chuyện.','Tách sang cuộc trò chuyện mới khi đổi dự án hoặc chủ đề để tránh dùng nhầm dữ liệu.'),
'Ví dụ: “Soạn email gửi đối tác, nội dung xin lịch họp vào tuần sau, giọng lịch sự, khoảng 150 từ.”',
'Ví dụ xử lý file trên Windows: “Đọc file Bao_cao.xlsx trong thư mục đã cấp. Liệt kê các sheet, kiểm tra ô thiếu và lập bảng tổng hợp. Tạo file đầu ra mới, giữ nguyên file gốc.”'),
s('5. Chọn AI và hiểu chi phí',
table(['Lựa chọn','Cách dùng / chi phí'],['NVIDIA, Cloud AI, AI trên máy','Được ứng dụng cấu hình miễn phí cho người dùng; khả năng và mức sẵn có phụ thuộc dịch vụ hoặc máy.'],['DeepSeek','Tính phí token theo bảng giá hiện hành trên server; cần số dư phù hợp.'],['OpenAI và các AI khác','Chỉ dùng khi đơn vị vận hành đã cấu hình và cho phép. Không mặc định tất cả đều miễn phí.']),
'Giá DeepSeek và phí duy trì có thể thay đổi theo Admin. Nếu phí duy trì bằng 0 đ, vẫn cần tiền token để dùng AI tính phí. Không có hạn mức token tháng do ứng dụng đặt không có nghĩa một yêu cầu có thể dài vô hạn.'),
s('6. Những việc nên kiểm tra trước khi tin kết quả',
steps('Đối chiếu tên file, mặt cắt, đơn vị và dữ liệu mà AI đã đọc.','Phân biệt kết quả đã lưu trong file với kết quả vừa chạy mới.','Kiểm tra file đầu ra có tồn tại và mở được.','Với mô hình kỹ thuật, kiểm tra hình học, vật liệu, điều kiện biên, pha tính, trạng thái tính toán và kết quả số.'),
'Khi AI cần thêm dữ kiện, cung cấp đúng phần còn thiếu. Không yêu cầu tự đoán chỉ tiêu đất, đơn vị hoặc kết quả để hoàn tất bài toán.'),
s('7. Đọc tiếp file nào?',
table(['File','Nội dung'],['02_Windows.docx','Cài đặt, chat, tài liệu, quyền, bộ nhớ và cập nhật'],['03_Android.docx','Cài APK, giọng nói, ví và kết nối Windows'],['04_Admin.docx','Người dùng, thanh toán, email, AI và đào tạo'],['05_PLAXIS_GeoStudio.docx','Chuẩn bị dữ liệu và kiểm tra kết quả kỹ thuật'],['06_Xu_ly_loi.docx','Đăng nhập, AI, QR, Git, APK và log']))
])

book('02_Windows.docx','SỬ DỤNG CHAT AI TRÊN WINDOWS','ChatAI Desktop 2.6.6 • người dùng và người vận hành',[
s('1. Cài đặt và mở ứng dụng',
steps('Dùng bộ cài ChatAI do đơn vị vận hành phát hành. Bộ cài có runtime Python nên máy đích không cần cài Python riêng.','Mở ChatAI từ biểu tượng trên Desktop hoặc Start Menu. Đợi màn hình khởi động kết thúc.','Kiểm tra địa chỉ server trong Cài đặt khi cần; chỉ dùng địa chỉ HTTPS do Admin cung cấp.','Đăng nhập tài khoản. Với AI trực tuyến, bắt đầu bằng NVIDIA; không bắt buộc cài Ollama.'),
'Nếu chạy từ mã nguồn: mở PowerShell tại thư mục dự án và chạy .\run.bat. Trình chạy có thể chuẩn bị môi trường; đọc thông báo nếu thiếu Python hoặc thư viện. Python 3.11/3.12 64-bit được hướng dẫn cho bản desktop này.',
'AI trên máy cần Ollama, dịch vụ đang hoạt động và mô hình đã tải. Chỉ cài phần này nếu thực sự muốn suy luận trên máy.'),
s('2. Các khu vực chính',
table(['Khu vực','Cách sử dụng'],['Thanh bên trái','Cuộc trò chuyện mới, tìm hội thoại, thư viện tài liệu và các tiện ích theo bản cài.'],['Vùng hội thoại','Đọc phản hồi, sao chép nội dung, kiểm tra file đầu ra và trạng thái thao tác.'],['Ô nhập phía dưới','Nhập yêu cầu; Enter gửi, Shift+Enter xuống dòng.'],['Thanh lựa chọn AI','Chọn AI trực tuyến hoặc AI trên máy và mô hình phù hợp.'],['Menu Tài khoản / Cài đặt','Hồ sơ, ví, thông báo, kết nối điện thoại, cấu hình và chức năng Admin nếu có quyền.'])),
s('3. Đọc file và làm việc với tài liệu',
steps('Đặt file vào thư mục được cấp quyền trong Cài đặt. Đối với Google Drive, bảo đảm file đã có sẵn ngoại tuyến trên máy.','Đính kèm hoặc nêu đường dẫn rõ ràng. Nêu tên file, sheet, vùng ô, chương hoặc mặt cắt cần đọc.','Yêu cầu AI báo tên nguồn và phạm vi đã đọc trước khi tổng hợp.','Khi cần sửa hoặc xuất tài liệu, chỉ định tên file đầu ra mới và yêu cầu giữ nguyên nguồn.','Mở file kết quả để kiểm tra nội dung và định dạng.'),
'Excel: nêu rõ sheet, vùng ô, đơn vị và cách xử lý ô trống/công thức. Word: ưu tiên DOCX. Với file DOC/PPT cũ hoặc macro, cần chuyển đổi hoặc công cụ tương thích; không coi là đã hỗ trợ mặc định.',
'PDF có lớp chữ và PDF scan là hai trường hợp khác nhau. Nếu PDF scan cần OCR/đọc ảnh, chọn luồng và mô hình thực sự hỗ trợ, kiểm tra phạm vi trang đã đọc. Không xem tuyên bố “đọc toàn bộ” là bằng chứng nếu không có số trang/nguồn.'),
s('4. Đọc ảnh và tìm kiếm mạng',
'Có thể đính kèm ảnh hoặc dán ảnh từ clipboard theo giao diện bản cài. Mô hình được chọn phải hỗ trợ ảnh; mô hình chỉ nhận văn bản không tự hiểu ảnh.',
'Nút Tìm web/Tìm kiếm mạng hỗ trợ tra cứu. Yêu cầu AI đưa liên kết nguồn và phân biệt dữ liệu tham khảo với yêu cầu của bạn. Không đưa mật khẩu, API key hoặc dữ liệu bí mật vào truy vấn web.',
'Ví dụ: “Tra hướng dẫn chính thức cho phiên bản phần mềm đang cài, trích nguồn và kiểm tra lệnh trước khi dùng.”'),
s('5. Cấp quyền cho AI làm việc',
steps('Mở Cài đặt → Office và công cụ / Điều khiển ứng dụng tùy chức năng.','Chọn các thư mục và ứng dụng thực sự cần cho công việc.','Nêu phạm vi được làm: đọc dữ liệu, tạo bản sao, lưu đầu ra; những việc cần giữ nguyên.','Theo dõi tiến trình và trả lời khi thiếu dữ kiện hoặc có thao tác cần xác nhận.','Dùng Dừng app AI khi muốn yêu cầu dừng. Công cụ đang chạy có thể cần trả quyền điều khiển trước khi dừng hoàn toàn.'),
'Không ghi đè mô hình PLAXIS đang mở. Lưu dự án hiện tại và chuẩn bị dự án riêng trước yêu cầu dựng mô hình mới. Quyền truy cập file không thay thế giấy phép hoặc quyền của phần mềm kỹ thuật.'),
s('6. Bộ nhớ và lịch sử',
table(['Loại','Ý nghĩa'],['Hội thoại','Nội dung trao đổi và tiến độ công việc trong cuộc trò chuyện.'],['Bộ nhớ cá nhân','Thông tin bạn chủ động lưu theo tài khoản, như cách trình bày hoặc ưu tiên.'],['Bài học thực hiện','Lỗi công cụ và bước sửa có bằng chứng; hỗ trợ tham khảo cho lần sau.'],['Bộ nhớ đào tạo chung','Bài học do Admin chia sẻ; không dùng thay dữ liệu công trình mới.']),
'Vào Cài đặt → Bộ nhớ cá nhân để quản lý mục ghi nhớ. Chỉ lưu thông tin cần thiết; không lưu khóa bí mật. Đồng bộ hội thoại/checkpoint không có nghĩa công việc sẽ tự chạy lại khi mở app.'),
s('7. Ví, thông báo và điện thoại',
'Vào menu tài khoản → Số dư và thanh toán để xem số dư, token theo tháng, lịch sử và QR. Thông báo có thể mở qua nút chuông/menu theo bản cài. Nếu chức năng không hiện, kiểm tra bản đang chạy và quyền tài khoản.',
'Kết nối điện thoại: menu Tài khoản → Kết nối điện thoại. Quy trình ghép và xem ảnh được hướng dẫn riêng trong 03_Android.docx.'),
s('8. Cập nhật bản mã nguồn',
steps('Đóng ChatAI và các tiến trình Git đang thao tác kho.','Mở PowerShell tại đúng thư mục dự án.','Chạy git status. Nếu có thay đổi cục bộ, sao lưu hoặc commit trước khi cập nhật.','Chạy git switch main rồi git pull origin main.','Chỉ khi pull thành công mới mở lại bằng .\run.bat.'),
'Bản EXE cần bộ cài mới; git pull trong thư mục mã nguồn không tự cập nhật EXE đã cài ở vị trí khác. Xử lý file khóa Git theo 06_Xu_ly_loi.docx, không dùng reset --hard để tránh mất thay đổi.')
])

book('03_Android.docx','SỬ DỤNG CHAT AI TRÊN ANDROID','Bản thử nghiệm 0.8 • chat, giọng nói và kết nối Windows',[
s('1. Tải và cài APK',
steps('Mở github.com/vuanh97nd/ChatAI/actions và chọn Build Android APK.','Chọn lượt build mới nhất phù hợp và kiểm tra toàn bộ lượt đã có dấu xanh.','Trong Artifacts, tải ChatAI-Android-test-…; giải nén ZIP lấy app-debug.apk.','Chuyển file sang điện thoại Android và mở bằng trình quản lý file.','Nếu Android yêu cầu, cho phép ứng dụng đang mở APK cài từ nguồn này. Sau khi cài, mở Chat AI.'),
'Không cần Flutter hoặc Android Studio khi chỉ tải APK có sẵn. Android 6.0 trở lên là mức tối thiểu hiện được cấu hình.',
'APK hiện là bản debug thử nghiệm. Nếu cập nhật báo khác chữ ký, sao lưu những gì cần giữ, ghi lại tài khoản và kiểm tra hướng dẫn trước khi gỡ bản cũ; gỡ app xóa dữ liệu cục bộ và khóa ghép nối. Bản phát hành ổn định cần khóa ký cố định.'),
s('2. Dùng thử, đăng nhập và điều hướng',
'Bản Android mới mở thẳng Chat, cho 3 lượt gửi tin nhắn NVIDIA theo mã thiết bị của bản cài. Lượt đã gửi tới dịch vụ, kể cả mất phản hồi/lỗi, vẫn được ghi nhận; từ lượt tiếp theo cần đăng nhập hoặc đăng ký. Trước đăng nhập không có ví, bộ nhớ, Admin hoặc điều khiển máy tính. Cần cài APK và Worker mới tương thích.',
'Đăng nhập/đăng ký dùng logo ChatAI xanh–tím trong bản mới. Nếu vẫn thấy robot trắng, có thể đang dùng APK cũ.',
table(['Tab','Chức năng'],['Chat','NVIDIA/DeepSeek, nhập câu hỏi, micro, đọc câu trả lời'],['Hội thoại','Mở lại các hội thoại Android đã lưu trên server'],['Bộ nhớ','Xem/thêm bộ nhớ cá nhân và xem thông báo'],['Tài khoản','Số dư, token, QR, đăng xuất; Admin có mục Quản trị'],['Máy tính','Ghép nối, giao việc, xem tiến trình/kết quả và yêu cầu ảnh']),
'Hội thoại Android và kho lịch sử desktop hiện tách riêng. Không thấy hội thoại Windows trong tab Hội thoại Android không có nghĩa đã bị mất. Ví và bộ nhớ cá nhân dùng chung tài khoản.'),
s('3. Chat và giọng nói',
steps('Chọn NVIDIA hoặc DeepSeek trong Chat.','Nhập câu hỏi hoặc bấm micro; cho phép ghi âm nếu muốn dùng giọng nói.','Nói tiếng Việt, đọc bản nháp, sửa nội dung nếu cần và bấm Gửi. Giọng nói không tự gửi yêu cầu.','Bấm Đọc câu trả lời để nghe; bấm Dừng đọc khi cần.','Chuyển trang hoặc đưa app xuống nền sẽ dừng nghe/đọc.'),
'Điện thoại cần dịch vụ nhận dạng và giọng đọc tiếng Việt. Nhận dạng có thể cần mạng và do dịch vụ của hệ điều hành xử lý. Đây là thao tác bằng nút bấm, chưa phải cuộc gọi giọng nói liên tục.'),
s('4. Nạp token qua QR',
steps('Mở Tài khoản, xem số dư và bảng giá hiện tại.','Nhập số tiền 20.000–10.000.000 đ, bội số 1.000 đ.','Bấm tạo QR và kiểm tra số tiền, ngân hàng, tên người nhận, nội dung chuyển khoản và thời hạn.','Chuyển đúng số tiền/nội dung. QR có hiệu lực theo thời hạn hiển thị; đơn hiện tại thường 30 phút.','Quay lại app, chờ server xác nhận và kiểm tra lịch sử nạp/số dư.'),
'QR không tự chứng minh đã thanh toán. Nếu chuyển sai số tiền/nội dung hoặc đơn hết hạn, liên hệ Admin kèm chứng từ. Khi không rõ kết quả mạng, kiểm tra đơn đã có trước khi tạo hoặc chuyển thêm lần nữa.',
'Phí duy trì do Admin đặt. Giá 0 đ nghĩa miễn phí duy trì, không miễn phí token DeepSeek. Admin hệ thống không có nút nạp như người dùng thường.'),
s('5. Ghép nối với Windows',
steps('Mở ChatAI Windows và đăng nhập cùng tài khoản với Android.','Windows: Tài khoản → Kết nối điện thoại hoặc Cài đặt → Nâng cao → Kết nối điện thoại.','Android: Máy tính → Quét QR. Cho phép camera khi cần; hoặc Dán mã JSON từ Windows.','Đối chiếu tên và mã điện thoại; xác nhận ghép một lần trên Windows. QR ghép có hạn 5 phút.','Giữ ChatAI Windows mở và bật Nhận việc từ điện thoại.','Nhập/nói yêu cầu trong tab Máy tính rồi Gửi việc.'),
'Ví dụ: “Đọc Excel và DXF đã có trong thư mục D:\\DuAn\\MatCatA, kiểm tra tên mặt cắt và đơn vị. Lập bảng dữ liệu đầu vào, tạo file mới và không sửa nguồn.”',
'Windows đang bận, có bản nháp hoặc đang xử lý việc khác có thể chưa nhận ngay. Mất mạng không tự chạy lại tác vụ đã nhận để tránh lặp thao tác.'),
s('6. Theo dõi, bổ sung và dừng',
table(['Trạng thái','Cách xử lý'],['Chờ máy tính','Kiểm tra máy online, ChatAI đang mở và cho phép nhận việc.'],['Đang thực hiện','Đọc tiến trình, không gửi lại cùng việc chỉ vì chưa thấy kết quả.'],['Cần trợ giúp','Bấm Bổ sung, trả lời đúng dữ kiện AI đang cần.'],['Đã tạm dừng','Kiểm tra trạng thái máy rồi Tiếp tục nếu phù hợp.'],['Đã kết thúc','Đọc phản hồi và các giới hạn; chưa đồng nghĩa mô hình kỹ thuật đã được chứng nhận.'],['Mất kết nối / có lỗi','Kiểm tra hội thoại và mô hình Windows trước khi gửi lại.']),
'Tạm dừng/Hủy yêu cầu dừng tại điểm an toàn; không bảo đảm ngắt ngay phần mềm đang tính. Khi PLAXIS Input đang mở, tác vụ mới đề cập PLAXIS bị chặn để tránh thay mô hình.'),
s('7. Yêu cầu ảnh màn hình',
steps('Windows: bật Cho phép điện thoại yêu cầu ảnh màn hình trong phiên này. Mặc định tắt.','Android: Máy tính → Yêu cầu ảnh màn hình.','Chờ ảnh, phóng to để xem; bấm Ẩn ảnh khi không cần.','Tắt quyền hoặc thu hồi điện thoại trên Windows khi ngừng sử dụng.'),
'Ảnh là màn hình chính, có thể chứa dữ liệu của ứng dụng khác. Chỉ chụp khi yêu cầu, không chụp định kỳ; ảnh mã hóa trên server và hết hạn sau 5 phút. Chức năng cần kho ảnh mã hóa đã cấu hình. Ảnh không tự gửi cho nhà cung cấp AI.'),
s('8. Thông báo và quyền truy cập',
steps('Đăng nhập, mở Tài khoản → Bật thông báo điện thoại.','Android 13 trở lên: chấp nhận quyền Thông báo nếu muốn nhận cảnh báo. Nếu từ chối, vẫn xem được mục thông báo trong app.','Dùng Tắt thông báo điện thoại khi cần; đăng xuất cũng hủy công việc nền và thông báo đang hiện.'),
'App kiểm tra thông báo tài khoản bằng WorkManager khoảng 15 phút/lần khi có mạng, Android có thể trì hoãn lúc tiết kiệm pin. Đây chưa phải push tức thời. Những thông báo cũ không tự báo hàng loạt ngay khi bật.',
'Quyền cần thiết: Internet cho kết nối; micro chỉ khi nói; camera chỉ khi quét QR; thông báo chỉ khi bật nhận. Không cần quyền đọc toàn bộ bộ nhớ điện thoại hoặc trợ năng. Chế độ dùng thử không nhận thông báo tài khoản.'),
s('9. Phạm vi hiện chưa có',
'Chưa gửi PDF/Excel/DXF từ điện thoại trực tiếp lên Windows, chưa tải file kết quả, chưa điều khiển chuột/phím trực tiếp, chưa có push tức thời qua FCM hay tự cập nhật APK. Đặt file sẵn vào thư mục Windows đã cấp quyền và lấy file kết quả trên máy tính.',
'Admin Android được hướng dẫn ở 04_Admin.docx. Không phải toàn bộ chức năng quản trị desktop đã có trên Android.')
])

book('04_Admin.docx','QUẢN TRỊ CHAT AI','Tài khoản, thanh toán, thông báo và cấu hình dịch vụ',[
s('1. Phân biệt quyền Admin',
'Đăng nhập tài khoản Admin hệ thống do đơn vị vận hành quản lý. Không có quyền Admin chỉ bằng cách sửa thông tin trên điện thoại; server kiểm tra từng yêu cầu.',
table(['Chức năng','Windows','Android hiện tại'],['Người dùng, token và số dư','Có','Có'],['Khóa/mở khóa, thu hồi phiên, xóa mềm/khôi phục','Có','Có'],['Gửi thông báo, nạp thủ công, sửa bảng giá','Có','Có'],['Ngân hàng và webhook QR','Có','Chưa có giao diện'],['Email/Resend và cấu hình AI','Có','Chưa có giao diện'],['Đào tạo bộ nhớ dùng chung','Có','Chưa có giao diện'])),
s('2. Quản lý người dùng trên Android',
steps('Tài khoản → Quản trị → Người dùng.','Tìm theo tên/email và chọn bộ lọc Hoạt động, Đã khóa, Đã xóa hoặc Tất cả.','Xem token tháng này, số dư và trạng thái; bấm người dùng để xem các tháng và nhật ký.','Chọn Khóa/Mở khóa hoặc Thu hồi phiên theo nhu cầu, đọc và xác nhận thao tác.','Xóa mềm chỉ khi cần; khôi phục nếu server còn cho phép trong thời hạn, hiện là 30 ngày.'),
'Token tính theo giờ Việt Nam, gồm usage API đã ghi nhận. AI trên máy và lượt không trả usage không được tự ước lượng để cộng vào thống kê. Phí tháng là số tiền sổ cái thực tế, không phải tổng token nhân giá mới.'),
s('3. Nạp thủ công',
steps('Xác minh giao dịch hoặc lý do hỗ trợ trước khi nạp.','Android: Quản trị → Thanh toán; Windows: Quản trị thanh toán.','Nhập tên tài khoản, số tiền và lý do. Luồng Admin Android hiện cho 1.000–5.000.000 đ.','Kiểm tra người nhận/số tiền ở hộp xác nhận và thực hiện.','Xem lại số dư và lịch sử để đối chiếu.'),
'Khi mất mạng, Android giữ mã yêu cầu và nội dung trong kho bảo mật. Thử lại đúng nội dung dùng cùng mã, không tự tạo lần nạp mới. Nếu yêu cầu trước chưa rõ kết quả, giải quyết yêu cầu đó trước khi đổi tài khoản/số tiền/lý do.'),
s('4. Phí duy trì và giá token',
steps('Mở Thanh toán, tải lại bảng giá.','Nhập phí duy trì 0–5.000.000 đ/30 ngày. Giá 0 đ là miễn phí duy trì.','Nhập giá DeepSeek 1–10.000.000 đ/triệu token.','Bấm Lưu bảng giá, kiểm tra giá hiển thị trên một tài khoản người dùng.'),
'Ví dụ nếu giá là 4.000 đ/triệu token: 1 triệu token phí 4.000 đ; 5 triệu phí 20.000 đ. Đây chỉ là ví dụ, không thay bảng giá đang lưu trên server. Phí API nhà cung cấp vẫn do chủ API chi trả.',
'Giá thay đổi không hồi tố lượt cũ. Đơn QR đã tạo có thông tin riêng; người dùng phải kiểm tra lại giá và số tiền trước khi thanh toán.'),
s('5. Cấu hình QR tự động trên Windows',
steps('Quản trị thanh toán: nhập mã ngân hàng VietQR, số tài khoản và tên người nhận.','Kết nối ngân hàng với SePay và tạo webhook tiền vào theo URL do ứng dụng hiển thị.','Luồng hiện tại dùng xác thực API Key; đặt cùng khóa trong SePay và cấu hình ChatAI. Header yêu cầu là Authorization: Apikey <khóa>.','Lưu cấu hình, bật QR tự động rồi chuyển khoản thử với nội dung và số tiền chính xác.','Đối chiếu ngân hàng, lịch sử đơn và ví trước khi mở cho người dùng.'),
'Không chỉ chọn HMAC-SHA256 trên SePay nếu Worker hiện chưa xử lý chữ ký đó. Kiểu xác thực phải khớp triển khai. Không gửi khóa webhook hoặc API key vào chat/tài liệu.',
'Webhook phải đúng tiền vào, tài khoản nhận, mã đơn, số tiền và thời hạn. Chuyển sai hoặc quá hạn cần đối soát; không lấy ảnh chụp giao dịch làm xác nhận tự động.'),
s('6. Gửi thông báo',
steps('Android: Quản trị → Thông báo.','Nhập * để gửi tất cả hoặc tên tài khoản để gửi riêng.','Nhập tiêu đề (tối đa 160 ký tự) và nội dung (tối đa 10.000 ký tự).','Đọc xác nhận rồi gửi; kiểm tra bằng tài khoản người nhận.'),
'Thông báo đọc trong ứng dụng; bản Android mới có thể báo trên điện thoại khi người dùng bật quyền và kiểm tra nền. Kiểm tra nền khoảng 15 phút/lần có thể bị trì hoãn, chưa phải push tức thời hoặc gửi email. Ví dụ: “Bảo trì server lúc 22:00–22:30. Vui lòng lưu công việc trước thời gian này.”'),
s('7. Email và AI trên Windows',
'Email: Tài khoản → Cấu hình email. Nhập tên người gửi, email thuộc tên miền đã xác minh trong Resend và API key, lưu rồi gửi thử. Để trống key khi lưu có thể giữ key cũ. Gửi thử được chấp nhận không chứng minh thư đã vào inbox.',
'Địa chỉ onboarding@resend.dev phục vụ thử nghiệm theo giới hạn Resend, không thay tên miền gửi đã xác minh cho khách hàng. Không dùng địa chỉ Gmail bất kỳ làm người gửi nếu chưa đáp ứng quy định dịch vụ.',
'AI: lưu cấu hình nhà cung cấp/mô hình trên server qua giao diện Admin hiện có. Kiểm tra kết nối thành công chưa chứng minh cấu hình đã lưu; lưu xong mở lại kiểm tra. Không đóng gói key vào APK.'),
s('8. Đào tạo bộ nhớ chung trên Windows',
steps('Tài khoản → Bộ nhớ đào tạo AI, bật chế độ chia sẻ bài học mới.','Chạy bài mẫu với nguồn rõ ràng và theo dõi bằng chứng công cụ.','Kiểm tra bài học lỗi/bước sửa/môi trường và mức kiểm chứng.','Thu hồi bài học không phù hợp; tắt đào tạo khi làm dự án riêng.'),
'Không biến lời AI nói “đã thành công” thành bằng chứng đã tính. Chia sẻ cách làm không chia sẻ thông số công trình làm mặc định cho người khác. Không đồng bộ bí mật hoặc dữ liệu riêng vào kho chung.'),
s('9. Theo dõi server',
'Kiểm tra Workers Logs theo mã lỗi, mức dùng D1/KV, trạng thái webhook và chi phí nhà cung cấp. Hạn mức nền tảng vẫn tồn tại dù ứng dụng tăng giới hạn riêng. Khi D1 hết hạn mức đọc, nhiều chức năng tài khoản/ví/AI có thể cùng lỗi; tối ưu truy vấn hoặc điều chỉnh gói sau khi xác định nguyên nhân.')
])

book('05_PLAXIS_GeoStudio.docx','GIAO VIỆC KỸ THUẬT CHO CHAT AI','PLAXIS 2D/3D và GeoStudio • chuẩn bị, thực hiện, kiểm tra',[
s('1. Điều kiện chung',
'ChatAI hỗ trợ tự động hóa trên máy đã có phần mềm và công cụ kết nối. Khả năng phụ thuộc phiên bản, giấy phép, API, giao diện và dữ liệu thực; không bảo đảm mọi bài có thể tự chạy trọn vẹn.',
steps('Sao lưu dữ liệu và mô hình hiện có.','Chuẩn bị đề bài, tài liệu, đơn vị, phiên bản phần mềm và tiêu chí đầu ra.','Cấp thư mục dữ liệu cho ChatAI.','Chuẩn bị dự án mới/bản sao; ghi rõ không ghi đè mô hình đang mở.','Chỉ cho phép tính sau khi đối chiếu đầu vào và phạm vi mô hình.')),
s('2. Mẫu yêu cầu PLAXIS',
'“Dùng PLAXIS 2D [phiên bản], dựng bài [tên/chương] theo manual [file]. Tạo project mới tại [đường dẫn], giữ nguyên mô hình khác. Đọc đủ hình học, vật liệu, nước, kết cấu, tải và các pha. Tự tra API đúng phiên bản, kiểm tra từng bước và đọc kết quả số. Nếu thiếu dữ liệu làm thay đổi mô hình thì hỏi, không tự thay bằng bài mẫu khác.”',
'Với 3D cần đủ kích thước theo các phương và tải/kết cấu 3D. Không lấy bài 2D đùn sang 3D nếu chiều dài hoặc điều kiện chưa có trong đề.'),
s('3. Kết nối PLAXIS',
steps('Mở đúng PLAXIS Input đã cài, không mở file Setup thay phần mềm.','Trong Expert → Configure remote scripting server, cấu hình server theo phiên bản và giấy phép.','Đối chiếu cổng và thông tin xác thực với ChatAI. Input/Output thường dùng 10000/10001 trong cấu hình dự án, nhưng phải kiểm tra thực tế.','Yêu cầu AI đọc project, phiên bản và các đối tượng hiện có trước khi tạo/sửa.','Khi lỗi lệnh, đọc phần đã hoàn thành; không chạy lại toàn bộ batch tạo đối tượng.'),
'Tra API cần cả lệnh gốc lẫn đối tượng con và tài liệu đúng phiên bản. Không kết luận một chức năng hoàn toàn không có chỉ từ một thuộc tính đọc-only. Nếu API không đáp ứng, xem phương án giao diện hoặc hỏi thao tác tối thiểu còn thiếu.'),
s('4. Kiểm tra mô hình trước tính',
table(['Nhóm','Kiểm tra'],['Hình học','Miền đất, mặt cắt, cao độ, chiều dài 3D, vùng đào/đắp và giao nhau'],['Vật liệu','Mô hình đất, dung trọng, độ cứng, c/φ hoặc Su, thoát nước, OCR/POP và đơn vị'],['Nước, biên, tải','Mực nước, điều kiện biên, tải mặt/kết cấu theo đề'],['Kết cấu','Tường, cọc, neo, chống, tunnel/interface/contraction đúng đối tượng'],['Các pha','Start from phase, kích hoạt/tắt và vật liệu của từng pha'],['Lưới','Loại phần tử, mật độ/refinement và vùng quan tâm']),
'Không đổi bài hố đào thành bờ đắp chỉ vì mẫu tự động chỉ hỗ trợ bờ đắp. Nếu đơn giản hóa, phải mô tả thành phần bỏ đi và được người dùng chấp nhận.'),
s('5. Dữ liệu GeoStudio cần chuẩn bị',
steps('Excel số liệu vật liệu: sheet, chỉ tiêu, đơn vị, nguồn thí nghiệm.','Bảng tổng hợp xử lý/BTH: tên mặt cắt, địa tầng, bề dày, hố khoan liên quan.','DXF: đường tự nhiên, đường thiết kế và các layer/đối tượng phân biệt rõ.','GSZ mẫu nếu có: dùng để tham khảo cấu trúc/trình bày; không dùng kết quả cũ thay tính mới.','Điều kiện nước, tải, giai đoạn, phương pháp ổn định và miền tìm trượt.'),
'Mẫu yêu cầu: “Mặt cắt [tên]. Đọc BTH để lấy địa tầng và bề dày; đọc SLTT để lấy chỉ tiêu có địa chỉ ô và đơn vị; đọc DXF để lấy hình học. GSZ chỉ là mẫu trình bày. Tạo bản mới, giữ nguồn, báo dữ kiện thiếu trước khi Solve.”'),
s('6. Địa tầng và Co/Su',
'Khi đề cho bề dày đứng không đổi và chấp nhận lớp song song đường tự nhiên, đường đáy lớp có thể dựng bằng dịch cao độ xuống theo tổng bề dày. Đây là dịch đứng, không phải OFFSET vuông góc; giữ các điểm gãy và kiểm tra vùng khép kín.',
'Với hồ sơ mà người dùng xác nhận Co trong SLTT là Su, cần dùng đúng mô hình không thoát nước và φ=0 khi được chỉ định. Không coi Co luôn là Su cho mọi workbook; phải kiểm tra nguồn thí nghiệm và quy ước dự án. Bài thoát nước dùng c′/φ′ đúng nguồn.',
'Không đổi ô trống, lỗi công thức hoặc thiếu cache thành số 0; không lấy vật liệu GSZ cũ để lấp chỗ thiếu nếu chưa được chấp nhận.'),
s('7. Chạy và đọc kết quả mới',
steps('Xác minh file dự án và analysis sẽ chạy.','Theo dõi trạng thái tính toán đến hoàn tất/lỗi.','Đọc dữ liệu đầu ra của phiên tính mới, ghi thời điểm và nguồn.','PLAXIS: xem chuyển vị, áp lực nước, nội lực và hệ số an toàn phù hợp mục tiêu.','GeoStudio: đọc Fs, mặt trượt nguy hiểm, phương pháp, nước/tải và phạm vi tìm trượt.','Lưu project mới, bảng đầu vào, log và báo cáo; đối chiếu lại dữ liệu nguồn.'),
'Giá trị Fs hữu hạn/dương không tự chứng minh hội tụ, mô hình đúng hoặc miền tìm trượt đủ rộng. Người phụ trách chuyên môn cần thẩm tra trước dùng cho thiết kế.'),
s('8. Tiếp tục sau lỗi',
'Giữ cùng hội thoại, nêu lỗi và yêu cầu đọc lại trạng thái thật. Các bước đã thực hiện, lỗi và bước chưa rõ phải tách biệt. Khi kết nối mất, không giả định chưa chạy rồi tạo lại đối tượng. AI nên tự tra cách xử lý; chỉ hỏi người dùng khi thiếu dữ liệu chuyên môn, có nhiều lựa chọn ảnh hưởng kết quả hoặc công cụ không thể thực hiện bước cần thiết.')
])

book('06_Xu_ly_loi.docx','XỬ LÝ LỖI VÀ CẬP NHẬT','Tra cứu nhanh cho người dùng và người hỗ trợ',[
s('1. Cách báo lỗi có ích',
steps('Ghi bản ứng dụng, Windows/Android và thời điểm lỗi.','Nêu thao tác ngay trước lỗi và điều mong đợi.','Chụp thông báo hoặc sao chép nguyên văn lỗi.','Cung cấp log liên quan đã che dữ liệu riêng/khóa bí mật.','Với mô hình kỹ thuật, nêu project/analysis và các bước đã chạy; không tự chạy lại toàn bộ trước khi kiểm tra.'),
'Không gửi mật khẩu, API key, khóa webhook, file login.dpapi, cấu hình chứa bí mật hoặc toàn bộ dữ liệu cá nhân khi không cần.'),
s('2. Đăng nhập và cấu hình AI',
table(['Hiện tượng','Kiểm tra / xử lý'],['Đã đăng nhập nhưng app lại hỏi','Kiểm tra phiên, mạng và server; không nhập thông tin vào cửa sổ lạ. Gửi log nếu lặp.'],['Email không đăng nhập được','Dùng tên đăng nhập; kiểm tra email có được xác minh/được server chấp nhận.'],['Dịch vụ AI không hợp lệ','Cập nhật đúng bản chạy; đối chiếu mô hình/nhà cung cấp trong cấu hình và crash.log.'],['Kiểm tra server thành công nhưng không lưu','Lưu cấu hình, mở lại kiểm tra; lấy lỗi lưu cụ thể. Kiểm tra kết nối không thay thao tác lưu.'],['AI trả rỗng / quá tải','Kiểm tra provider, model, usage và mã lỗi; không gửi lại tính phí liên tục khi chưa rõ kết quả.'])),
s('3. QR và số dư',
table(['Hiện tượng','Cách xử lý'],['Đổi mệnh giá nhưng QR vẫn cũ','Kiểm tra số tiền của đơn và nội dung QR; không chuyển theo ô nhập nếu đơn còn là số tiền khác.'],['Đã chuyển nhưng chưa cộng','Đối chiếu đúng người nhận, số tiền, nội dung và hạn; Admin xem webhook/giao dịch.'],['Phí duy trì đã đặt 0 nhưng hiển thị giá cũ','Làm mới ví, cập nhật đúng client/Worker; không tạo đơn gia hạn theo nhãn cũ.'],['App chậm khi mở thanh toán','Giữ một cửa sổ/đơn, chờ phản hồi; gửi log nếu lặp hoặc tự tắt.'],['Số dư đang giữ chỗ','Admin đối soát usage thật; không tự xác nhận 0 token khi chưa có bằng chứng.'])),
s('4. Giọng nói và kết nối điện thoại',
'Không nghe: kiểm tra quyền Micro, dịch vụ nhận dạng tiếng Việt và mạng. Không đọc: kiểm tra giọng TTS tiếng Việt. Từ chối quyền vẫn có thể nhập chữ.',
'Máy tính offline: mở ChatAI Windows, đăng nhập đúng cùng tài khoản, kiểm tra Internet và bật nhận việc. QR ghép hết hạn: tạo lại. Không xem ảnh: kiểm tra quyền chụp phiên Windows và kho ảnh mã hóa server.',
'Ứng dụng mất kết nối khi đang tính: kiểm tra máy thật và hội thoại trước khi gửi lại. Không tự lặp lệnh tạo mô hình.'),
s('5. Cập nhật Git trong PowerShell',
'Chạy từng lệnh sau tại thư mục dự án của bạn, thay đường dẫn nếu khác:',
steps('cd "G:\\My Drive\\Dev\\Chat-AI\\Chat AI 2.6.6"','git status','git switch main','git pull origin main','.\\run.bat'),
'Nếu git status có thay đổi cục bộ, sao lưu/commit trước. Khi cần stash gồm file mới, dùng git stash push -u -m "Backup truoc cap nhat"; giữ bản sao riêng nếu dữ liệu quan trọng và xử lý xung đột khi khôi phục.',
'Nếu báo ORIG_HEAD.lock/packed-refs.lock: đóng GitHub Desktop và tiến trình Git đang thao tác, kiểm tra Get-Process git* -ErrorAction SilentlyContinue. Chỉ khi xác định không còn thao tác Git mới xóa đúng file khóa cũ theo thông báo, ví dụ Remove-Item ".git\\ORIG_HEAD.lock". Không xóa file ORIG_HEAD/packed-refs hoặc toàn bộ .git.',
'Kho trên Google Drive có thể bị đồng bộ khóa; nếu tái diễn, cân nhắc checkout trên ổ cục bộ ngoài thư mục đồng bộ.'),
s('6. APK build hoặc cài lỗi',
steps('GitHub Actions → Build Android APK → lượt cần kiểm tra → job android.','Mở bước có dấu đỏ, đọc lỗi cụ thể; dòng info không nhất thiết làm build thất bại.','Nếu Analyze and test lỗi, gửi tên kiểm thử và chi tiết exception.','Nếu Build installable test APK lỗi, gửi lỗi Gradle/SDK/resource.','Chỉ tải Artifact sau lượt thành công; cài APK mới không có nghĩa server đã được triển khai cùng bản.'),
'Build vàng/In progress nghĩa đang chạy. Lần đầu có thể lâu do tải SDK và thư viện. Build xanh xác nhận tạo gói, không thay kiểm tra trên điện thoại thật.'),
s('7. Lấy log Windows',
'Đứng đúng thư mục ứng dụng; đường dẫn tương đối dùng thư mục hiện tại, không phải thư mục mà bạn muốn.',
steps('Get-Content ".\\data\\crash.log" -Encoding UTF8 -Tail 50','Get-Content ".\\data\\startup.log" -Encoding UTF8 -Tail 50','Get-Content ".\\build-logs\\runtime.log" -Tail 30 -Wait'),
'Lệnh cuối chỉ dùng khi thực sự có log build runtime; Ctrl+C để ngừng theo dõi. Nếu đang ở C:\\Windows\\system32 thì chuyển về thư mục dự án trước. Log có thể chứa tên file hoặc nội dung riêng, cần xem trước khi gửi.'),
s('8. Lỗi server và hạn mức Cloudflare',
'Ghi endpoint và mã lỗi, Admin tìm trong Workers Logs. D1 vượt hạn mức đọc có thể làm nhiều endpoint cùng lỗi. Thời gian reset theo UTC trong thông báo; đổi sang giờ Việt Nam bằng cộng 7 giờ.',
'Kiểm tra truy vấn, index, polling và gói dịch vụ. Nâng gói không tự sửa vòng gọi lặp hoặc truy vấn đọc quá nhiều. Không tạo lại database hay xóa ví/lịch sử chỉ để thử sửa lỗi.'),
s('9. Checklist trước khi đưa cho người dùng rộng rãi',
steps('Đăng ký/đăng nhập/đăng xuất và đổi tài khoản không lẫn dữ liệu.','Chat NVIDIA, DeepSeek thiếu/đủ số dư; token và phí đúng bảng giá.','QR và webhook thật, thử gửi lại không cộng hai lần.','Ghép đúng/sai tài khoản, thu hồi điện thoại, dừng/tiếp tục và mất mạng.','Logo, bàn phím, chữ lớn, quyền micro/camera và khởi động trên điện thoại thật.','Bản sao dự án kỹ thuật giữ nguyên nguồn và đọc đúng kết quả mới.'))
])

for filename,title,subtitle,sections in BOOKS:
 d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
 sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7)
 for level,size in [(1,17),(2,13)]:
  st=d.styles[f'Heading {level}'];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('2459A6')
  st.paragraph_format.keep_with_next=True
 header=sec.header.paragraphs[0];header.text='CHAT AI  |  HƯỚNG DẪN SỬ DỤNG';header.style='Caption'
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 footer.add_run('ChatAI • 11/10/2026  |  Trang ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ROOT/'logo_chat_ai.png'),width=Cm(3))
 p=d.add_paragraph(title);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.style='Title'
 p=d.add_paragraph(subtitle);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=d.add_paragraph('Bộ tài liệu sử dụng • cập nhật 11/10/2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 d.add_paragraph('Phạm vi: Desktop 2.6.6 và Android thử nghiệm 0.8; bao gồm luồng dùng thử mới trong mã nguồn hướng dẫn. Tên nút có thể khác nhẹ theo bản cài. Chức năng cần client và server tương thích; chỉ dẫn dưới đây không chứng nhận đã kiểm thử mọi môi trường thực tế.')
 d.add_heading('Tra cứu trong tài liệu',level=1)
 for heading,_ in sections:d.add_paragraph(heading)
 d.add_page_break()
 for heading,items in sections:
  d.add_heading(heading,level=1)
  for item in items:
   if isinstance(item,str):d.add_paragraph(item)
   elif item[0]=='steps':
    for n,t in enumerate(item[1],1):d.add_paragraph(f'{n}. {t}')
   elif item[0]=='table':
    t=d.add_table(rows=1,cols=len(item[1]));t.style='Light Shading Accent 1'
    for i,v in enumerate(item[1]):t.rows[0].cells[i].text=v
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in item[2]:
     cells=t.add_row().cells
     for i,v in enumerate(row):cells[i].text=v
     trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    d.add_paragraph()
 d.add_heading('Nguồn đối chiếu',level=1)
 d.add_paragraph('README dự án; mobile/README.md và giao diện Android; docs/phone-connection.md; docs/payment-setup.md; docs/email-setup.md; docs/ai-training-memory.md; docs/plaxis-3d.md; docs/geoslope-workflow.md. Tài liệu tổng hợp cách sử dụng, không chép thông số hoặc kết quả của một công trình làm mặc định cho công trình khác.')
 d.core_properties.title=title;d.core_properties.author='ChatAI';d.core_properties.subject=subtitle
 d.save(OUT/filename)
with zipfile.ZipFile(OUT/'Bo_huong_dan_ChatAI_2026-10-11.zip','w',zipfile.ZIP_DEFLATED) as z:
 for filename,*_ in BOOKS:z.write(OUT/filename,arcname=filename)
print('Created',len(BOOKS),'Word documents and ZIP')
