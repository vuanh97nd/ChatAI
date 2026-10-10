"""Build separate illustrated technical user guides (schematics, not screenshots)."""
from pathlib import Path
import math, zipfile
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/huong-dan-word';IMG=OUT/'hinh-ky-thuat';IMG.mkdir(exist_ok=True)
plt.rcParams['font.family']='DejaVu Sans'

def flow(filename,title,labels):
 fig,ax=plt.subplots(figsize=(11,6));ax.set_xlim(0,12);ax.set_ylim(0,7);ax.axis('off')
 ax.text(6,6.7,title,ha='center',fontsize=17,color='#2459a6',weight='bold')
 positions=[(.4,4.1),(3.3,4.1),(6.2,4.1),(9.1,4.1),(9.1,1.4),(6.2,1.4),(3.3,1.4),(.4,1.4)]
 for i,(label,(x,y)) in enumerate(zip(labels,positions)):
  ax.add_patch(FancyBboxPatch((x,y),2.4,1.5,boxstyle='round,pad=.13',facecolor='#eff4ff',edgecolor='#426bb5',linewidth=1.5))
  ax.text(x+.2,y+1.15,str(i+1),fontsize=14,weight='bold',color='#426bb5')
  ax.text(x+1.2,y+.65,label,ha='center',va='center',fontsize=11)
  if i<7:
   nx,ny=positions[i+1]
   if y==ny:
    start=(x+2.55,y+.75) if nx>x else (x-.15,y+.75);end=(nx-.2,ny+.75) if nx>x else (nx+2.6,ny+.75)
   else:start=(x+1.2,y-.1);end=(nx+1.2,ny+1.7)
   ax.annotate('',xy=end,xytext=start,arrowprops={'arrowstyle':'->','color':'#64748b','lw':1.5})
 ax.text(6,.45,'SƠ ĐỒ MINH HỌA — KHÔNG PHẢI ẢNH CHỤP GIAO DIỆN',ha='center',fontsize=10,color='#64748b')
 fig.savefig(IMG/filename,dpi=160,bbox_inches='tight');plt.close(fig)
flow('plaxis2d_quy_trinh.png','PLAXIS 2D 2024.2 — trình tự làm việc với ChatAI',['Đọc đề / manual\nChọn Plane strain…','Kết nối Input\nĐọc project hiện tại','Hình học\nĐịa tầng / kết cấu','Vật liệu / nước\nĐối chiếu đơn vị','Mesh\nKiểm tra độ mịn','Staged construction\nKiểm tra từng phase','Calculate\nTheo dõi trạng thái','Output\nUy / nội lực / ΣMsf'])
flow('plaxis3d_quy_trinh.png','PLAXIS 3D 2024 — trình tự làm việc với ChatAI',['Đề bài 3D\nĐủ X, Y, Z','Kết nối Input\nChuẩn bị project mới','Soil volume\nBorehole / địa tầng','Structures\nTải / nước / vật liệu','Mesh 3D\nRefinement','Staged construction\nThi công từng pha','Calculate\nKiểm tra hoàn tất','Output\nUz / Ux / Uy / nội lực'])
flow('geostudio_quy_trinh.png','GeoStudio 2025.1.1 / SLOPE/W — trình tự dữ liệu',['BTH xử lý\nChọn đúng mặt cắt','SLTT\nChỉ tiêu / đơn vị','DXF\nTự nhiên / thiết kế','GSZ mẫu\nChỉ tham khảo cấu trúc','Vùng đất mới\nGán vật liệu','Nước / tải\nMiền tìm trượt','Solve\nTheo dõi phiên mới','Kết quả\nFs / mặt trượt / log'])
flow('kiem_chung.png','Các mức bằng chứng — không nhảy qua bước',['Nguồn đầu vào\nTên file / sheet / ô','Lệnh đọc / tạo\nKết quả công cụ','Đối tượng hiện có\nKhông tạo trùng','Đối chiếu mô hình\nHình học / vật liệu','Pha / analysis\nNước / tải / biên','Tính hoàn tất?\nLỗi / hội tụ / cảnh báo','Đọc kết quả mới\nProject / thời điểm','Thẩm tra chuyên môn\nTrước dùng thiết kế'])
# Conceptual diagram: vertical thickness, not a real DXF or design profile.
fig,ax=plt.subplots(figsize=(10,4.6));x=[0,2,4,6,8,10];y=[2,2.4,2.1,2.6,2.2,2]
colors=['#e7cba2','#b9d4a4','#b4c6e7']
for k,c in enumerate(colors):
 upper=[a-k for a in y];lower=[a-k-1 for a in y]
 ax.fill_between(x,upper,lower,color=c,edgecolor='#475569',linewidth=1)
 ax.text(8.1,upper[4]-.6,f'Lớp {k+1}',fontsize=11)
ax.plot(x,y,color='#2563eb',lw=2,label='Đường tự nhiên')
ax.annotate('',xy=(4,y[2]-2),xytext=(4,y[2]),arrowprops={'arrowstyle':'<->','lw':1.5,'color':'#9d174d'})
ax.text(4.25,.95,'Bề dày theo phương đứng\nKhông OFFSET vuông góc',fontsize=11,color='#9d174d')
ax.set_title('Minh họa địa tầng dịch đứng theo đường tự nhiên',fontsize=15);ax.axis('off');ax.legend(loc='upper left')
fig.text(.5,.03,'Hình khái niệm; không dùng hình hoặc tọa độ này làm dữ liệu công trình.',ha='center',fontsize=10)
fig.savefig(IMG/'geostudio_dia_tang.png',dpi=160,bbox_inches='tight');plt.close(fig)
fig=plt.figure(figsize=(9,5));ax=fig.add_subplot(projection='3d')
for z in [0,-1]:
 ax.plot([0,1,1,0,0],[0,0,1,1,0],[z]*5,color='#2459a6')
for a,b in [(0,0),(1,0),(1,1),(0,1)]:ax.plot([a,a],[b,b],[-1,0],color='#2459a6')
ax.set_xlabel('X: kích thước ngang');ax.set_ylabel('Y: kích thước còn lại');ax.set_zlabel('Z: cao độ')
ax.set_xticks([]);ax.set_yticks([]);ax.set_zticks([]);ax.set_title('PLAXIS 3D: kiểm tra đủ ba phương và đơn vị')
fig.text(.5,.015,'Sơ đồ trục quy ước; không phải hình học dự án hay ảnh PLAXIS.',ha='center',fontsize=10)
fig.savefig(IMG/'plaxis3d_truc.png',dpi=160,bbox_inches='tight');plt.close(fig)

def newdoc(title,version):
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2)
 d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(11)
 d.styles['Normal'].paragraph_format.space_after=Pt(7)
 for name,size in [('Heading 1',16),('Heading 2',13)]:
  d.styles[name].font.name='Calibri';d.styles[name].font.size=Pt(size);d.styles[name].font.color.rgb=RGBColor.from_string('2459A6')
 p=sec.header.paragraphs[0];p.text='CHAT AI • HƯỚNG DẪN KỸ THUẬT';p.style='Caption'
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('11/10/2026 | Trang ')
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ROOT/'logo_chat_ai.png'),width=Cm(2.5))
 d.add_heading(title,0);d.add_paragraph(version)
 d.add_paragraph('Hướng dẫn sử dụng ChatAI để giao việc, theo dõi và kiểm tra kết quả. Không thay manual phần mềm, hồ sơ khảo sát hoặc thẩm tra thiết kế.')
 d.add_paragraph('Lưu ý về hình: bản này có sơ đồ quy trình/khái niệm tự vẽ, không phải ảnh chụp giao diện. Chưa bổ sung ảnh thao tác thực tế vì nguồn manual trả HTTP 403 và môi trường soạn không có phần mềm Windows. Các bước có ghi tên chức năng cần đối chiếu manual cài cùng phiên bản.')
 return d

def sec(d,title,*paras):
 d.add_heading(title,1)
 for p in paras:d.add_paragraph(p)
def steps(d,*items):
 for n,x in enumerate(items,1):d.add_paragraph(f'{n}. {x}')
def prompt(d,title,text):
 d.add_heading(title,2);p=d.add_paragraph(text);p.paragraph_format.left_indent=Cm(.4)
def tbl(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
  t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 d.add_paragraph()
def figure(d,name,caption):
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(IMG/name),width=Cm(16))
 p=d.add_paragraph(caption);p.style='Caption'
def common_start(d,app):
 sec(d,'1. Chuẩn bị máy, tài khoản và dự án',f'Phần mềm {app} phải cài trên Windows, có giấy phép và mở đúng chương trình Input/GeoStudio, không dùng file Setup thay chương trình. ChatAI phải có công cụ tương ứng và được cấp quyền vào thư mục dữ liệu.')
 steps(d,'Sao lưu dự án đang mở và tài liệu gốc.','Tạo thư mục Nguon, LamViec, KetQua và NhatKy riêng cho bài mới.','Kiểm tra file Google Drive đã có sẵn trên máy, không chỉ là tệp trực tuyến.','Trong ChatAI, cấp thư mục dự án và cấu hình phần mềm cần dùng.','Đăng nhập server, chọn AI phù hợp. Đưa manual/đề bài và đường dẫn dữ liệu vào cùng hội thoại.')
 sec(d,'2. Thống nhất phạm vi trước khi thực hiện','Nêu rõ tên bài/mặt cắt, phiên bản, đơn vị, nguồn đầu vào, vị trí lưu project mới và các kết quả cần xuất. Cho phép AI tự đọc nguồn và tra API; chỉ hỏi khi thiếu dữ kiện ảnh hưởng bài toán hoặc cần can thiệp mà công cụ không làm được.')
 prompt(d,'Câu mở đầu nên dùng',f'“Thực hiện bài [tên] bằng {app}, lưu project mới tại [đường dẫn]. Dữ liệu nguồn [danh sách]. Giữ nguyên nguồn và mô hình đang mở. Tự tra cú pháp đúng phiên bản, báo từng bước đã kiểm tra, chỉ hỏi khi thiếu dữ kiện cần thiết. Không thay bài bằng mẫu khác.”')
def common_end(d,app):
 sec(d,'11. Lưu hồ sơ bàn giao')
 tbl(d,['Đầu ra','Nội dung cần có'],[['Project','File mới mở được, tên analysis/phase rõ'],['Bảng đầu vào','Thông số, đơn vị, nguồn file/sheet/ô/trang'],['Nhật ký','Các bước đã làm, lỗi và bước sửa có bằng chứng'],['Kết quả','Bảng số, đơn vị, analysis/phase, thời điểm và trạng thái tính'],['Hình kết quả','Đúng biến, đúng hệ số phóng đại và đúng phiên tính'],['Giới hạn','Dữ kiện thiếu, giả thiết được chấp nhận, cảnh báo và việc cần thẩm tra']])
 prompt(d,'Yêu cầu bàn giao','“Lập báo cáo: mục tiêu, dữ liệu nguồn, hình học, vật liệu, nước/tải/biên, các pha/analysis, trạng thái tính và kết quả mới. Nêu bước chưa xác minh. Cho đường dẫn project, log và file kết quả; không báo hoàn tất chỉ vì mở được phần mềm.”')
 sec(d,'12. Khi AI dừng hoặc lỗi')
 tbl(d,['Hiện tượng','Cách tiếp tục'],[['Không kết nối được','Kiểm tra chương trình, giấy phép, server/cổng và bản đang chạy; không tự dựng lại.'],['Sai cú pháp hoặc thuộc tính không có','Tra đối tượng con và tài liệu đúng phiên bản; không lặp lệnh lỗi y nguyên.'],['Một batch chạy dở','Đọc các đối tượng đã tạo và bước chưa rõ, tiếp tục từ trạng thái thật.'],['Bị giới hạn ở bài mẫu','Giữ mục tiêu bài, tra API/giao diện khác; không âm thầm bỏ cấu kiện.'],['Đã tính nhưng không đọc được kết quả','Kiểm tra analysis/phase/file Output và trạng thái hoàn tất, không lấy kết quả mẫu.'],['Mất mạng khi đang chạy','Kiểm tra phần mềm và hội thoại trên Windows trước khi gửi lại.']])
 sec(d,'13. Dùng từ điện thoại','Đăng nhập cùng tài khoản, ghép QR trong Windows → Tài khoản → Kết nối điện thoại và Android → Máy tính. Windows phải đang mở, nhận việc và có dữ liệu sẵn trong thư mục được cấp. Android hiện chưa gửi file trực tiếp lên Windows hoặc tải file kết quả.',
'Khi PLAXIS Input đang mở, receiver điện thoại chặn tác vụ mới có tên PLAXIS để tránh thay mô hình; thực hiện công việc PLAXIS trên desktop hoặc chuẩn bị phiên riêng. Tạm dừng/hủy chờ điểm dừng an toàn, không chắc ngắt ngay bộ tính.')
 sec(d,'14. Checklist thẩm tra trước dùng kết quả')
 steps(d,'Đúng tên công trình, mặt cắt, phiên bản và đơn vị.','Không lấy vật liệu/tải/nước/kết quả mẫu thay đầu vào thật.','Đúng project và phiên tính mới; không đọc nhầm file cũ.','Kiểm tra lưới hoặc miền tìm trượt, cảnh báo và độ nhạy.','Người có chuyên môn đối chiếu với mục tiêu/tiêu chuẩn thiết kế.')
 figure(d,'kiem_chung.png','Hình: chuỗi bằng chứng cần kiểm tra; sơ đồ minh họa, không phải chứng nhận đúng thiết kế.')
 sec(d,'15. Ảnh thực tế cần bổ sung để thành hướng dẫn bấm nút',f'Cần ảnh chụp {app} đúng phiên bản hoặc manual cài theo chương trình: cửa sổ khởi tạo, vật liệu, nước/tải, pha/analysis, tính và kết quả. Có thể che dữ liệu riêng. Khi có nguồn, bổ sung ảnh đánh dấu từng control; không đoán vị trí nút hoặc dùng giao diện giả.')
 sec(d,'Nguồn đối chiếu','Mã nguồn và tài liệu ChatAI: docs/plaxis-3d.md, docs/geoslope-workflow.md, docs/phone-connection.md; phiên bản do người dùng xác nhận. Manual chính thức chưa tải được trong phiên soạn. Các thông số không được trình bày như giá trị mặc định thiết kế.')

for dim,version,filename in [('2D','2024.2','07_PLAXIS_2D_2024_2.docx'),('3D','2024','08_PLAXIS_3D_2024.docx')]:
 app=f'PLAXIS {dim} {version}';d=newdoc(f'CHAT AI VỚI {app}',f'Windows • {app} • hướng dẫn theo quy trình')
 common_start(d,app);figure(d,f'plaxis{dim.lower()}_quy_trinh.png','Hình 1. Trình tự giao việc và kiểm tra; số trên sơ đồ là thứ tự bước, không phải số nút giao diện.')
 sec(d,'3. Mở và kiểm tra Remote Scripting Server')
 steps(d,'Mở đúng PLAXIS Input và kiểm tra phiên bản ở thông tin chương trình.','Expert → Configure remote scripting server: cấu hình theo giấy phép và phiên bản đang dùng.','Đối chiếu host/cổng/thông tin xác thực của ChatAI; không đưa mật khẩu vào báo cáo hoặc ảnh chia sẻ.','Yêu cầu công cụ đọc Project, phiên bản và các đối tượng đang có. Nếu khác dự án mong muốn thì dừng thay đổi.','Output có kết nối riêng; kiểm tra cổng thực tế. Input/Output thường dùng 10000/10001 trong cấu hình dự án, không bắt buộc mọi máy như vậy.')
 prompt(d,'Kiểm tra kết nối','“Chỉ đọc: kiểm tra Input/Output, phiên bản, Project và đối tượng hiện có. Không tạo project mới, không xóa hay sửa mô hình. Báo đúng cổng và trạng thái kết nối; che thông tin xác thực.”')
 sec(d,'4. Đọc đề và lập bảng đầu vào','Yêu cầu AI liệt kê chương/trang đã đọc. Với PDF scan, xác minh OCR đủ trang; đoạn thiếu không được tự điền. Lập bảng giữ nguồn và đơn vị trước khi tạo hình học.')
 tbl(d,['Nhóm','Dữ kiện cần','Nguồn phải ghi'],[['Dự án','Đơn vị, kiểu mô hình, miền tính','Trang/điều khoản'],['Địa tầng','Cao độ, bề dày, borehole, nước','Hố khoan/bảng'],['Vật liệu','Mô hình, drainage, γ, độ cứng, sức chống cắt','Bảng và đơn vị'],['Kết cấu/tải','Vị trí, độ cứng, tải và liên kết','Hình/bảng'],['Thi công','Thứ tự, start phase, trạng thái kết cấu/nước','Mục pha tính']])
 if dim=='2D':
  sec(d,'5. Dựng hình học 2D','Chọn Plane strain hoặc Axisymmetric theo đề, không tự chọn chỉ để khớp mẫu. Xác định trục X và Y/cao độ; kiểm tra biên miền và đơn vị trước khi tạo điểm/đường/vùng.')
  steps(d,'Đọc tọa độ miền và địa tầng từ nguồn.','Tạo borehole/lớp hoặc hình học theo bài; kiểm tra cao độ và vùng khép kín.','Dựng cấu kiện đúng loại: plate, embedded beam, anchor/interface… theo yêu cầu và API thực.','Đối với tunnel: kiểm tra Tunnel designer, mặt cắt, bán kính/offset, plate, interface và contraction theo manual bài đang làm.','Sau mỗi nhóm lệnh, đọc lại số đối tượng, vị trí và tên; không coi generatetunnel trả OK là đã có mặt cắt đúng.')
  prompt(d,'Kiểm tra tunnel khi gặp giới hạn','“Đọc Tunnel và CrossSection cùng các đối tượng con, đối chiếu lệnh trong reference đúng phiên bản. Kiểm tra segment và hình đã sinh thật. Nếu chưa có mặt cắt, chưa chia lưới/tính; tìm cách thực hiện phù hợp, không kết luận toàn API không hỗ trợ chỉ từ một thuộc tính read-only.”')
 else:
  sec(d,'5. Dựng hình học 3D','Kiểm tra đủ X, Y và Z; Z là phương cao độ/chuyển vị đứng trong quy ước PLAXIS 3D. Bài 2D không xác định chiều dài đùn sang 3D. Cần kích thước hố/bờ đắp, miền xa biên và vị trí kết cấu đúng bài.')
  steps(d,'Đọc kích thước mặt bằng và cao độ.','Tạo borehole/soil volume theo phân lớp.','Dựng surface/line/volume và gán cấu kiện theo API hoặc giao diện có kiểm chứng.','Với hố đào, giữ tường, strut, neo, tải mặt, nước và giai đoạn theo đề; không đổi thành embankment.','Đọc lại tọa độ, mặt/vùng và liên kết để kiểm tra giao nhau, thể tích và vật liệu.')
  figure(d,'plaxis3d_truc.png','Hình 2. Trục X/Y/Z cần kiểm tra; sơ đồ không phải mô hình công trình.')
 sec(d,'6. Gán vật liệu, nước và điều kiện biên','Phân biệt E của Mohr–Coulomb với E50/Eoed/Eur của Hardening Soil. Nếu công cụ bài mẫu chỉ nhận E, không mặc nhiên thay mô hình đất của đề bằng Mohr–Coulomb. γ phải đúng dung trọng; c/φ và Su phải đúng điều kiện thoát nước.')
 steps(d,'Tạo bộ vật liệu có tên riêng và nguồn rõ.','Đọc lại SoilModel, DrainageType và các tham số chính.','Gán đúng từng soil cluster/cấu kiện, không chỉ tạo vật liệu trong thư viện.','Kiểm tra nước, head, pore pressure và tải đúng pha.','Lập bảng chênh lệch đầu vào–mô hình; thiếu thông số thì hỏi thay vì đoán.')
 prompt(d,'Đối chiếu mô hình','“Đối chiếu hình học, vật liệu đã gán, nước/tải/biên với bảng nguồn. Báo thuộc tính nào được kiểm chứng, thuộc tính nào chưa. Không ghi ‘đúng manual’ cho toàn bài nếu chỉ kiểm tra tên vật liệu.”')
 sec(d,'7. Chia lưới','Chuyển sang chế độ Mesh khi hình học đã đủ. Mật độ và refinement theo đề/vùng quan tâm. Không áp một mức Fine mặc định cho mọi bài; cần xem độ nhạy ở vùng ứng suất/chuyển vị lớn.')
 steps(d,'Kiểm tra hình học trước Generate mesh.','Đặt mật độ/refinement phù hợp, tạo lưới.','Đọc cảnh báo, kiểm tra khu vực tunnel/chân tường/đáy đào/tiếp xúc.','Lưu lưới và mô hình; ghi lựa chọn đã dùng.')
 sec(d,'8. Các pha thi công và kích hoạt','Lập bảng phase trước khi tạo. Start from phase quyết định nhánh tính; không nối pha chỉ vì đứng sau trong danh sách. Tra cách activate/deactivate và thuộc tính theo phase đúng API, không lặp set .Active khi đối tượng không có thuộc tính đó.')
 tbl(d,['Pha','Bắt đầu từ','Đất/kết cấu','Tải/nước','Calculation type'],[['Initial','Theo đề','Trạng thái ban đầu','Theo đề','Theo manual'],['Phase 1…n','Tên pha nguồn rõ','Kích hoạt/tắt đúng thi công','Thay đổi có nguồn','Theo đề'],['Safety nếu cần','Trạng thái cần kiểm tra','Giữ đúng mô hình nhánh','Theo giả thiết','Safety']])
 prompt(d,'Kiểm tra pha trước Calculate','“Đọc lại từng pha: tên, Start from phase, kiểu tính, vật liệu, đất đào/đắp, kết cấu, tải/nước. Liệt kê khác biệt với bảng phase. Chỉ Calculate sau khi không còn lỗi đầu vào làm đổi mục tiêu.”')
 sec(d,'9. Chạy tính và theo dõi')
 steps(d,'Lưu project mới và ghi tên phiên/pha sẽ tính.','Chạy Calculate, theo dõi trạng thái từng pha.','Nếu pha lỗi, ghi thông báo và nguyên nhân có bằng chứng; không tự tăng độ cứng hoặc đổi số liệu để ép hội tụ.','Kiểm tra các pha thực sự hoàn tất; ghi cả cảnh báo.','Mở Output của đúng dự án và đúng pha mới tính.')
 sec(d,'10. Đọc và giải thích kết quả')
 vertical='Uy' if dim=='2D' else 'Uz';horizontal='Ux' if dim=='2D' else 'Ux và Uy tại cùng điểm/nút'
 tbl(d,['Mục','Cần báo'],[['Chuyển vị',f'Phương đứng {vertical}, phương ngang {horizontal}; đơn vị, dấu và vị trí'],['Kết cấu','Nội lực/chuyển vị theo cấu kiện và phase'],['Nước','Áp lực nước phù hợp bài, phân biệt tổng và dư'],['An toàn','ΣMsf khi Safety được thực hiện phù hợp; không coi mọi pha đều cho Fs'],['Min/max','Phạm vi đối tượng/nút và phase đang tổng hợp']])
 prompt(d,'Xuất kết quả số',f'“Đọc kết quả pha [tên] của dự án mới. Tổng hợp {vertical}, {horizontal}, nội lực và hệ số an toàn nếu có; ghi điểm/vị trí, đơn vị, trạng thái tính và nguồn. Phân biệt hình biến dạng phóng đại với chuyển vị thật.”')
 common_end(d,app);d.save(OUT/filename)

d=newdoc('CHAT AI VỚI GEOSTUDIO / SLOPE/W','GeoStudio 2025.1.1 • Excel/BTH → DXF → mô hình → Solve')
common_start(d,'GeoStudio 2025.1.1 / SLOPE/W');figure(d,'geostudio_quy_trinh.png','Hình 1. Thứ tự nguồn dữ liệu; sơ đồ minh họa.')
sec(d,'3. Chọn mặt cắt từ bảng tổng hợp xử lý','BTH/THXL quyết định mặt cắt và phân lớp. Tìm theo tên/lý trình, không mặc định cùng một hàng cho mọi file. Đọc hố khoan, mã lớp, thứ tự/bề dày, cao độ, phương án xử lý và giai đoạn.')
steps(d,'Liệt kê sheet và tìm nhãn tên mặt cắt.','Đọc hàng/cột liên quan, kể cả ô gộp và chú thích.','Lập bảng mã lớp–bề dày–nguồn ô–đơn vị.','Kiểm tra lớp trống, lớp có bề dày 0 và tổng bề dày; không tự lấy 0 thay ô lỗi.','So sánh tên mặt cắt với DXF và GSZ mẫu, báo nếu không khớp.')
prompt(d,'Lệnh đọc BTH','“Đọc [file BTH], tìm mặt cắt [tên]. Liệt kê địa tầng theo thứ tự, bề dày, hố khoan và địa chỉ ô. Không lấy Fs có sẵn làm kết quả tính mới. Nếu có nhiều dòng cùng tên thì cho biết khác biệt.”')
sec(d,'4. Đọc SLTT và ghép vật liệu','Ghép mã/tên lớp của BTH với bảng chỉ tiêu, không dùng cột gần nhất để lấp thiếu. Ghi file/sheet/ô, giá trị và đơn vị. Ô công thức cần giá trị được xác nhận, không tự đổi ô thiếu cache thành 0.')
tbl(d,['Mã lớp','γ','c′/φ′ hoặc Su','Nguồn','Mô hình/điều kiện'],[['[mã]','[giá trị, kN/m³]','[chỉ tiêu, kPa/độ]','[sheet!ô]','[theo đề]']])
steps(d,'Kiểm tra đơn vị dung trọng và cường độ.','Xác minh Co là Su hay c theo loại thí nghiệm/quy ước người dùng.','Nếu dùng Su với bài không thoát nước φ=0, giữ đúng mô hình tương ứng; không cộng φ khác.','Nếu dùng c′/φ′, kiểm tra điều kiện nước/áp lực lỗ rỗng và nguồn DST theo hồ sơ.','Thiếu chỉ tiêu một lớp phải hỏi hoặc nêu rõ phương án đã được chấp nhận.')
prompt(d,'Lệnh ghép chỉ tiêu','“Đọc [SLTT], ghép chỉ tiêu cho các lớp của [mặt cắt]. Báo nguồn từng ô và ý nghĩa Co. Trong hồ sơ này [nếu đã xác nhận: Co=Su]. Không áp quy ước sang lớp/file khác khi chưa xác minh.”')
sec(d,'5. Đọc DXF và phân biệt các đường')
steps(d,'Đọc danh sách layer và loại đối tượng.','Xác định đường tự nhiên, thiết kế, ranh giới khảo sát và bản vẽ phụ bằng nhãn/nguồn.','Kiểm tra đơn vị, tọa độ, block, đường cong/bulge và nhiều mặt cắt.','Đối chiếu cao độ/chiều rộng với BTH và đề.','Lưu bảng đối tượng/layer/handle đã dùng; không thay hình thực bằng hình thang đơn giản.')
prompt(d,'Lệnh đọc DXF','“Kiểm tra [DXF], liệt kê layer/handle và đường tự nhiên–thiết kế của [mặt cắt]. Giữ điểm gãy/đường cong đúng nguồn, báo đơn vị và phép chuyển tọa độ nếu có. Chưa dựng vùng đất khi chưa xác nhận đúng đường.”')
sec(d,'6. Dựng địa tầng và vùng đắp/đào','Nếu BTH cho bề dày đứng không đổi và người dùng chấp nhận lớp song song đường tự nhiên: x đáy bằng x tự nhiên; y đáy lớp k bằng y tự nhiên trừ tổng bề dày lớp 1…k. Đây là dịch đứng, không phải OFFSET vuông góc.')
figure(d,'geostudio_dia_tang.png','Hình 2. Minh họa dịch đứng; không phải mặt cắt Km 134+300 hay dữ liệu để tính.')
steps(d,'Dùng đúng đường tự nhiên và bề dày BTH.','Tạo ranh giới lớp, giữ điểm gãy và khép vùng theo biên.','Vùng đắp giữa đường thiết kế và tự nhiên; kiểm tra vùng đào/giao nhau.','Nếu có ranh khảo sát thực, ưu tiên ranh đó theo đề thay vì dịch đều.','Không tự kéo dài miền hoặc bổ sung lớp nền dưới cùng ngoài dữ kiện.')
sec(d,'7. Tạo project/analysis và gán vật liệu','Mở GeoStudio 2025.1.1 có giấy phép SLOPE/W. Tạo file mới hoặc bản sao ở LamViec; GSZ mẫu chỉ tham khảo cấu trúc/trình bày, không ghi đè hay sao kết quả để giả là Solve mới.')
steps(d,'Chọn analysis SLOPE/W phù hợp bài; tra chức năng thực tế trên phiên bản đang dùng.','Nhập/vẽ hình học đã chuẩn hóa, kiểm tra vùng khép kín.','Tạo vật liệu theo SLTT và gán từng vùng; phân biệt vật liệu thư viện với vật liệu đã gán.','Đọc lại bảng vùng–vật liệu–nguồn để đối chiếu.','Lưu project và kiểm tra mở lại được.')
prompt(d,'Lệnh tạo mô hình','“Tạo [đường dẫn mới] theo địa tầng BTH, hình học DXF và chỉ tiêu SLTT đã đối chiếu. GSZ mẫu chỉ tham khảo cấu trúc. Đọc lại vật liệu thực gán từng vùng; không sao tải/nước/Fs của mẫu nếu không có nguồn.”')
sec(d,'8. Nước, tải và miền tìm trượt')
tbl(d,['Mục','Cần xác định'],[['Nước','Mực nước/áp lực lỗ rỗng và nguồn; không dùng nước của GSZ mẫu mặc định'],['Tải','Vị trí, cường độ, đơn vị và giai đoạn'],['Phương pháp','Theo đề/tiêu chuẩn và khả năng chương trình'],['Miền tìm trượt','Kiểu tìm kiếm, giới hạn vào/ra và độ bao phủ vùng nguy hiểm'],['Tình huống','Giai đoạn thi công, xử lý nền và các tổ hợp phải so sánh']])
sec(d,'9. Kiểm tra trước Solve')
steps(d,'Đúng tên project/analysis/mặt cắt.','Đủ vùng đất, vật liệu gán và đơn vị.','Đúng nước, tải, phương pháp và miền tìm trượt.','Không có dữ kiện bị thay bởi mẫu hoặc giá trị đoán.','Lưu ảnh/bảng mô hình đầu vào và chỉ chạy khi các khác biệt quan trọng được xử lý.')
prompt(d,'Lệnh kiểm tra trước tính','“Lập checklist vùng đất, vật liệu, nước/tải và miền tìm trượt theo nguồn. Nêu phần chưa rõ, không đổi dữ liệu để ép Fs giống mẫu. Chỉ Solve đúng analysis mới đã kiểm tra.”')
sec(d,'10. Solve và đọc Fs mới')
steps(d,'Ghi project, analysis, thời điểm và đường dẫn kết quả mới.','Thực hiện Solve, theo dõi hoàn tất/lỗi/cảnh báo.','Đọc kết quả của phiên mới; kết quả sẵn trong GSZ mẫu phải ghi là kết quả đã lưu.','Tổng hợp Fs nhỏ nhất hợp lệ và mặt trượt tương ứng; kiểm tra các mặt không hợp lệ và phạm vi tìm kiếm.','Ghi phương pháp, nước/tải, trạng thái và hình kết quả có nguồn.','Thử độ nhạy miền tìm trượt/đầu vào nếu mục tiêu thẩm tra yêu cầu.')
prompt(d,'Lệnh báo cáo kết quả','“Đọc kết quả mới của analysis [tên]. Báo Fs, mặt trượt nguy hiểm, phương pháp, phạm vi tìm trượt, nước/tải, số trường hợp hợp lệ nếu có, trạng thái tính và nguồn file. Không dùng Fs dương/hữu hạn làm bằng chứng duy nhất là bài đã đúng.”')
common_end(d,'GeoStudio 2025.1.1 / SLOPE/W');d.save(OUT/'09_GeoStudio_2025_1_1.docx')
with zipfile.ZipFile(OUT/'Ba_huong_dan_PLAXIS_GeoStudio.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in ('07_PLAXIS_2D_2024_2.docx','08_PLAXIS_3D_2024.docx','09_GeoStudio_2025_1_1.docx'):z.write(OUT/f,arcname=f)
print('Created three illustrated Word guides and ZIP')
